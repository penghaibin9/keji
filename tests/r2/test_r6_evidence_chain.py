"""R6 procurement 8.5/10.5/10.7/12.5 regression, synthetic SQLite only."""
import io,json,uuid
from copy import deepcopy
from datetime import date
import pytest
from sqlalchemy import select
from openpyxl import load_workbook
from app.models import Assessment,Rule,Period,User,Org,Attachment,Qualification,Task,TaskAttempt
from app.core import digest
from app.services import facts
from app.worker import run_once
from app.grants import sql, APPEND_ONLY_TABLES
from conftest import auth,post
from test_api import make_assessment,qual,png,REASON


def ready(c):
    x=make_assessment(c)
    for who in ('teacher','depthead','collegehead','schooladmin'):
        auth(c,who)
        r=post(c,f"/api/assessments/{x['id']}/actions",{'action':'confirm','expected_version':x['version'],**REASON})
        assert r.status_code==200,r.text;x=r.json()
    return x


def archive(c):
    x=ready(c)
    r=post(c,f"/api/assessments/{x['id']}/actions",{'action':'archive','expected_version':x['version'],**REASON})
    assert r.status_code==200,r.text
    return r.json()


def approved_qualification(c):
    auth(c,'teacher');q=qual(c)
    r=c.post(f"/api/attachments/qualification/{q['id']}",files={'file':('当期证据.png',png(),'image/png')},
        data={'reason':'隔离证据上传','expected_version':q['version']},headers={'Idempotency-Key':str(uuid.uuid4())})
    assert r.status_code==200,r.text;attachment=r.json()
    r=post(c,f"/api/qualifications/{q['id']}/actions",{'action':'submit','expected_version':2,**REASON});assert r.status_code==200,r.text
    auth(c,'depthead');r=post(c,f"/api/qualifications/{q['id']}/actions",{'action':'approve','expected_version':3,**REASON});assert r.status_code==200,r.text
    return q,attachment


def export(c,app,kind,filters=None):
    r=post(c,'/api/exports',{'kind':kind,'filters':filters or {},'confirmed':kind=='handover'});assert r.status_code==200,r.text
    task=r.json();r=run_once(app.state.settings);assert r['status']=='DONE',r
    r=c.get(f"/api/tasks/{task['id']}/download");assert r.status_code==200,r.text
    return r.content


def test_archive_freezes_identity_period_rule_and_final_event(client2,app2):
    c=client2;x=archive(c);before=c.get(f"/api/assessments/{x['id']}/print").json()
    with app2.state.sessions.begin() as db:
        db.get(User,'teacher').name='更名后的教师'
        db.get(Org,'dept1').name='更名后的教研室'
        db.get(Org,'school').name='更名后的学校'
        db.get(Period,'period2026').name='更名后的周期'
        db.get(Rule,x['rule_id']).name='更名后的规则'
    for suffix in ('','/print'):
        r=c.get(f"/api/assessments/{x['id']}"+suffix);assert r.status_code==200,r.text;d=r.json()
        assert d['teacher']['name']=='测试教师一' and d['identity_frozen']
        assert d['period']==before['period'] and d['rule']==before['rule']
        assert d['organization_chain']==before['organization_chain']
        assert d['events'][-1]['action']=='考核归档' and d['events'][-1]['actor_id']=='schooladmin'
        assert d['assessment']['archive_sha256']==x['archive_sha256']
    proof=c.get(f"/api/assessments/{x['id']}/evidence");assert proof.status_code==200
    assert digest(proof.json()['archive'])==proof.json()['archive_sha256']
    assert 'storage_key' not in proof.text
    wb=load_workbook(io.BytesIO(export(c,app2,'assessments')),read_only=True)
    rows=list(wb.active.values);wb.close()
    assert rows[1][1]=='测试教师一' and isinstance(rows[1][3],(int,float)) and rows[1][3]==100


@pytest.mark.parametrize('mutation',['hash','score','detail','payload','source'])
def test_corrupted_archive_fails_detail_print_evidence_and_export(client2,app2,mutation):
    c=client2;x=archive(c)
    with app2.state.sessions.begin() as db:
        obj=db.get(Assessment,x['id'])
        if mutation=='hash':obj.archive_sha256='0'*64
        elif mutation=='score':obj.score=1
        elif mutation=='detail':obj.detail=[]
        elif mutation=='source':obj.source_snapshot={}
        else:obj.archive={**obj.archive,'reason':'静默改写'}
    for suffix in ('','/print','/evidence'):
        assert c.get(f"/api/assessments/{x['id']}"+suffix).status_code==409
    r=post(c,'/api/exports',{'kind':'assessments','filters':{},'confirmed':False})
    # Depending on validation location, reject before enqueue or produce a failed task.
    if r.status_code==200:
        assert run_once(app2.state.settings)['status']=='FAILED'
        assert c.get(f"/api/tasks/{r.json()['id']}/download").status_code==409
    else:assert r.status_code==409


def test_old_valid_archive_is_explicitly_unfrozen_not_silently_backfilled(client2,app2):
    c=client2;x=archive(c)
    with app2.state.sessions.begin() as db:
        obj=db.get(Assessment,x['id']);old=deepcopy(obj.archive)
        for key in ('schema_version','teacher','organization_chain','archived_by_name'):old.pop(key,None)
        obj.archive=old;obj.archive_sha256=digest(old)
        db.get(User,'teacher').name='当前姓名不能补历史'
    r=c.get(f"/api/assessments/{x['id']}/print");assert r.status_code==200
    assert r.json()['identity_frozen'] is False and '未冻结' in r.json()['teacher']['name']
    assert '当前姓名不能补历史' not in r.text


def test_process_print_label_and_no_proof_before_archive(client2):
    c=client2;x=make_assessment(c)
    d=c.get(f"/api/assessments/{x['id']}/print").json()
    assert not d['immutable'] and d['document_status']=='DRAFT_NOT_FINAL'
    assert c.get(f"/api/assessments/{x['id']}/evidence").status_code==409


def test_archival_attachment_uses_frozen_reference_and_live_permissions(client2,app2):
    c=client2;q,a=approved_qualification(c);x=archive(c)
    path=f"/api/assessments/{x['id']}/attachments/{a['id']}/download"
    auth(c,'teacher');assert c.get(path).content==png()
    with app2.state.sessions.begin() as db:
        # Simulate later replacement of current association, not historical bytes.
        db.get(Qualification,q['id']).attachment_ids=[]
    assert c.get(path).content==png()
    assert c.get(path.replace(a['id'],'not-associated')).status_code==404
    auth(c,'otherteacher');assert c.get(path).status_code==403
    auth(c,'depthead');assert c.get(path).status_code==200
    with app2.state.sessions.begin() as db:db.get(User,'teacher').org_id='dept2'
    assert c.get(path).status_code==403
    auth(c,'schooladmin');assert c.get(path).status_code==200


def test_archive_refuses_missing_evidence_and_rolls_back(client2,app2):
    c=client2;q,a=approved_qualification(c);x=ready(c)
    with app2.state.sessions() as db:key=db.get(Attachment,a['id']).storage_key
    (app2.state.settings.storage_dir/key).unlink()
    r=post(c,f"/api/assessments/{x['id']}/actions",{'action':'archive','expected_version':x['version'],**REASON})
    assert r.status_code==503,r.text
    with app2.state.sessions() as db:
        obj=db.get(Assessment,x['id']);assert obj.state=='READY' and obj.version==x['version'] and not obj.archive


def test_archive_download_uses_original_hash_not_rewritten_attachment_hash(client2,app2):
    from hashlib import sha256
    c=client2;q,a=approved_qualification(c);x=archive(c)
    with app2.state.sessions.begin() as db:
        item=db.get(Attachment,a['id']);key=item.storage_key
        changed=b'!'*item.size;(app2.state.settings.storage_dir/key).write_bytes(changed)
        item.sha256=sha256(changed).hexdigest()
    r=c.get(f"/api/assessments/{x['id']}/attachments/{a['id']}/download")
    assert r.status_code==503,r.text


def test_next_period_cannot_use_tampered_archived_score(client2,app2):
    from fastapi import HTTPException
    x=archive(client2)
    with app2.state.sessions.begin() as db:
        db.get(Assessment,x['id']).score=999
        period=Period(id='p2027',name='下期',start_date=date(2027,1,1),end_date=date(2027,12,31));db.add(period);db.flush()
        with pytest.raises(HTTPException) as exc:facts(db,'teacher',period)
        assert exc.value.status_code==409


@pytest.mark.parametrize('change',['disabled','role','organization'])
def test_export_worker_records_failure_when_authorization_changes(client2,app2,change):
    c=client2;auth(c,'schooladmin')
    task=post(c,'/api/exports',{'kind':'records','filters':{},'confirmed':False}).json()
    with app2.state.sessions.begin() as db:
        if change=='disabled':db.get(User,'schooladmin').active=False
        elif change=='role':db.get(User,'schooladmin').role_id='role-college'
        else:db.get(User,'teacher').org_id='dept2';db.get(User,'schooladmin').auth_version+=1
    assert run_once(app2.state.settings)['status']=='FAILED'
    with app2.state.sessions() as db:
        t=db.get(Task,task['id']);assert t.attempts==1 and t.lease_token is None and not t.result
        assert db.scalar(select(TaskAttempt).where(TaskAttempt.task_id==t.id)).result=='ERROR'


def test_runtime_mysql_grants_make_evidence_append_only_not_installer():
    statements=sql()
    assert len(APPEND_ONLY_TABLES)==10
    for table in APPEND_ONLY_TABLES:
        lines=[x for x in statements.splitlines() if f'`{table}`' in x]
        assert len(lines)==1 and 'GRANT SELECT, INSERT ON ' in lines[0] and 'UPDATE' not in lines[0] and 'DELETE' not in lines[0]
    assert "GRANT SELECT ON `nondegree`.`installation_provenance` TO 'nondegree_app'@'%';" in statements


def test_archived_teacher_list_uses_frozen_names_and_period(client2,app2):
    x=archive(client2)
    with app2.state.sessions.begin() as db:
        db.get(User,'teacher').name='现在改名'
        db.get(Period,'period2026').name='现在改周期名'
    r=client2.get('/api/assessments');assert r.status_code==200
    row=r.json()['items'][0]
    assert row['teacher_display_name']=='测试教师一'
    assert row['period_display_name']=='2026隔离测试周期'


def test_organization_export_uses_frozen_target_not_current_name(client2,app2):
    from test_r3_org_assessment import make_rule,calculate,action
    c=client2;auth(c,'schooladmin');make_rule(c,'dept','dept1');x=calculate(c,'dept','dept1')
    auth(c,'depthead');x=action(c,x,'review').json()
    auth(c,'schooladmin');x=action(c,x,'school_audit').json();r=action(c,x,'archive');assert r.status_code==200,r.text
    with app2.state.sessions.begin() as db:db.get(Org,'dept1').name='新部门名'
    wb=load_workbook(io.BytesIO(export(c,app2,'organization_assessments')),read_only=True)
    rows=list(wb.active.values);wb.close()
    assert rows[1][0]=='测试教研室一'
