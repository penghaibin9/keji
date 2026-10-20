"""10.9 progress quick filters, 10.7 tied ranking + printable ranking, 5.2 qualification summary, 7.4/7.5 comparison."""
from datetime import date,timedelta
from decimal import Decimal
from sqlalchemy import select
from app.models import Audit,Assessment,TrainingRecord,Qualification
from conftest import auth,post
from test_api import make_assessment,REASON,rule,activate
from test_r3_org_assessment import make_rule,calculate,action
from test_social_training import proof


def test_teacher_progress_filters_follow_role_stage(client2):
    x=make_assessment(client2);ident=x['id']
    auth(client2,'teacher');assert client2.get('/api/assessments?progress=pending').json()['total']==1
    r=post(client2,f'/api/assessments/{ident}/actions',{'action':'confirm','expected_version':x['version'],**REASON});v=r.json()['version']
    assert client2.get('/api/assessments?progress=pending').json()['total']==0 and client2.get('/api/assessments?progress=done').json()['total']==1
    auth(client2,'depthead');assert client2.get('/api/assessments?progress=pending').json()['total']==1 and client2.get('/api/assessments?progress=done').json()['total']==0
    r=post(client2,f'/api/assessments/{ident}/actions',{'action':'return','expected_version':v,**REASON});assert r.status_code==200
    assert client2.get('/api/assessments?progress=returned').json()['total']==1 and client2.get('/api/assessments?progress=pending').json()['total']==0
    auth(client2,'collegehead');assert client2.get('/api/assessments?progress=pending').json()['total']==0
    assert client2.get('/api/assessments?progress=everything').status_code==422


def test_tied_competition_rank_and_printable_ranking(client2,app2):
    auth(client2,'schooladmin');x=rule(client2);assert activate(client2,x).status_code==200
    r=post(client2,'/api/assessments/calculate',{'period_id':'period2026','teacher_ids':['teacher','otherteacher'],**REASON});assert r.status_code==200,r.text
    items=client2.get('/api/assessments?period_id=period2026').json()['items']
    assert [i['rank'] for i in items]==[1,1] and [i['scope_rank'] for i in items]==[1,2]  # both 10 hours -> 100
    p=client2.get('/api/assessment-ranking/print?period_id=period2026');assert p.status_code==200,p.text
    body=p.json();assert body['rank_method']=='competition' and [i['rank'] for i in body['items']]==[1,1] and body['items'][0]['teacher_name']
    with app2.state.sessions() as db:assert db.scalar(select(Audit).where(Audit.action=='打印教师考核排名表'))
    auth(client2,'teacher');assert client2.get('/api/assessment-ranking/print').json()['total']==1


def test_org_returned_flag_progress_and_print(client2,app2):
    auth(client2,'schooladmin');rule_obj=make_rule(client2,'dept','dept1');obj=calculate(client2,'dept','dept1')
    auth(client2,'depthead');assert client2.get('/api/organization-assessments?progress=pending').json()['total']==1
    obj=action(client2,obj,'review').json();assert client2.get('/api/organization-assessments?progress=done').json()['total']==1
    auth(client2,'schooladmin');assert client2.get('/api/organization-assessments?progress=pending').json()['total']==1
    obj=action(client2,obj,'return').json()
    listed=client2.get('/api/organization-assessments?progress=returned').json();assert listed['total']==1 and listed['items'][0]['returned'] is True
    auth(client2,'depthead');assert client2.get('/api/organization-assessments').json()['items'][0]['returned'] is True
    assert client2.get('/api/organization-assessments/'+obj['id']).json()['returned'] is True
    proof(client2,'org_assessment',obj['id'],obj['version'])
    listed=client2.get('/api/organization-assessments?progress=returned').json()
    assert listed['total']==1 and listed['items'][0]['returned'] is True
    detail=client2.get('/api/organization-assessments/'+obj['id']).json()
    assert detail['returned'] is True and detail['events'][-1]['action']=='上传组织考核佐证材料'
    from io import BytesIO
    from openpyxl import load_workbook
    from test_r6_evidence_chain import export
    filters={'period_id':'period2026','target_kind':'dept','progress':'returned'}
    wb=load_workbook(BytesIO(export(client2,app2,'organization_assessments',filters)),read_only=True)
    assert len(list(wb.active.values))==2;wb.close()
    response=action(client2,obj,'review');assert response.status_code==200,response.text
    assert client2.get('/api/organization-assessments/'+obj['id']).json()['returned'] is False
    assert client2.get('/api/organization-assessments?progress=returned').json()['total']==0
    wb=load_workbook(BytesIO(export(client2,app2,'organization_assessments',filters)),read_only=True)
    assert len(list(wb.active.values))==1;wb.close()
    auth(client2,'schooladmin')
    p=client2.get('/api/organization-assessment-ranking/print',params={'period_id':'period2026','target_kind':'dept','rule_id':rule_obj['id']});assert p.status_code==200,p.text
    assert p.json()['items'][0]['rank']==1 and p.json()['rule_version']==rule_obj['version_no']
    with app2.state.sessions() as db:assert db.scalar(select(Audit).where(Audit.action=='打印组织考核排名表'))


def test_qualification_summary_distribution(client2,app2):
    auth(client2,'teacher')
    for no,state,expires in [('A','SUBMITTED',None),('B','RETURNED',None),('C','APPROVED',date.today()+timedelta(days=5)),('D','APPROVED',None),('E','APPROVED',date.today()-timedelta(days=1))]:
        with app2.state.sessions.begin() as db:
            db.add(Qualification(teacher_id='teacher',category='职业资格',name='证'+no,certificate_no=no,issuer='x',obtained=date(2020,1,1),expires=expires,state=state))
    d={x['key']:x['count'] for x in client2.get('/api/qualification-summary').json()['items']}
    assert d=={'SUBMITTED':1,'RETURNED':1,'EXPIRING':1,'VALID':1,'EXPIRED':1,'DRAFT':0}
    # 列表的“即将到期”“已过期”页签与概况同一口径。
    for key,names in [('EXPIRING',['证C']),('EXPIRED',['证E']),('APPROVED',['证C','证D','证E'])]:
        listed=client2.get('/api/qualifications?state='+key).json();assert sorted(x['name'] for x in listed['items'])==names,key
    assert client2.get('/api/qualifications?state=VALIDISH').status_code==422
    auth(client2,'otherteacher');assert client2.get('/api/qualification-summary').json()['total']==0


def test_comparison_uses_same_stats_source(client2,app2):
    auth(client2,'schooladmin')
    with app2.state.sessions.begin() as db:
        db.add(TrainingRecord(external_key='old1',teacher_id='teacher',event_date=date(2025,9,1),project_code='P1',project='p',course_code='C1',course='c',class_code='CL',class_name='n',hours=1,participants=2,income=50))
    assert client2.get('/api/statistics/comparison').json()['available'] is False
    r=client2.get('/api/statistics/comparison',params={'start':'2026-07-01','end':'2026-12-31'});assert r.status_code==200,r.text;c=r.json()
    assert c['current']['income']==client2.get('/api/statistics',params={'start':'2026-07-01','end':'2026-12-31'}).json()['income']=='300.00'
    assert c['windows']['previous']=={'start':'2025-12-29','end':'2026-06-30','days':184} and c['mom']['income'] is None
    assert c['year_ago']['income']=='50.00' and c['yoy']['income']=='500.00'
    s=client2.get('/api/statistics').json();assert 'assessment_average' in s['colleges'][0]
