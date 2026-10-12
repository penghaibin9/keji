"""采购2.3/6.5、3.4、9.3/9.5/9.6、10.3、14.6：七处审计问题的隔离回归。"""
from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import json
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from conftest import auth, post, put
from test_social_training import setup_social, recognized, detail, action, proof, row, import_rows
from test_point_review import review, sent_review
from app.core import serial
from app.models import Org, User, Setting, Attachment, PointEntry, Qualification, WorkflowEvent
from app.social_models import SocialProject, SocialUnit, SocialPolicy, SocialClaim, SocialAnnual, SocialResource
from app.reference_data import TrainingProjectMaster, TrainingCourseMaster
from app import social_training as st
from app.worker import run_once


def sibling(db):
    db.add(Org(id='sibling',code='sibling',name='同院外室',kind='dept',parent_id='college1'));db.flush()
    db.get(User,'otherteacher').org_id='sibling'


@pytest.mark.parametrize('username',['teacher','depthead'])
def test_2_3_course_resources_share_scope_in_all_read_paths(setup_social,app2,client2,username):
    with app2.state.sessions.begin() as db:
        sibling(db)
        db.add(TrainingProjectMaster(id='unrelated',code='UNRELATED',name='无关项目'));db.flush()
        db.add(TrainingCourseMaster(code='UNRELATED-C',name='无关课程',project_code='UNRELATED'))
        for ident,course,teacher in [('hidden','UNRELATED-C','otherteacher'),('same-course-other','SOC-C1','otherteacher'),('own','SOC-C1','teacher'),('shared','SOC-C1',None)]:
            db.add(SocialResource(id=ident,org_id='college1',kind='course_resource',code=ident,name=ident,data={'teacher_id':teacher,'course_code':course,'resource_type':'video','duration_minutes':'12'}))
    auth(client2,'schooladmin')
    for ident in ('hidden','own'):
        proof(client2,'social_resource',ident,detail(client2,'social_resource',ident)['item']['version'])
    with app2.state.sessions() as db:
        files={x.object_id:x.id for x in db.scalars(select(Attachment).where(Attachment.object_type=='social_resource'))}
    auth(client2,username)
    workspace=client2.get('/api/social/workspace?section=courses').json()
    assert 'UNRELATED-C' not in [c['code'] for c in workspace['courses']]
    assert {x['id'] for x in workspace['resources']}=={'own','shared'}
    assert set(workspace['course_stats'])=={'SOC-C1'} and workspace['course_stats']['SOC-C1']['resources']==2
    assert client2.get('/api/social/courses/UNRELATED-C/resources').json()['total']==0
    assert client2.get('/api/social/courses/SOC-C1/resources').json()['total']==2
    assert client2.get('/api/social/courses/SOC-C1/resources?page=2').json()['items']==[]
    stats=client2.get('/api/social/resource-stats')
    if username=='teacher':assert stats.status_code==403
    else:assert stats.json()['items'][0]['resources']==2 and Decimal(stats.json()['items'][0]['minutes'])==Decimal('24')
    for ident in ('hidden','same-course-other'):
        for suffix in ('','/export'):
            assert client2.get('/api/social/objects/social_resource/'+ident+suffix).status_code==403
    assert client2.get('/api/attachments/'+files['hidden']+'/download').status_code==403
    assert client2.get('/api/attachments/'+files['own']+'/download').status_code==200
    assert detail(client2,'social_resource','own')['item']['id']=='own'
    for user in ('collegehead','schooladmin'):
        auth(client2,user)
        assert client2.get('/api/social/courses/UNRELATED-C/resources').json()['total']==1
        assert client2.get('/api/attachments/'+files['hidden']+'/download').status_code==200


def test_2_3_6_5_dept_project_export_and_frozen_claims_hide_other_department(recognized,app2,client2):
    claim_id,_=recognized
    with app2.state.sessions.begin() as db:
        sibling(db);claim=db.get(SocialClaim,claim_id);pid=claim.project_id
        other=row('other-unit',1);other['teachers'][0]['teacher_id']='otherteacher';other['learners'][0]['reference']='SIBLING-SECRET'
        db.add(SocialUnit(project_id=pid,external_key='other-unit',data=other));db.flush()
        mixed=db.scalar(select(SocialUnit).where(SocialUnit.project_id==pid,SocialUnit.external_key!='other-unit'))
        mixed.data={**mixed.data,'teachers':mixed.data['teachers']+[{'teacher_id':'otherteacher','staff_type':'internal','role':'instructor'}]}
        trial=deepcopy(claim.snapshot['trial']);trial['basis']['facts']=st.project_facts(db,db.get(SocialProject,pid))
        claim.snapshot=st.freeze(trial);original=deepcopy(claim.snapshot)
        entries=db.scalars(select(PointEntry).where(PointEntry.teacher_id=='teacher',PointEntry.kind=='SOCIAL')).all()
        for entry in entries:entry.source_snapshot={**entry.source_snapshot,'recognition':deepcopy(original)}
        entry_id=entries[0].id;amount=sum(x.amount for x in entries)
        file_id=db.scalar(select(Attachment.id).where(Attachment.object_type=='social_project',Attachment.object_id==pid))
    auth(client2,'depthead')
    d=detail(client2,'social_project',pid)
    assert d['readonly'] and d['facts']['units'] and not d['events'] and not d['attachments']
    assert d['item']['closeout']=={} and 'expected_income' not in d['item']['data']
    assert all(not u['data']['learners'] and all(t['teacher_id']=='teacher' for t in u['data']['teachers']) for u in d['facts']['units'])
    assert client2.get('/api/attachments/'+file_id+'/download').status_code==403
    for url in ('/api/social/objects/social_project/'+pid,'/api/social/objects/social_claim/'+claim_id,'/api/social/workspace?section=recognition','/api/points/'+entry_id,'/api/points?period_id=period2026'):
        r=client2.get(url);assert r.status_code==200,r.text
        assert all(secret not in r.text for secret in ('SIBLING-SECRET','otherteacher','finance_reference'))
    exported=client2.get('/api/social/objects/social_project/'+pid+'/export');assert exported.status_code==200
    with BytesIO(exported.content) as stream:
        book=load_workbook(stream,read_only=True)
        assert all('SIBLING-SECRET' not in str(value) for sheet in book for values in sheet.iter_rows(values_only=True) for value in values)
        summary=list(book['参与摘要'].iter_rows(values_only=True))
        assert len(summary)==2 and summary[1][0]=='unit1' and summary[1][-1]=='主讲' and summary[1][-2]
        book.close()
    auth(client2,'collegehead');assert 'SIBLING-SECRET' in client2.get('/api/social/objects/social_project/'+pid).text
    assert client2.get('/api/attachments/'+file_id+'/download').status_code==200
    with app2.state.sessions() as db:
        assert db.get(SocialClaim,claim_id).snapshot==original
        assert sum(db.scalars(select(PointEntry.amount).where(PointEntry.teacher_id=='teacher',PointEntry.kind=='SOCIAL')))==amount


@pytest.mark.parametrize('returned',[False,True])
def test_9_5_void_unbooked_correction_then_restart_review(recognized,app2,client2,returned):
    claim_id,_=recognized;ident=sent_review(client2,claim_id)
    auth(client2,'teacher');fix=action(client2,'point-reviews','social_resource',ident,'agree').json()['item']['data']['correction_id']
    auth(client2,'collegehead')
    assert action(client2,'corrections','social_resource',fix,'void').status_code==403
    if returned:assert action(client2,'corrections','social_resource',fix,'return').status_code==200
    auth(client2,'depthead');item=detail(client2,'social_resource',fix)['item']
    fields={k:item['data'][k] for k in ('claim_id','period_id','teaching','auxiliary','decision_reference','reason','basis')}
    changed=put(client2,'/api/social/corrections/'+fix,{**fields,'teaching':'-2','expected_version':item['version']})
    assert changed.status_code in (403,409)
    assert review(client2,claim_id).status_code==409
    result=action(client2,'corrections','social_resource',fix,'void');assert result.status_code==200,result.text
    assert result.json()['item']['state']=='VOID'
    assert detail(client2,'social_resource',ident)['review_context']['correction_state']=='VOID'
    assert action(client2,'corrections','social_resource',fix,'submit').status_code==409
    with app2.state.sessions() as db:
        assert not st.open_point_reviews(db,'teacher')
        assert sum(db.scalars(select(PointEntry.amount).where(PointEntry.teacher_id=='teacher',PointEntry.kind=='SOCIAL')))==Decimal('1.56')
        assert db.scalar(select(WorkflowEvent.id).where(WorkflowEvent.object_id==fix,WorkflowEvent.action=='积分更正：已作废'))
    new=sent_review(client2,claim_id);auth(client2,'teacher')
    new_fix=action(client2,'point-reviews','social_resource',new,'agree').json()['item']['data']['correction_id']
    for user in ('collegehead','reviewer','vice','schooladmin'):
        auth(client2,user);r=action(client2,'corrections','social_resource',new_fix,'approve');assert r.status_code==200,r.text
    auth(client2,'depthead');assert action(client2,'corrections','social_resource',new_fix,'void').status_code==409
    with app2.state.sessions() as db:
        assert sum(db.scalars(select(PointEntry.amount).where(PointEntry.teacher_id=='teacher',PointEntry.kind=='SOCIAL')))==Decimal('1.00')


@pytest.mark.parametrize('window',[1,5,30])
def test_3_4_8_4_qualification_reminder_uses_same_config(client2,app2,window):
    today=st.today()
    with app2.state.sessions.begin() as db:
        setting=db.get(Setting,'qualification');setting.value={**setting.value,'expiry_days':window}
        for days in (-1,window,window+1):
            db.add(Qualification(id='expiry-'+str(days),teacher_id='teacher',category='职业资格',name='隔离资质',certificate_no='AUDIT'+str(days),issuer='隔离机构',obtained=today-timedelta(days=300),expires=today+timedelta(days=days),state='APPROVED'))
    auth(client2,'teacher')
    todos=client2.get('/api/todos?type=qualification&size=100').json()
    mine=client2.get('/api/social/mine?year='+str(today.year)).json()
    expected={'expiry--1','expiry-'+str(window)}
    assert {x['id'] for x in todos['items'] if x['id'].startswith('expiry-')}==expected
    assert {x['id'] for x in mine['reminders'] if x['id'].startswith('expiry-')}==expected


@pytest.mark.parametrize('mode',['sole_dean','sole_secretary','inactive_alternate','out_of_scope_alternate'])
def test_10_3_send_requires_eligible_nonrecused_staff(recognized,app2,client2,mode):
    claim_id,_=recognized
    creator='collegehead' if mode!='sole_secretary' else 'depthead'
    with app2.state.sessions.begin() as db:
        setting=db.get(Setting,'social_training_duties')
        setting.value={**setting.value,'secretary':['depthead','schooladmin']}
        if mode=='sole_secretary':setting.value={**setting.value,'secretary':['depthead']}
        elif mode=='inactive_alternate':
            db.get(User,'schooladmin').active=False;setting.value={**setting.value,'dean':['collegehead','schooladmin']}
        elif mode=='out_of_scope_alternate':setting.value={**setting.value,'dean':['collegehead','otherdept']}
    auth(client2,creator);draft=review(client2,claim_id).json()['item'];proof(client2,'social_resource',draft['id'],draft['version'])
    d=detail(client2,'social_resource',draft['id']);assert d['workflow_blockers']
    result=action(client2,'point-reviews','social_resource',draft['id'],'send');assert result.status_code==409 and '替代人员' in result.text
    auth(client2,'teacher');assert client2.get('/api/social/objects/social_resource/'+draft['id']).status_code==404
    with app2.state.sessions.begin() as db:
        db.get(User,'schooladmin').active=True
        setting=db.get(Setting,'social_training_duties');setting.value={**setting.value,'dean':['collegehead','schooladmin'],'secretary':['depthead','schooladmin']}
    auth(client2,creator);assert not detail(client2,'social_resource',draft['id'])['workflow_blockers']
    assert action(client2,'point-reviews','social_resource',draft['id'],'send').status_code==200
    auth(client2,'teacher');result=action(client2,'point-reviews','social_resource',draft['id'],'explain');assert result.status_code==200
    auth(client2,creator);assert action(client2,'point-reviews','social_resource',draft['id'],'uphold').status_code==403
    auth(client2,'schooladmin');assert action(client2,'point-reviews','social_resource',draft['id'],'uphold').status_code==200


@pytest.mark.parametrize('missing',['file','metadata'])
def test_9_3_9_6_archive_verifies_inherited_evidence_atomically(recognized,app2,client2,missing):
    ident=sent_review(client2,recognized[0]);auth(client2,'teacher')
    fix=action(client2,'point-reviews','social_resource',ident,'agree').json()['item']['data']['correction_id']
    for user in ('collegehead','reviewer','vice'):
        auth(client2,user);assert action(client2,'corrections','social_resource',fix,'approve').status_code==200
    with app2.state.sessions.begin() as db:
        attachment=db.scalar(select(Attachment).where(Attachment.object_id==ident))
        file_id=attachment.id;original_key=attachment.storage_key;original_object=attachment.object_id
        if missing=='file':attachment.storage_key='0'*64
        else:attachment.object_id='missing-review'
    auth(client2,'schooladmin');r=action(client2,'corrections','social_resource',fix,'approve');assert r.status_code in (409,503),r.text
    assert detail(client2,'social_resource',fix)['item']['state']=='APPROVED'
    with app2.state.sessions.begin() as db:
        assert sum(db.scalars(select(PointEntry.amount).where(PointEntry.teacher_id=='teacher',PointEntry.kind=='SOCIAL')))==Decimal('1.56')
        attachment=db.get(Attachment,file_id);attachment.storage_key=original_key;attachment.object_id=original_object
    assert action(client2,'corrections','social_resource',fix,'approve').json()['item']['state']=='ARCHIVED'


def test_14_6_handover_contains_all_social_rows_and_attachment_links(recognized,app2,client2):
    claim_id,policy_id=recognized;sent_review(client2,claim_id)
    with app2.state.sessions.begin() as db:
        db.add(SocialAnnual(id='handover-annual',org_id='college1',period_id='period2026',policy_id=policy_id,owner_id='collegehead',data={'self_report':'隔离年度自评','zero':'0'}))
    auth(client2,'schooladmin');task=post(client2,'/api/exports',{'kind':'handover','filters':{},'confirmed':True}).json()
    assert run_once(app2.state.settings)['status']=='DONE'
    response=client2.get('/api/tasks/'+task['id']+'/download');assert response.status_code==200
    with ZipFile(BytesIO(response.content)) as archive,app2.state.sessions() as db:
        manifest=json.loads(archive.read('manifest.json'));tables={x['model_table']:x for x in manifest['tables']}
        for model in (SocialProject,SocialUnit,SocialPolicy,SocialClaim,SocialAnnual,SocialResource):
            table=tables[model.__tablename__];expected=db.scalars(select(model)).all()
            exported=[json.loads(line) for line in archive.read(table['typed_path']).decode().splitlines()]
            assert len(exported)==table['rows']==len(expected)>0
            assert {x['id']:x for x in exported}=={x.id:{k:serial(getattr(x,k)) for k in table['columns']} for x in expected}
            assert table['csv_path'] in archive.namelist()
        attachments=json.loads(archive.read('attachments_manifest.json'))
        assert any(x['object_type']=='social_resource' for x in attachments)
        for item in attachments:assert len(archive.read(item['archive_path']))==item['size']
        assert 'password_hash' not in tables['users']['columns']
