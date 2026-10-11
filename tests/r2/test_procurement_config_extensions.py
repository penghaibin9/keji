"""3.1 custom dictionary categories / extra parameters and 3.4 export display names."""
import io
from openpyxl import load_workbook
from app.worker import run_once
from conftest import auth,post,put

REASON={'reason':'自定义配置测试','basis':'合成隔离数据'}


def settings(c):
    return {x['key']:x for x in c.get('/api/settings').json()['items']}


def test_admin_adds_custom_dictionary_and_parameter_with_version_and_audit(client2):
    c=client2;auth(c,'schooladmin')
    r=post(c,'/api/settings',{'key':'dict:training_level','value':{'name':'培训层次','items':['初级','中级','高级'],'description':'项目分级'},**REASON})
    assert r.status_code==200,r.text;assert r.json()['version']==1
    r=post(c,'/api/settings',{'key':'param:max_class_size','value':{'name':'单班人数上限','type':'int','value':60},**REASON})
    assert r.status_code==200,r.text
    # Validation: duplicate, wrong value type, unknown structure, non-custom key.
    assert post(c,'/api/settings',{'key':'dict:training_level','value':{'name':'重复','items':['a']},**REASON}).status_code==409
    assert post(c,'/api/settings',{'key':'param:bad_int','value':{'name':'错误','type':'int','value':'六十'},**REASON}).status_code==422
    assert post(c,'/api/settings',{'key':'dict:dup_items','value':{'name':'重复项','items':['a','a']},**REASON}).status_code==422
    assert post(c,'/api/settings',{'key':'branding','value':{'school_name':'x'},**REASON}).status_code==422
    assert post(c,'/api/settings',{'key':'param:extra','value':{'name':'x','type':'text','value':'y','other':1},**REASON}).status_code==422
    s=settings(c);obj=s['dict:training_level']
    r=put(c,'/api/settings/dict:training_level',{'value':{'name':'培训层次','items':['初级','中级','高级','特级']},'expected_version':obj['version'],'reason':'增加特级','basis':'学校文件'})
    assert r.status_code==200,r.text;assert r.json()['version']==2
    assert put(c,'/api/settings/dict:training_level',{'value':{'name':'培训层次','items':[]},'expected_version':2,**REASON}).status_code==422
    hist=c.get('/api/settings/dict:training_level/history').json()['items']
    assert [h['action'] for h in hist]==['修改基础配置','新增学校自定义配置']
    assert hist[0]['reason']=='增加特级' and hist[0]['before']['value']['items']==['初级','中级','高级'] and hist[0]['after']['value']['items'][-1]=='特级'
    assert hist[1]['before']=={} and hist[1]['after']['key']=='dict:training_level'
    # Published to every portal user; non-admins cannot create.
    auth(c,'teacher');d=c.get('/api/dictionaries').json()
    assert d['categories']['training_level']['items'][-1]=='特级' and d['parameters']['max_class_size']['value']==60
    assert post(c,'/api/settings',{'key':'dict:other','value':{'name':'x','items':['a']},**REASON}).status_code==403
    auth(c,'depthead');assert post(c,'/api/settings',{'key':'dict:other','value':{'name':'x','items':['a']},**REASON}).status_code==403


def test_frontend_offers_custom_setting_entry():
    from pathlib import Path
    js=(Path(__file__).resolve().parents[2]/'frontend/app.js').read_text(encoding='utf-8')
    assert "'setting-create':settingCreate" in js and "api('/settings',{method:'POST'" in js


def test_excel_exports_use_configured_state_display_names(client2,app2):
    c=client2;auth(c,'teacher')
    r=post(c,'/api/qualifications',{'teacher_id':'teacher','category':'职业资格','name':'导出证书','certificate_no':'EXP-1','issuer':'机构','obtained':'2025-01-01',**REASON});assert r.status_code==200,r.text
    auth(c,'schooladmin');obj=settings(c)['business_dictionary']
    labels={**obj['value']['status_labels'],'DRAFT':'草稿（学校口径）'}
    assert put(c,'/api/settings/business_dictionary',{'value':{**obj['value'],'status_labels':labels},'expected_version':obj['version'],**REASON}).status_code==200
    task=post(c,'/api/exports',{'kind':'qualifications','filters':{},'confirmed':False}).json()
    assert run_once(app2.state.settings)['status']=='DONE'
    wb=load_workbook(io.BytesIO(c.get('/api/tasks/'+task['id']+'/download').content),read_only=True);rows=list(wb.active.values);wb.close()
    assert rows[0][8]=='状态' and rows[1][8]=='草稿（学校口径）'
    task=post(c,'/api/exports',{'kind':'audit','filters':{},'confirmed':False}).json()
    assert run_once(app2.state.settings)['status']=='DONE'
    wb=load_workbook(io.BytesIO(c.get('/api/tasks/'+task['id']+'/download').content),read_only=True);rows=list(wb.active.values);wb.close()
    assert {r[9] for r in rows[1:]}<={'成功','已拦截','异常'} and 'SUCCESS' not in {r[9] for r in rows[1:]}
