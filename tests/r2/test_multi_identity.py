"""一人多身份、组织页设负责人、缺负责人上交上一级、同一人连审标注、解锁。

隔离夹具：dept1 有 depthead、college1 有 collegehead；dept2 有 otherdept、college2 没有负责人。
"""
from datetime import datetime, timedelta
from sqlalchemy import select
from app.models import Rule, Assessment, User, WorkflowEvent, Audit
from conftest import auth, post, put


def _assessment(app2, ident, teacher, state='DEPARTMENT'):
    with app2.state.sessions.begin() as db:
        if db.get(Rule, 'mi-rule') is None:
            db.add(Rule(id='mi-rule', kind='assessment', name='多身份测试规则', version_no=1, period_id='period2026',
                config={'components': []}, fingerprint='d' * 64, state='ACTIVE', created_by='schooladmin', reason='测试'))
            db.flush()
        db.add(Assessment(id=ident, teacher_id=teacher, period_id='period2026', rule_id='mi-rule', score=80,
            detail=[], source_snapshot={}, state=state, version=1))


def _act(client, ident, action='confirm'):
    version = client.get('/api/assessments/' + ident).json()['assessment']['version']
    return post(client, f'/api/assessments/{ident}/actions', {'action': action, 'expected_version': version, 'reason': '核对无误'})


def _user(client, ident):
    return next(x for x in client.get('/api/users?size=200').json()['items'] if x['id'] == ident)


def _add_extra(client, ident, extras):
    x = _user(client, ident)
    return put(client, '/api/users/' + ident, {'name': x['name'], 'title': x['title'], 'org_id': x['org_id'], 'role_id': x['role_id'],
        'active': x['active'], 'is_admin': x['is_admin'], 'expected_version': x['version'], 'reason': '兼任调整',
        'extra_identities': extras})


def _switch(client, identity_id):
    r = post(client, '/api/auth/identity', {'identity_id': identity_id})
    assert r.status_code == 200, r.text
    me = client.get('/api/auth/me').json(); client.headers['x-csrf-token'] = me['csrf']
    return me


def test_teacher_with_dept_and_college_identities_switches_scope(client2):
    auth(client2, 'schooladmin')
    r = _add_extra(client2, 'teacher', [{'role_id': 'role-dept', 'org_id': 'dept1'}, {'role_id': 'role-college', 'org_id': 'college1'}])
    assert r.status_code == 200, r.text
    assert [e['role_kind'] for e in _user(client2, 'teacher')['extra_identities']] == ['dept', 'college']
    # 层级不匹配、重复、学校管理员作为兼任都拒绝
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-college', 'org_id': 'dept1'}]).status_code == 403
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-dept', 'org_id': 'dept1'}] * 2).status_code == 422
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-school', 'org_id': 'school'}]).status_code == 422
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-teacher', 'org_id': 'dept1'}]).status_code == 422

    me = auth(client2, 'teacher')
    assert me['role']['kind'] == 'teacher'
    options = me['identities']
    assert [x['role_kind'] for x in options] == ['teacher', 'dept', 'college'] and options[0]['current']
    dept_id = options[1]['id']
    me = _switch(client2, dept_id)
    assert me['role']['kind'] == 'dept' and me['org']['id'] == 'dept1' and me['primary_role']['kind'] == 'teacher'
    assert client2.get('/api/assessment-summary/latest').status_code == 403  # 负责人身份下不是教师本人视角
    me = _switch(client2, '')
    assert me['role']['kind'] == 'teacher'
    # 不能切到别人的身份
    assert post(client2, '/api/auth/identity', {'identity_id': 'not-mine'}).status_code == 404


def test_switch_is_audited_and_removed_identity_forces_login(client2):
    auth(client2, 'schooladmin')
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-dept', 'org_id': 'dept1'}]).status_code == 200
    me = auth(client2, 'teacher')
    _switch(client2, me['identities'][1]['id'])
    with client2.app.state.sessions.begin() as db:
        assert db.scalar(select(Audit).where(Audit.action == '切换身份', Audit.actor_id == 'teacher')) is not None
    other = client2.__class__(client2.app, raise_server_exceptions=False); other.headers['origin'] = 'http://testserver'
    auth(other, 'schooladmin')
    assert _add_extra(other, 'teacher', []).status_code == 200
    assert client2.get('/api/auth/me').status_code == 401


def test_org_page_sets_heads_and_vacancy_hands_up(client2, app2):
    _assessment(app2, 'a-other', 'otherteacher', 'COLLEGE')
    auth(client2, 'schooladmin')
    orgs = {o['id']: o for o in client2.get('/api/organizations').json()['items']}
    assert [h['user_id'] for h in orgs['dept1']['heads']] == ['depthead']
    assert orgs['college2']['heads'] == [] and 'heads' not in orgs['school']
    # college2 缺负责人：学院这一级由学校代办
    d = client2.get('/api/assessments/a-other').json()
    assert d['handler'] == {'kind': 'school', 'acting_for': '代缺位学院负责人办理'}
    todos = client2.get('/api/todos').json()['items']
    assert any(t['id'] == 'a-other' and '代缺位学院负责人办理' in t['title'] for t in todos)
    r = _act(client2, 'a-other'); assert r.status_code == 200, r.text
    with app2.state.sessions.begin() as db:
        ev = db.scalar(select(WorkflowEvent).where(WorkflowEvent.object_id == 'a-other').order_by(WorkflowEvent.created_at.desc()))
        assert ev.after['handling']['acting_for'] == '代缺位学院负责人办理'

    # 组织页给 college2 设负责人：只能选本学院老师
    cands = client2.get('/api/organizations/college2/head-candidates').json()['items']
    assert [c['id'] for c in cands] == ['otherteacher']
    assert post(client2, '/api/organizations/college2/heads', {'user_id': 'teacher', 'reason': '设负责人'}).status_code == 422
    r = post(client2, '/api/organizations/college2/heads', {'user_id': 'otherteacher', 'reason': '设负责人'})
    assert r.status_code == 200, r.text
    assert post(client2, '/api/organizations/college2/heads', {'user_id': 'otherteacher', 'reason': '重复'}).status_code == 409
    orgs = {o['id']: o for o in client2.get('/api/organizations').json()['items']}
    assert [(h['user_id'], h['primary']) for h in orgs['college2']['heads']] == [('otherteacher', False)]
    # 主身份负责人不能在组织页移除
    assert post(client2, '/api/organizations/dept1/heads/remove', {'user_id': 'depthead', 'reason': '调整'}).status_code == 409
    assert post(client2, '/api/organizations/college2/heads/remove', {'user_id': 'otherteacher', 'reason': '调整'}).status_code == 200


def test_dept_vacancy_goes_to_college_then_same_person_marked(client2, app2):
    _assessment(app2, 'a-teacher', 'teacher')
    with app2.state.sessions.begin() as db:
        db.get(User, 'depthead').active = False   # 教研室一缺负责人
    auth(client2, 'collegehead')
    d = client2.get('/api/assessments/a-teacher').json()
    assert d['handler'] == {'kind': 'college', 'acting_for': '代缺位教研室负责人办理'}
    assert _act(client2, 'a-teacher').status_code == 200          # 代教研室办理
    r = _act(client2, 'a-teacher'); assert r.status_code == 200, r.text  # 学院这一级，同一人连审
    assert r.json()['state'] == 'SCHOOL'
    with app2.state.sessions.begin() as db:
        events = db.scalars(select(WorkflowEvent).where(WorkflowEvent.object_id == 'a-teacher').order_by(WorkflowEvent.created_at)).all()
        assert events[0].after['handling'] == {'stage': 'DEPARTMENT', 'acting_for': '代缺位教研室负责人办理'}
        assert events[1].after['handling'] == {'stage': 'COLLEGE', 'same_person': '同一人连审'}
    # 学校不能越级办理教研室这一级：教研室有负责人时
    _assessment(app2, 'a-normal', 'otherteacher')
    auth(client2, 'schooladmin')
    assert _act(client2, 'a-normal').status_code == 403


def test_own_assessment_never_self_reviewed_across_identities(client2, app2):
    auth(client2, 'schooladmin')
    assert _add_extra(client2, 'teacher', [{'role_id': 'role-dept', 'org_id': 'dept1'}]).status_code == 200
    _assessment(app2, 'a-self', 'teacher')
    me = auth(client2, 'teacher')
    _switch(client2, me['identities'][1]['id'])
    assert _act(client2, 'a-self').status_code == 403


def test_unlock_keeps_password(client2, app2):
    with app2.state.sessions.begin() as db:
        u = db.get(User, 'teacher'); u.failures = 10; u.locked_until = datetime.utcnow() + timedelta(minutes=10)
    auth(client2, 'schooladmin')
    x = _user(client2, 'teacher'); assert x['locked'] is True
    r = post(client2, '/api/users/teacher/unlock', {'expected_version': x['version'], 'reason': '本人确认'})
    assert r.status_code == 200, r.text
    assert _user(client2, 'teacher')['locked'] is False
    auth(client2, 'teacher')


def test_temporary_identity_expires_after_its_last_day(app2,client2):
    from app.core import school_today
    from app.identities import KEY,expired,parse_valid_until
    from app.models import Setting
    import pytest
    from fastapi import HTTPException
    today=school_today()
    assert not expired({'valid_until':today.isoformat()}) and expired({'valid_until':(today-timedelta(days=1)).isoformat()}) and not expired({})
    for bad in ((today-timedelta(days=1)).isoformat(),(today+timedelta(days=400)).isoformat(),'2026/1/1','abc'):
        with pytest.raises(HTTPException):parse_valid_until(bad)
    auth(client2,'schooladmin')
    assert _add_extra(client2,'teacher',[{'role_id':'role-dept','org_id':'dept1','valid_until':(today-timedelta(days=1)).isoformat()}]).status_code==422
    r=_add_extra(client2,'teacher',[{'role_id':'role-dept','org_id':'dept1','valid_until':(today+timedelta(days=3)).isoformat()}])
    assert r.status_code==200,r.text
    assert len(auth(client2,'teacher')['identities'])==2
    # 到期（把到期日改成昨天，模拟时间过去）：不再能切换，也不算负责人；管理员列表仍看得到并标注已到期
    with app2.state.sessions.begin() as db:
        row=db.get(Setting,KEY);items=[dict(x) for x in row.value['items']]
        for x in items:x['valid_until']=(today-timedelta(days=1)).isoformat()
        row.value={'items':items};row.version+=1
    assert auth(client2,'teacher').get('identities',[])==[]
    auth(client2,'schooladmin')
    extra=_user(client2,'teacher')['extra_identities'][0]
    assert extra['expired'] is True and extra['valid_until']
    heads=next(o for o in client2.get('/api/organizations').json()['items'] if o['id']=='dept1')['heads']
    assert all(h['user_id']!='teacher' for h in heads)


def test_permission_list_export_has_one_row_per_identity(client2):
    from io import BytesIO
    from openpyxl import load_workbook
    from datetime import timedelta
    from app.core import school_today
    auth(client2,'schooladmin')
    until=(school_today()+timedelta(days=2)).isoformat()
    assert _add_extra(client2,'teacher',[{'role_id':'role-dept','org_id':'dept1','valid_until':until},{'role_id':'role-college','org_id':'college1'}]).status_code==200
    r=client2.get('/api/delivery/permission-list.xlsx');assert r.status_code==200,r.text
    sheet=load_workbook(BytesIO(r.content)).active
    rows=[[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    mine=[x for x in rows if x[0]=='teacher']
    assert [x[6] for x in mine]==['主身份','临时兼任','长期兼任'] and mine[1][10]==until and '、' in mine[0][11]
    auth(client2,'teacher')
    assert client2.get('/api/delivery/permission-list.xlsx').status_code==403
