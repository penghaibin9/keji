import io,zipfile,json,uuid
from openpyxl import load_workbook
from sqlalchemy import select,func
from app.models import TrainingRecord,User,ImportJob,ImportRow,Task,PointEntry,Period
from datetime import datetime
from app.worker import run_once
from conftest import auth,post

REASON={'reason':'Excel闭环测试','basis':'合成测试，不是学校正式数据'}


def fill_template(client,kind,values):
    r=client.get('/api/imports/templates/'+kind);assert r.status_code==200,r.text
    wb=load_workbook(io.BytesIO(r.content));ws=wb.active;ws.append(values);out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def upload_preview(client,kind,data,name='import.xlsx'):
    return client.post(f'/api/imports/{kind}/preview',files={'file':(name,data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},headers={'Idempotency-Key':str(uuid.uuid4())})


def test_records_excel_template_preview_confirm_and_repeat_guard(client2,app2):
    auth(client2,'depthead')
    data=fill_template(client2,'records',['EXCEL-001','J001','2026-09-10','PX','项目X','CX','课程X','CLX','班级X',2.5,12,321.5,'Excel导入'])
    r=upload_preview(client2,'records',data);assert r.status_code==200,r.text
    job=r.json();assert job['errors']==0 and job['state']=='PREVIEW'
    detail=client2.get(f"/api/imports/{job['id']}").json();assert detail['rows']['total']==1 and detail['rows']['items'][0]['errors']==[]
    r=post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON});assert r.status_code==200,r.text
    done=r.json();assert done['state']=='COMMITTED' and done['receipt']['written']==1
    with app2.state.sessions() as db:
        x=db.scalar(select(TrainingRecord).where(TrainingRecord.external_key=='EXCEL-001'));assert x and x.teacher_id=='teacher' and str(x.income)=='321.50'
    assert post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':done['version'],**REASON}).status_code==409


def test_excel_formula_and_bad_teacher_become_error_rows_not_business_data(client2,app2):
    auth(client2,'depthead')
    data=fill_template(client2,'records',['=1+1','NOPE','2026-09-10','PX','项目X','CX','课程X','CLX','班级X',2,1,1,'Excel导入'])
    r=upload_preview(client2,'records',data);assert r.status_code==200,r.text
    job=r.json();assert job['errors']==1
    e=client2.get(f"/api/imports/{job['id']}/errors");assert e.status_code==200
    wb=load_workbook(io.BytesIO(e.content),read_only=True);rows=list(wb.active.iter_rows(values_only=True));wb.close()
    assert '错误原因' in rows[0] and ('公式' in str(rows[1][-1]) or '教师' in str(rows[1][-1]))
    assert post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON}).status_code==409
    with app2.state.sessions() as db:assert db.scalar(select(func.count()).select_from(TrainingRecord).where(TrainingRecord.external_key=='=1+1'))==0


def test_teacher_account_excel_import_is_school_only_and_defaults_disabled(client2,app2):
    auth(client2,'depthead')
    assert client2.get('/api/imports/templates/teachers').status_code==403
    auth(client2,'schooladmin')
    data=fill_template(client2,'teachers',['excel_teacher','Excel教师','J900','dept1','讲师'])
    r=upload_preview(client2,'teachers',data);assert r.status_code==200,r.text
    job=r.json();assert job['errors']==0
    r=post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON});assert r.status_code==200,r.text
    with app2.state.sessions() as db:
        u=db.scalar(select(User).where(User.username=='excel_teacher'));assert u and not u.active and u.must_change_password and u.staff_no=='J900'


def test_import_schema_errors_show_business_chinese_without_program_details(client2):
    auth(client2,'depthead')
    samples=[('qualifications',['J001','职业资格','合成资质','IMPORT-BAD-DATE','合成机构','2026-09-01','2026-08-01'],'有效期不能早于取得日期'),
             ('participations',['IMPORT-BAD-HOURS','J001','2026-09-01','PX','合成参训','CX','合成课程','1.234','Excel导入'],'参训学时格式或取值不符合要求')]
    for kind,values,message in samples:
        response=upload_preview(client2,kind,fill_template(client2,kind,values));assert response.status_code==200,response.text
        job=response.json();assert job['errors']==1
        detail=client2.get('/api/imports/'+job['id']).json()
        assert detail['rows']['items'][0]['errors']==[message]
        assert not any(value in str(detail['rows']['items'][0]['errors']) for value in ['validation error','QualificationIn','ParticipationIn','https://','type=','input_value'])
        exported=client2.get('/api/imports/'+job['id']+'/errors');assert exported.status_code==200
        workbook=load_workbook(io.BytesIO(exported.content),data_only=False);rows=list(workbook.active.values)
        assert rows[1][-1]==message;workbook.close()


def test_legacy_import_errors_are_read_in_chinese_without_rewriting_the_saved_record(client2,app2):
    auth(client2,'depthead')
    values=['J001','职业资格','合成历史预览','IMPORT-LEGACY-DATE','合成机构','2026-09-01','2026-08-01']
    job=upload_preview(client2,'qualifications',fill_template(client2,'qualifications',values)).json()
    legacy='1 validation error for QualificationIn\nValue error, 有效期不能早于取得日期 [type=value_error]\nFor further information visit https://errors.pydantic.dev/example'
    with app2.state.sessions.begin() as db:
        row=db.scalar(select(ImportRow).where(ImportRow.job_id==job['id']));row.errors=[legacy];ident=row.id
    response=client2.get('/api/imports/'+job['id']);assert response.status_code==200,response.text
    message='该行数据校验未通过，请核对日期、数值和必填项后重新上传'
    assert response.json()['rows']['items'][0]['errors']==[message]
    assert 'validation error' not in response.text and 'https://' not in response.text
    exported=client2.get('/api/imports/'+job['id']+'/errors');assert exported.status_code==200
    workbook=load_workbook(io.BytesIO(exported.content));assert list(workbook.active.values)[1][-1]==message;workbook.close()
    with app2.state.sessions() as db:assert db.get(ImportRow,ident).errors==[legacy]


def test_record_export_worker_and_download_are_real_xlsx(client2,app2):
    auth(client2,'schooladmin')
    r=post(client2,'/api/exports',{'kind':'records','filters':{'project_code':'P1'},'confirmed':False});assert r.status_code==200,r.text
    task=r.json();result=run_once(app2.state.settings);assert result['status']=='DONE'
    d=client2.get(f"/api/tasks/{task['id']}/download");assert d.status_code==200,d.text
    wb=load_workbook(io.BytesIO(d.content),read_only=True,data_only=False);rows=list(wb.active.iter_rows(values_only=True));wb.close()
    assert rows[0][0]=='业务日期' and len(rows)==2 and rows[1][4]=='P1'


def test_points_export_uses_business_names_local_time_and_hidden_trace_columns(client2,app2):
    with app2.state.sessions.begin() as db:
        period_name=db.get(Period,'period2026').name
        for ident,teacher,key,value,snapshot in [('visible','teacher','social_teaching_claim','0.07',{}),('zero','teacher','social_fix_auxiliary_correction','0',{'correction_id':'correction'}),('outside','otherteacher','social_teaching_other','999',{})]:
            db.add(PointEntry(id=ident,teacher_id=teacher,period_id='period2026',item_key=key,kind='SOCIAL',amount=value,raw_value=value,formula='按归档材料核对',source_snapshot=snapshot,actor_id='schooladmin',reason='合成测试',basis='合成依据',created_at=datetime(2026,10,17,21,24,7)))
    auth(client2,'teacher')
    task=post(client2,'/api/exports',{'kind':'points','filters':{},'confirmed':False}).json()
    assert run_once(app2.state.settings)['status']=='DONE'
    response=client2.get('/api/tasks/'+task['id']+'/download');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content),data_only=True);ws=wb.active;rows=list(ws.values);headers=rows[0]
    assert len(rows)==3 and {r[3] for r in rows[1:]}=={'社会培训授课积分','社会培训辅助积分更正'}
    assert {r[2] for r in rows[1:]}=={period_name} and sum(r[5] for r in rows[1:])==0.07
    assert rows[1][10]=='2026-10-18T05:24:07' and rows[1][15]=='2026-10-17T21:24:07'
    for index,title in enumerate(headers,1):
        from openpyxl.utils import get_column_letter
        assert ws.column_dimensions[get_column_letter(index)].hidden==title.startswith('实施人员使用：')
    assert {r[14] for r in rows[1:]}=={'social_teaching_claim','social_fix_auxiliary_correction'}
    wb.close()


def test_handover_export_is_open_zip_and_excludes_credentials(client2,app2):
    auth(client2,'schooladmin')
    r=post(client2,'/api/exports',{'kind':'handover','filters':{},'confirmed':True});assert r.status_code==200,r.text
    task=r.json();assert run_once(app2.state.settings)['status']=='DONE'
    d=client2.get(f"/api/tasks/{task['id']}/download");assert d.status_code==200
    with zipfile.ZipFile(io.BytesIO(d.content)) as z:
        names=set(z.namelist());assert {'manifest.json','database_structure.json','data/training_records.csv','data/users.csv'}<=names
        text=z.read('data/users.csv').decode('utf-8-sig');assert 'password_hash' not in text and 'FixtureOnly' not in text
        manifest=json.loads(z.read('manifest.json'));assert manifest['format'].startswith('Yueke NonDegree open handover')
