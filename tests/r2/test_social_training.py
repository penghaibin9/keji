"""School policy boundaries use isolated data and normal authenticated requests."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import uuid
import pytest
from PIL import Image
from sqlalchemy import select
from app.models import User, Role, Setting, PointEntry, WorkflowEvent, Notice
from app.social_models import SocialClaim, SocialProject, SocialUnit, SocialResource
from app.reference_data import TrainingProjectMaster
from app import social_training as st
from conftest import auth, post, put
from support import HASH


def payload():
    return {'reason':'隔离测试项目申报','data':{'code':'SOC-P1','name':'隔离技能培训','cohort':'第一期','org_id':'college1','owner_id':'teacher',
        'category':'community','high_risk':False,'expected_income':'0','expected_participants':30,'expected_outcomes':'完成技能培训',
        'start_date':'2026-09-01','end_date':'2026-09-02','mode':'module','class_type':'technical','level':'province','level_basis':'隔离省级专项认定依据','scene':'dispatched',
        'auxiliary_standard':'隔离已备案细则','auxiliary_daily_cap':'2','venue':{'type':'campus','place':'隔离楼101教室'}}}


def row(key='unit1',count=30):
    return {'external_key':key,'course_code':'SOC-C1','class_code':'SOC-CL1','module':'技能一','event_date':'2026-09-01','mode':'module','evidence_note':'隔离完整参训佐证',
        'learners':[{'reference':f'L{i}','full_time_student':False,'graduate_exception':False,'signed':True,'hours_met':True,'full_attendance':True,'passed':True,'priority_group':True,'exclusion_reason':''} for i in range(count)],
        'teachers':[{'teacher_id':'teacher','staff_type':'internal','role':'lecturer'}]}


def proof(client,kind,ident,version):
    image=BytesIO();Image.new('RGB',(2,2),'white').save(image,format='PNG')
    response=client.post('/api/attachments/'+kind+'/'+ident,headers={'Idempotency-Key':str(uuid.uuid4())},data={'expected_version':str(version),'reason':'隔离正式依据附件'},files={'file':('evidence.png',image.getvalue(),'image/png')})
    assert response.status_code==200,response.text


def detail(client,kind,ident):
    r=client.get('/api/social/objects/'+kind+'/'+ident);assert r.status_code==200,r.text;return r.json()


def action(client,route,kind,ident,operation,**kwargs):
    item=detail(client,kind,ident)['item']
    return post(client,'/api/social/'+route+'/'+ident+'/actions',{'action':operation,'expected_version':item['version'],'reason':'隔离业务确认','basis':'隔离书面依据',**kwargs})


@pytest.fixture
def setup_social(app2,client2):
    with app2.state.sessions.begin() as db:
        for ident in ('reviewer','vice'):
            db.add(User(id=ident,username=ident,name='隔离审核人',org_id='school',role_id='role-school',password_hash=HASH,must_change_password=False,active=True))
        db.add(Setting(key='social_training_duties',value={'dean':['collegehead'],'coordinator':['reviewer'],'vice_president':['vice'],'secretary':['depthead'],'finance':['reviewer'],'academic':['reviewer'],'discipline':['reviewer'],'committee':['vice']}))
    auth(client2,'schooladmin')
    r=post(client2,'/api/social/projects',payload());assert r.status_code==200,r.text
    p=r.json()['item'];proof(client2,'social_project',p['id'],p['version'])
    r=action(client2,'projects','social_project',p['id'],'submit');assert r.status_code==200,r.text
    auth(client2,'reviewer');r=action(client2,'projects','social_project',p['id'],'file');assert r.status_code==200,r.text
    r=post(client2,'/api/social/courses',{'code':'SOC-C1','name':'技能课程','project_code':'SOC-P1','reason':'隔离课程登记'});assert r.status_code==200,r.text
    return p['id']


def import_rows(client,pid,rows):
    data={'rows':rows,'expected_version':detail(client,'social_project',pid)['item']['version'],'reason':'隔离参训导入'}
    preview=post(client,'/api/social/projects/'+pid+'/units/preview',data)
    if preview.status_code!=200:return preview
    return post(client,'/api/social/projects/'+pid+'/units/confirm',{**data,'preview_hash':preview.json()['preview_hash']})


def close_project(client,pid):
    r=action(client,'projects','social_project',pid,'close_submit',closeout={'closed_on':'2026-09-02','outcome_report':'完整结项报告','satisfaction':'96','actual_income':'10000','funding_source':'other_training','retained':'2000','personnel':'6000','finance_reference':'财务确认凭证','location':'onsite','implicit_cost':'0','settlement_on':'2026-09-03','materials_checked':['plan','filing_form','roster','attendance','hours','process','assessment','outcomes']})
    assert r.status_code==200,r.text
    r=action(client,'projects','social_project',pid,'close');assert r.status_code==200,r.text


def policy(client,complete=False,**overrides):
    auth(client,'schooladmin')
    cfg={'interpretation_reference':'隔离书面解释','calendar_start':'2026-01-01','calendar_end':'2027-12-31','calendar_reference':'隔离学校工作日历'}
    if complete:cfg.update(cap_mode='ratio_cap',person_day_unit='daily',precision=2,rounding='half_up',effective_date='2026-08-01',date_basis='closeout',calendar_start='2026-01-01',calendar_end='2027-12-31',calendar_reference='隔离学校工作日历',holidays=[],workdays=[],auxiliary_overlap='exclusive',title_scope='teaching',title_precision=2,title_rounding='half_up')
    cfg.update(overrides)
    r=post(client,'/api/social/policies',{'period_id':'period2026','name':'隔离确认口径','data':cfg,'reason':'隔离规则登记'});assert r.status_code==200,r.text
    item=r.json()['item'];proof(client,'social_policy',item['id'],item['version'])
    auth(client,'reviewer');r=action(client,'policies','social_policy',item['id'],'activate');assert r.status_code==200,r.text
    return item['id']


def test_workday_calendar_missing_and_adjusted_weekend():
    with pytest.raises(Exception,match='学校工作日历'):st.add_workdays(date(2026,10,1),1,{})
    cfg={'calendar_start':'2026-10-01','calendar_end':'2026-10-31','calendar_reference':'确认校历','holidays':['2026-10-02'],'workdays':['2026-10-03']}
    assert st.add_workdays(date(2026,10,1),1,cfg)==date(2026,10,3)
    assert st.add_workdays(date(2026,10,1),2,cfg)==date(2026,10,5)


@pytest.fixture
def project_candidates(app2):
    with app2.state.sessions.begin() as db:
        for index in range(207):
            ident=f'option-{index:03d}';code=f'OPT-{index:03d}'
            db.add(TrainingProjectMaster(id=ident,code=code,name='隔离候选项目'+str(index)))
        db.flush()
        for index in range(207):
            ident=f'option-{index:03d}';org='college2' if index==206 else 'college1'
            owner='depthead' if index in (0,205) else 'otherteacher' if index==206 else 'teacher'
            data={**payload()['data'],'code':f'OPT-{index:03d}','name':'隔离候选项目'+str(index),'cohort':'检索期次'+str(index),'org_id':org,'owner_id':owner}
            db.add(SocialProject(id=ident,master_id=ident,org_id=org,owner_id=owner,cohort=data['cohort'],data=data))
        db.flush()
        for index,teacher in ((0,'teacher'),(205,'teacher-suffix'),(206,'teacher')):
            for unit in range(2):
                db.add(SocialUnit(project_id=f'option-{index:03d}',external_key=str(unit),data={'teachers':[{'teacher_id':teacher}]}))


def test_project_options_search_and_pagination_beyond_two_hundred(project_candidates,client2,monkeypatch):
    assert client2.get('/api/social/project-options').status_code==401
    auth(client2,'teacher')
    def forbid_unit_scan(*args):raise AssertionError('项目候选不得逐项目加载参训明细')
    monkeypatch.setattr(st,'units',forbid_unit_scan)
    first=client2.get('/api/social/project-options').json()
    assert first['total']==205 and first['page']==1 and first['size']==25 and len(first['items'])==25
    pages=[client2.get('/api/social/project-options',params={'page':page,'size':100}).json() for page in (1,2,3)]
    assert [len(p['items']) for p in pages]==[100,100,5]
    assert len({item['id'] for p in pages for item in p['items']})==205
    assert client2.get('/api/social/project-options',params={'page':4,'size':100}).json()['items']==[]
    for query in ('OPT-000','隔离候选项目0','检索期次0'):
        found=client2.get('/api/social/project-options',params={'q':query}).json()
        assert found['total']==1 and found['items'][0]['id']=='option-000'
        assert found['items'][0]['code']=='OPT-000' and 'data' not in found['items'][0]
    for query in ('OPT-205','OPT-206','%','_'):
        found=client2.get('/api/social/project-options',params={'q':query}).json()
        assert found['total']==0 and found['items']==[]
    for params in ({'size':101},{'page':0},{'q':'长'*101}):
        assert client2.get('/api/social/project-options',params=params).status_code==422
    workspace=client2.get('/api/social/workspace',params={'section':'projects','page':2}).json()
    assert workspace['pagination']['total']==205 and len(workspace['projects'])==5
    found=client2.get('/api/social/workspace',params={'section':'projects','q':'检索期次0'}).json()
    assert found['pagination']['total']==1 and found['projects'][0]['id']=='option-000'


def test_project_options_enforce_current_role_scope_and_permission(project_candidates,client2,app2):
    for user,expected in (('collegehead',206),('otherteacher',1),('otherdept',1),('schooladmin',207)):  # 教研室只读本室教师负责的项目
        auth(client2,user)
        result=client2.get('/api/social/project-options')
        assert result.status_code==200,result.text
        assert result.json()['total']==expected
        if user=='otherteacher':assert result.json()['items'][0]['id']=='option-206'
    with app2.state.sessions.begin() as db:
        role=db.get(Role,'role-school');role.permissions=[p for p in role.permissions if p!='records.read']
    assert client2.get('/api/social/project-options').status_code==403


def test_teacher_resource_org_names_preserve_college_scope(client2,app2):
    with app2.state.sessions.begin() as db:
        for org in ('college1','college2'):
            db.add(SocialResource(org_id=org,kind='base',code='BASE-'+org,name='隔离培训基地',data={'location':'校内','onsite':True,'recognition':'隔离认定依据'}))
    auth(client2,'teacher')
    orgs=client2.get('/api/organizations').json()['items']
    assert [org['id'] for org in orgs]==['dept1']
    response=client2.get('/api/social/workspace',params={'section':'bases'})
    assert response.status_code==200,response.text
    resources=response.json()['resources']
    assert len(resources)==1
    assert resources[0]['org_id']=='college1' and resources[0]['org_name']=='测试学院一'
    assert response.json()['pagination']['total']==1
    assert client2.get('/api/social/workspace',params={'section':'bases','q':'BASE-college2'}).json()['resources']==[]
    assert [org['id'] for org in client2.get('/api/organizations').json()['items']]==['dept1']


def test_import_is_preview_then_atomic_and_no_double_teacher_count(setup_social,client2):
    pid=setup_social;r=row();r['teachers'].append({'teacher_id':'depthead','staff_type':'external','role':'instructor'})
    result=import_rows(client2,pid,[r]);assert result.status_code==200,result.text
    assert result.json()['effective']==30
    assert import_rows(client2,pid,[row()]).status_code==409
    assert detail(client2,'social_project',pid)['facts']['effective']==30


def test_scope_and_same_mode_are_enforced(setup_social,client2):
    pid=setup_social;r=row();r['mode']='person_day'
    assert import_rows(client2,pid,[r]).status_code==422
    auth(client2,'otherdept');assert client2.get('/api/social/objects/social_project/'+pid).status_code==403


def test_unconfirmed_policy_never_books_formal_points(setup_social,client2,app2):
    pid=setup_social;assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid);pol=policy(client2)
    auth(client2,'teacher');r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'本人隔离申报'});assert r.status_code==200,r.text
    claim=r.json()['item'];assert any('封顶' in b for b in r.json()['trial']['blockers'])
    assert r.json()['trial']['total'] is None
    proof(client2,'social_claim',claim['id'],claim['version'])
    assert action(client2,'claims','social_claim',claim['id'],'submit').status_code==409
    with app2.state.sessions() as db:assert db.scalar(select(func_count(PointEntry.id)).where(PointEntry.kind=='SOCIAL'))==0


def func_count(column):
    from sqlalchemy import func
    return func.count(column)


@pytest.mark.parametrize('title_precision,expected_title',[(2,'0.02'),(0,'0'),(None,None)])
def test_full_recognition_roles_and_no_duplicate_booking(setup_social,client2,app2,monkeypatch,title_precision,expected_title):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3))
    pid=setup_social;assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid);pol=policy(client2,True,title_precision=title_precision)
    auth(client2,'teacher');r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'本人隔离申报'});assert r.status_code==200,r.text
    claim=r.json()['item'];assert r.json()['trial']['teaching']=='1.56'
    proof(client2,'social_claim',claim['id'],claim['version']);r=action(client2,'claims','social_claim',claim['id'],'submit');assert r.status_code==200,r.text
    auth(client2,'schooladmin');assert action(client2,'claims','social_claim',claim['id'],'approve').status_code==403
    for user in ('collegehead','reviewer','vice','depthead'):
        auth(client2,user);r=action(client2,'claims','social_claim',claim['id'],'approve');assert r.status_code==200,r.text
    assert r.json()['item']['state']=='ARCHIVED'
    assert action(client2,'claims','social_claim',claim['id'],'approve').status_code==409
    with app2.state.sessions() as db:
        entries=db.scalars(select(PointEntry).where(PointEntry.kind=='SOCIAL')).all()
        assert sum((e.amount for e in entries),Decimal(0))==Decimal('1.56')
    auth(client2,'teacher')
    year=client2.get('/api/social/ledger-summary').json()['items'][0]['years']['2026']
    assert year['title_score']==expected_title
    assert year['title_pending'] is (title_precision is None)


@pytest.mark.parametrize('mode,expected_status',[(None,409),('original_submission',200),('resubmission_date',409)])
def test_returned_claim_deadline_requires_confirmed_basis(setup_social,client2,monkeypatch,mode,expected_status):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid)
    pol=policy(client2,True,resubmission_mode=mode)
    auth(client2,'teacher')
    r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'隔离退回补正期限验证'})
    assert r.status_code==200,r.text
    claim=r.json()['item'];proof(client2,'social_claim',claim['id'],claim['version'])
    assert action(client2,'claims','social_claim',claim['id'],'submit').status_code==200
    auth(client2,'collegehead');assert action(client2,'claims','social_claim',claim['id'],'return').status_code==200
    monkeypatch.setattr(st,'today',lambda:date(2026,10,8));auth(client2,'teacher')
    result=action(client2,'claims','social_claim',claim['id'],'submit')
    assert result.status_code==expected_status,result.text
    if mode=='original_submission':assert result.json()['item']['data']['submitted_on']=='2026-09-03'
    else:assert detail(client2,'social_claim',claim['id'])['item']['state']=='RETURNED'


def test_exclusions_and_graduate_exception():
    r=row(count=4);r['learners'][0]['full_time_student']=True;r['learners'][1]['passed']=False
    r['learners'][2]['full_time_student']=True;r['learners'][2]['graduate_exception']=True
    assert st.unit_counts(r)=={'effective':2,'excluded':2,'priority':2}


def test_annual_review_publicity_and_missing_targets(setup_social,client2,monkeypatch):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid)
    pol=policy(client2,True,annual_model='four_indicators',bonus_pool='all_bonus_shared_ten',total_cap='110',grade_mode='four',step_mode='full_steps',satisfaction_inclusive=False,achievement_dedup='highest_per_reference',organization_basis='project_units')
    auth(client2,'collegehead')
    data={'org_id':'college1','period_id':'period2026','policy_id':pol,'reason':'年度自评隔离验证','data':{'self_report':'完整自评报告','income_target':None,'participants_target':30,'priority_target':30,'target_reference':'年度任务确认','achievements':[{'reference':'A1','name':'省级成果','level':'province','confirmed_on':'2026-09-02','recognition_reference':'隔离省级成果确认决定'},{'reference':'A1','name':'同一成果国家级认定','level':'national','confirmed_on':'2026-09-02','recognition_reference':'隔离国家级成果确认决定'}]}}
    r=post(client2,'/api/social/annuals',data);assert r.status_code==200,r.text
    item=r.json()['item'];ident=item['id'];assert any('收入年度目标' in b for b in r.json()['trial']['blockers'])
    proof(client2,'social_annual',ident,item['version']);assert action(client2,'annuals','social_annual',ident,'submit').status_code==409
    data['expected_version']=detail(client2,'social_annual',ident)['item']['version'];data['data']['income_target']='10000'
    r=put(client2,'/api/social/annuals/'+ident,data);assert r.status_code==200,r.text
    assert r.json()['trial']['total']=='90.00';assert r.json()['trial']['achievement_base']=='10.00'
    assert action(client2,'annuals','social_annual',ident,'submit').status_code==200
    auth(client2,'reviewer')
    for step in ('finance','academic','discipline','coordinator'):
        r=action(client2,'annuals','social_annual',ident,step);assert r.status_code==200,r.text
    auth(client2,'vice');r=action(client2,'annuals','social_annual',ident,'publish');assert r.status_code==200,r.text
    auth(client2,'schooladmin');assert action(client2,'annuals','social_annual',ident,'archive').status_code==403
    # A school secretary can archive only after the complete publicity window.
    from app.models import Setting
    with client2.app.state.sessions.begin() as db:
        s=db.get(Setting,'social_training_duties');s.value={**s.value,'secretary':['vice']}
    auth(client2,'vice');assert action(client2,'annuals','social_annual',ident,'archive').status_code==409
    monkeypatch.setattr(st,'today',lambda:date(2026,9,12))
    r=action(client2,'annuals','social_annual',ident,'archive');assert r.status_code==200,r.text
    assert r.json()['item']['state']=='ARCHIVED'


@pytest.fixture
def recognized(setup_social,client2,monkeypatch,request):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid);pol=policy(client2,True,**getattr(request,'param',{}))
    auth(client2,'teacher');r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'隔离积分申请'})
    assert r.status_code==200,r.text
    claim=r.json()['item'];proof(client2,'social_claim',claim['id'],claim['version'])
    assert action(client2,'claims','social_claim',claim['id'],'submit').status_code==200
    for user in ('collegehead','reviewer','vice','depthead'):
        auth(client2,user);r=action(client2,'claims','social_claim',claim['id'],'approve');assert r.status_code==200,r.text
    return claim['id'],pol


def test_correction_is_reviewed_append_only_and_summary_tracks_it(recognized,client2,app2):
    ident,_=recognized;original=detail(client2,'social_claim',ident)['item']['snapshot']
    auth(client2,'teacher');r=post(client2,'/api/social/corrections',{'claim_id':ident,'period_id':'period2026','teaching':'-0.56','auxiliary':'0','decision_reference':'学校正式更正决定','reason':'原始签到核实后更正'})
    assert r.status_code==200,r.text
    correction=r.json()['item'];proof(client2,'social_resource',correction['id'],correction['version'])
    assert action(client2,'corrections','social_resource',correction['id'],'submit').status_code==200
    assert action(client2,'corrections','social_resource',correction['id'],'approve').status_code==403
    for user in ('collegehead','reviewer','vice','depthead'):
        auth(client2,user);r=action(client2,'corrections','social_resource',correction['id'],'approve');assert r.status_code==200,r.text
    assert action(client2,'corrections','social_resource',correction['id'],'approve').status_code==409
    assert detail(client2,'social_claim',ident)['item']['snapshot']==original
    auth(client2,'teacher');summary=client2.get('/api/social/ledger-summary').json()['items'][0]
    assert Decimal(summary['teaching'])==Decimal('1.00')
    assert Decimal(summary['years']['2026']['title_score'])==Decimal('0.01')
    auth(client2,'schooladmin')
    r=post(client2,'/api/points/adjust',{'teacher_id':'teacher','period_id':'period2026','item_key':'manual_review','amount':'1','reason':'试图跳过社会培训更正','basis':'隔离'})
    assert r.status_code==409,r.text
    with app2.state.sessions() as db:
        assert sum(db.scalars(select(PointEntry.amount).where(PointEntry.kind=='SOCIAL')))==Decimal('1.00')


def test_bimonth_and_annual_reports_use_recognized_snapshot(recognized,client2,app2,monkeypatch):
    _,pol=recognized
    with app2.state.sessions.begin() as db:
        setting=db.get(Setting,'social_training_duties');setting.value={**setting.value,'secretary':['vice']}
    auth(client2,'vice')
    r=post(client2,'/api/social/reports',{'org_id':'college1','policy_id':pol,'year':2026,'start_month':9,'end_month':10,'reason':'双月积分梳理报送'})
    assert r.status_code==200,r.text
    report=r.json()['item'];assert report['data']['rows'][0]['total']=='1.56'
    assert action(client2,'reports','social_resource',report['id'],'submit').status_code==200
    assert detail(client2,'social_resource',report['id'])['events'][-1]['action']=='积分汇总：待归口复核'
    with app2.state.sessions.begin() as db:
        old_event=db.scalar(select(WorkflowEvent).where(WorkflowEvent.object_id==report['id'],WorkflowEvent.action=='积分汇总：待归口复核'))
        old_id=old_event.id;old_event.action='积分汇总：待院长初审'
    assert detail(client2,'social_resource',report['id'])['events'][-1]['action']=='积分汇总：待归口复核'
    with app2.state.sessions() as db:assert db.get(WorkflowEvent,old_id).action=='积分汇总：待院长初审'
    auth(client2,'reviewer');assert action(client2,'reports','social_resource',report['id'],'approve').status_code==409
    monkeypatch.setattr(st,'today',lambda:date(2026,11,1));assert action(client2,'reports','social_resource',report['id'],'approve').status_code==200
    auth(client2,'vice');r=post(client2,'/api/social/reports',{'org_id':'school','policy_id':pol,'year':2026,'start_month':1,'end_month':12,'reason':'年度全校积分公示'})
    assert r.status_code==200,r.text
    ident=r.json()['item']['id'];assert action(client2,'reports','social_resource',ident,'submit').status_code==200
    monkeypatch.setattr(st,'today',lambda:date(2027,1,1));auth(client2,'reviewer')
    r=action(client2,'reports','social_resource',ident,'approve');assert r.status_code==200,r.text
    auth(client2,'teacher');assert action(client2,'reports','social_resource',ident,'dispute').status_code==200
    auth(client2,'vice');assert action(client2,'reports','social_resource',ident,'archive').status_code==409
    monkeypatch.setattr(st,'today',lambda:date(2027,1,10));assert action(client2,'reports','social_resource',ident,'archive').status_code==409
    auth(client2,'reviewer');assert action(client2,'reports','social_resource',ident,'resolve').status_code==200
    auth(client2,'vice');assert action(client2,'reports','social_resource',ident,'archive').status_code==200


def test_annual_report_cannot_archive_points_changed_after_publication(recognized,client2,app2,monkeypatch):
    claim_id,pol=recognized
    with app2.state.sessions.begin() as db:
        setting=db.get(Setting,'social_training_duties');setting.value={**setting.value,'secretary':['vice','depthead']}
    monkeypatch.setattr(st,'today',lambda:date(2027,1,1));auth(client2,'vice')
    r=post(client2,'/api/social/reports',{'org_id':'school','policy_id':pol,'year':2026,'start_month':1,'end_month':12,'reason':'年度全校积分公示'})
    assert r.status_code==200,r.text
    report=r.json()['item'];original=report['data']['rows']
    assert action(client2,'reports','social_resource',report['id'],'submit').status_code==200
    auth(client2,'reviewer');published=action(client2,'reports','social_resource',report['id'],'approve');assert published.status_code==200
    old_notice=published.json()['item']['data']['notice_id']
    with app2.state.sessions.begin() as db:
        legacy=db.get(SocialResource,report['id']);legacy.data={k:v for k,v in legacy.data.items() if k!='notice_id'}
    auth(client2,'teacher')
    r=post(client2,'/api/social/corrections',{'claim_id':claim_id,'period_id':'period2026','teaching':'-0.56','auxiliary':'0','decision_reference':'公示期间正式核实更正决定','reason':'已公示数据核实后更正'})
    assert r.status_code==200,r.text
    correction=r.json()['item'];proof(client2,'social_resource',correction['id'],correction['version'])
    assert action(client2,'corrections','social_resource',correction['id'],'submit').status_code==200
    for user in ('collegehead','reviewer','vice','depthead'):
        auth(client2,user);r=action(client2,'corrections','social_resource',correction['id'],'approve');assert r.status_code==200,r.text
    monkeypatch.setattr(st,'today',lambda:date(2027,1,10));auth(client2,'vice')
    r=action(client2,'reports','social_resource',report['id'],'archive')
    assert r.status_code==409,r.text
    assert '重新生成并公示' in r.json()['detail']
    retained=detail(client2,'social_resource',report['id'])['item']
    assert retained['state']=='PUBLIC' and retained['data']['rows']==original
    r=post(client2,'/api/social/reports',{'org_id':'school','policy_id':pol,'year':2026,'start_month':1,'end_month':12,'reason':'根据已批准更正重新生成年度公示'})
    assert r.status_code==200,r.text
    replacement=r.json()['item'];replacement_id=replacement['id']
    assert replacement_id!=report['id'] and replacement['data']['rows'][0]['total']=='1.00'
    assert action(client2,'reports','social_resource',replacement_id,'submit').status_code==200
    auth(client2,'teacher');assert action(client2,'reports','social_resource',report['id'],'dispute').status_code==200
    auth(client2,'reviewer');blocked=action(client2,'reports','social_resource',replacement_id,'approve')
    assert blocked.status_code==409 and '旧公示仍有异议未处理' in blocked.json()['detail']
    assert detail(client2,'social_resource',report['id'])['item']['state']=='PUBLIC'
    assert detail(client2,'social_resource',replacement_id)['item']['state']=='SUBMITTED'
    with app2.state.sessions() as db:assert db.get(Notice,old_notice).published
    assert action(client2,'reports','social_resource',report['id'],'resolve').status_code==200
    with app2.state.sessions.begin() as db:
        notice=db.get(Notice,old_notice)
        db.add(Notice(id='ambiguous-report-notice',title=notice.title,body=notice.body,category=notice.category,org_id=notice.org_id,author_id=notice.author_id,published=True,created_at=notice.created_at))
    ambiguous=action(client2,'reports','social_resource',replacement_id,'approve')
    assert ambiguous.status_code==409 and '通知关联不明确' in ambiguous.json()['detail']
    assert detail(client2,'social_resource',report['id'])['item']['state']=='PUBLIC'
    assert detail(client2,'social_resource',replacement_id)['item']['state']=='SUBMITTED'
    with app2.state.sessions.begin() as db:
        assert db.get(Notice,old_notice).published
        db.delete(db.get(Notice,'ambiguous-report-notice'))
    monkeypatch.setattr(st,'today',lambda:date(2028,1,10));auth(client2,'reviewer')
    assert action(client2,'reports','social_resource',replacement_id,'approve').status_code==409
    assert detail(client2,'social_resource',report['id'])['item']['state']=='PUBLIC'
    assert detail(client2,'social_resource',replacement_id)['item']['state']=='SUBMITTED'
    with app2.state.sessions() as db:assert db.get(Notice,old_notice).published
    monkeypatch.setattr(st,'today',lambda:date(2027,1,10))
    auth(client2,'reviewer');r=action(client2,'reports','social_resource',replacement_id,'approve')
    assert r.status_code==200,r.text
    assert r.json()['item']['data']['public_start']=='2027-01-10'
    assert r.json()['item']['data']['public_end']=='2027-01-15'
    retired=detail(client2,'social_resource',report['id'])['item']
    assert retired['state']=='VOID' and retired['data']['superseded_by']==replacement_id
    assert retired['data']['rows']==original and retired['data']['sha256']==report['data']['sha256']
    assert report['id'] in r.json()['item']['data']['replaces']
    with app2.state.sessions() as db:
        assert not db.get(Notice,retired['data']['notice_id']).published
        assert db.get(Notice,r.json()['item']['data']['notice_id']).published
    new_notice=r.json()['item']['data']['notice_id'];auth(client2,'teacher')
    assert client2.get('/api/notices/'+old_notice).status_code==404
    assert client2.get('/api/notices/'+new_notice).status_code==200
    assert client2.get('/api/social/objects/social_resource/'+report['id']).status_code==403
    assert report['id'] not in {x['id'] for x in client2.get('/api/social/workspace?section=recognition').json()['resources']}
    auth(client2,'reviewer')
    assert report['id'] not in {x['id'] for x in client2.get('/api/todos').json()['items']}
    assert action(client2,'reports','social_resource',report['id'],'dispute').status_code==409
    exported=client2.get('/api/social/objects/social_resource/'+report['id']+'/export')
    assert exported.status_code==200,exported.text
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(exported.content),data_only=True)
    assert dict(list(book['业务概要'].values)[1:])['办理状态']=='已被新公示替代'
    assert list(book['已认定积分汇总'].values)[1][-1]=='1.56'
    auth(client2,'vice')
    for day in (date(2027,1,10),date(2027,1,14),date(2027,1,15)):
        monkeypatch.setattr(st,'today',lambda:day)
        r=action(client2,'reports','social_resource',replacement_id,'archive')
        assert r.status_code==409 and '5个完整工作日' in r.json()['detail'],r.text
    monkeypatch.setattr(st,'today',lambda:date(2027,1,16))
    r=action(client2,'reports','social_resource',replacement_id,'archive')
    assert r.status_code==200,r.text
    assert r.json()['item']['state']=='ARCHIVED'
    assert detail(client2,'social_resource',report['id'])['item']['data']['rows']==original
    assert action(client2,'reports','social_resource',report['id'],'archive').status_code==409


def test_course_import_atomicity_project_todos_and_export(setup_social,client2):
    pid=setup_social;course={'code':'SOC-C2','name':'第二课程','project_code':'SOC-P1','reason':'隔离批量登记'}
    data={'rows':[course],'reason':'隔离批量登记'}
    preview=post(client2,'/api/social/courses/import/preview',data);assert preview.status_code==200,preview.text
    assert post(client2,'/api/social/courses/import/confirm',{**data,'preview_hash':'bad'}).status_code==409
    r=post(client2,'/api/social/courses/import/confirm',{**data,'preview_hash':preview.json()['preview_hash']});assert r.status_code==200,r.text
    assert post(client2,'/api/social/courses/import/preview',data).status_code==409
    assert import_rows(client2,pid,[row()]).status_code==200
    assert import_rows(client2,pid,[row('changed-key')]).status_code==409
    rows=client2.get('/api/todos?type=social_project').json()['items'];assert any(x['id']==pid for x in rows)
    from openpyxl import load_workbook
    exported=client2.get('/api/social/objects/social_project/'+pid+'/export');assert exported.status_code==200
    book=load_workbook(BytesIO(exported.content));assert book['有效人次依据'].max_row==31


@pytest.mark.parametrize('count,cap,expected',[(30,'ratio_cap','1.56'),(31,'ratio_cap','1.61'),(31,'over_standard_two','3.12'),(60,'ratio_cap','3.12'),(61,'ratio_cap','3.12')])
def test_cap_boundaries_and_weights(setup_social,client2,monkeypatch,count,cap,expected):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row(count=count)]).status_code==200;close_project(client2,pid);pol=policy(client2,True,cap_mode=cap)
    auth(client2,'teacher');r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'边界候选试算'})
    assert r.status_code==200,r.text
    assert r.json()['trial']['teaching']==expected


def test_auxiliary_daily_cap_and_legacy_overlap_block_formal_results(setup_social,client2,app2,monkeypatch):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200;close_project(client2,pid);pol=policy(client2,True)
    with app2.state.sessions.begin() as db:
        db.add(PointEntry(teacher_id='teacher',period_id='period2026',item_key='legacy',kind='BASE',amount=1,raw_value='1',formula='旧规则隔离数据',source_snapshot={},actor_id='schooladmin',reason='原积分'))
    auth(client2,'teacher');r=post(client2,'/api/social/claims',{'project_id':pid,'period_id':'period2026','policy_id':pol,'recognition_form_reference':'学校制式认定表与佐证附件','reason':'隔离辅助工作','auxiliary':[{'work_date':'2026-09-01','amount':'3','description':'培训资料整理'}]})
    assert r.status_code==200,r.text
    blockers=r.json()['trial']['blockers'];assert any('个人每日' in b for b in blockers);assert any('旧规则' in b for b in blockers)


def test_formal_findings_require_duty_and_cannot_be_overwritten_as_complaints(recognized,client2,app2):
    ident,pol=recognized;auth(client2,'reviewer');claim=detail(client2,'social_claim',ident)['item'];pid=claim['project_id']
    data={'expected_version':detail(client2,'social_project',pid)['item']['version'],'kind':'violation','case_reference':'DECISION-1','received_on':'2026-09-03','policy_id':pol,'invalidate_cohort':True,'annual_veto':True,'reason':'录入学校正式处理决定','basis':'隔离学校正式书面决定'}
    auth(client2,'schooladmin');assert post(client2,'/api/social/projects/'+pid+'/notes',data).status_code==403
    auth(client2,'reviewer');r=post(client2,'/api/social/projects/'+pid+'/notes',data);assert r.status_code==200,r.text
    assert detail(client2,'social_project',pid)['facts']['effective']==0
    data.update(expected_version=r.json()['version'],kind='complaint',invalidate_cohort=False,annual_veto=False)
    assert post(client2,'/api/social/projects/'+pid+'/notes',data).status_code==409
    with app2.state.sessions() as db:
        from fastapi import HTTPException
        from app.core import Actor
        from app.models import Role,Org
        user=db.get(User,'schooladmin');actor=Actor(db,user,db.get(Role,user.role_id),db.get(Org,user.org_id))
        with pytest.raises(HTTPException,match='待更正'):st.report_rows(db,actor,'college1',2026,9,10)


def test_missing_closeout_materials_and_chinese_project_search(setup_social,client2,monkeypatch):
    monkeypatch.setattr(st,'today',lambda:date(2026,9,3));pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200
    r=action(client2,'projects','social_project',pid,'close_submit',closeout={'closed_on':'2026-09-02','outcome_report':'材料待补','satisfaction':84,'rectification':'补齐材料并整改','actual_income':0,'finance_reference':'零收入核验','location':'onsite'})
    assert r.status_code==200,r.text
    assert action(client2,'projects','social_project',pid,'close').status_code==409
    result=client2.get('/api/social/workspace?section=projects&q=隔离技能').json()
    assert result['pagination']['total']==1;assert result['projects'][0]['id']==pid


def test_project_venue_required_offsite_needs_record_and_prep_checklist(client2,app2):
    auth(client2,'schooladmin')
    base=payload()
    no_venue={**base,'data':{k:v for k,v in base['data'].items() if k!='venue'}}
    assert post(client2,'/api/social/projects',no_venue).status_code==422
    offsite={**base,'data':{**base['data'],'venue':{'type':'offsite','place':'市技师学院实训楼'}}}
    assert post(client2,'/api/social/projects',offsite).status_code==422  # 校外缺一事一议
    offsite['data']['venue'].update(offsite_reason='校内无该设备',offsite_ref='校长办公会 2026-08-20')
    r=post(client2,'/api/social/projects',offsite);assert r.status_code==200,r.text
    p=r.json()['item'];assert p['data']['venue']['offsite_ref']=='校长办公会 2026-08-20'
    # 清单只提醒：勾选后保留，改草稿不丢
    r=put(client2,'/api/social/projects/'+p['id']+'/prep',{'reason':'准备','expected_version':p['version'],'items':{'plan':True,'notice':True}});assert r.status_code==200,r.text
    item=r.json()['item'];assert item['data']['prep']=={'plan':True,'notice':True}
    edit={**offsite,'expected_version':item['version']}
    r=put(client2,'/api/social/projects/'+p['id'],edit);assert r.status_code==200,r.text
    assert r.json()['item']['data']['prep']=={'plan':True,'notice':True}
    assert put(client2,'/api/social/projects/'+p['id']+'/prep',{'reason':'x','items':{'bogus':True}}).status_code==422
