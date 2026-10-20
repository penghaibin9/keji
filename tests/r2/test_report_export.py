"""导出上报表：预检缺项、只算已备案及以后的项目、按区间重叠取项目，两张表的数从系统里算。"""
from io import BytesIO
from copy import deepcopy
from datetime import date,datetime
from decimal import Decimal
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from app.reference_data import TrainingCourseMaster
from app.social_models import SocialProject, SocialResource
from conftest import auth, put
from test_social_training import setup_social  # noqa: F401
from test_project_apply import QUALITY
from app import social_report

RANGE = {'start': '2023-07-01', 'end': '2026-12-31'}


def fill(app2, pid, **data):
    with app2.state.sessions.begin() as db:
        p = db.get(SocialProject, pid)
        p.data = {**p.data, 'total_hours': '16', 'audiences': ['企业职工'], 'audience_other': '退役军人', 'service_types': ['national'], 'teaching_methods': ['classroom'], 'quality': QUALITY, **data}
        return p.data['code'], p.org_id


def test_precheck_lists_missing_items_then_exports_two_sheets(setup_social, client2, app2):
    c = client2
    pid = setup_social
    auth(c, 'teacher')
    assert c.get('/api/social/report-export/check', params=RANGE).status_code == 403
    auth(c, 'reviewer')
    r = c.get('/api/social/report-export/check', params=RANGE)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d['count'] == 1 and d['blocked'] == 1 and not d['ready'] and '请填写总学时' in d['items'][0]['issues']
    assert c.get('/api/social/report-export', params=RANGE).status_code == 409
    assert c.get('/api/social/report-export/check', params={'start': '2026-12-31', 'end': '2026-01-01'}).status_code == 422
    code, org = fill(app2, pid)
    # 本项目课程：一门讲授 8 学时、课件已审查，一门实训 24 学时、课件未审查，带 30 分钟视频。
    with app2.state.sessions.begin() as db:
        db.add(TrainingCourseMaster(code='SOC-C8', name='实训课', project_code=code, source='测试'))
        for course, hours, mode, checked in [('SOC-C1', '8', 'lecture', True), ('SOC-C8', '24', 'practice', False)]:
            db.add(SocialResource(kind='course_apply', org_id=org, code='CI-' + course, name=course, state='ACTIVE',
                                  data={'course_code': course, 'hours': hours, 'teaching_mode': mode, 'courseware': {'checked': checked}}))
        db.add(SocialResource(kind='course_resource', org_id=org, code='V-1', name='视频', state='DRAFT',
                              data={'course_code': 'SOC-C8', 'resource_type': 'video', 'duration_minutes': '30'}))
        db.add(SocialResource(kind='base', org_id=org, code='B-OK', name='实训基地', state='DRAFT', data={'usages': ['practice'], 'established_on': '2026-05-01'}))
        db.add(SocialResource(kind='base', org_id=org, code='B-APPLY', name='审批中基地', state='COLLEGE', data={'application': True, 'usages': ['practice']}))
    d = c.get('/api/social/report-export/check', params=RANGE).json()
    assert d['ready'] and d['blocked'] == 0
    assert c.get('/api/social/report-export/check', params={'start': '2020-01-01', 'end': '2020-12-31'}).json()['count'] == 0
    r = c.get('/api/social/report-export', params=RANGE)
    assert r.status_code == 200, r.text
    book = load_workbook(BytesIO(r.content))
    assert book.sheetnames == ['项目清单', '绩效数据']
    head, row = [x.value for x in book['项目清单'][1]], [x.value for x in book['项目清单'][2]]
    got = dict(zip(head, row))
    assert got['服务面向类型'] == '国家战略' and got['总学时'] == 16 and got['教学实施方式'] == '课堂教学'
    assert got['视频资源总时长（分钟）'] == 30 and got['实践教学基地数量（个）'] == 0
    perf = dict(zip([x.value for x in book['绩效数据'][1]], [x.value for x in book['绩效数据'][2]]))
    assert perf['统计范围'] == '全校' and perf['项目数（个）'] == 1 and perf['实践教学基地数量（个）'] == 1 and perf['实践教学基地名称'] == '实训基地'
    assert perf['非讲授式教学学时占比（%）'] == 75 and perf['近一年新增实践教学基地数（个）'] == 1
    # 学院只导本院；另一个学院看不到这个项目。
    auth(c, 'collegehead')
    assert c.get('/api/social/report-export/check', params=RANGE).json()['count'] == 1
    auth(c, 'otherdept')
    assert c.get('/api/social/report-export/check', params=RANGE).status_code == 403


def test_projects_not_yet_filed_are_left_out(setup_social, client2, app2):
    c = client2
    fill(app2, setup_social)
    with app2.state.sessions.begin() as db:
        db.get(SocialProject, setup_social).state = 'COLLEGE'
    auth(c, 'reviewer')
    d = c.get('/api/social/report-export/check', params=RANGE).json()
    assert d['count'] == 0 and not d['ready']
    assert c.get('/api/social/report-export', params=RANGE).status_code == 409


@pytest.mark.parametrize('declared,expected',[('60',60),('0',0),(None,30)])
def test_project_and_performance_video_minutes_share_one_basis(declared,expected):
    project=SimpleNamespace(data={'name':'合成项目','code':'P1','total_hours':'8','video_minutes':declared},cohort='')
    courses={'C1':{'project_code':'P1','minutes':Decimal(30),'hours':Decimal(8),'mode':'practice','created_at':datetime(2026,1,1)}}
    before=deepcopy(project.data)
    db=SimpleNamespace(scalars=lambda query:SimpleNamespace(all=lambda:[]))
    actor=SimpleNamespace(org_ids=['school'],role=SimpleNamespace(kind='school'))
    row=social_report.project_row(project,courses)
    performance=social_report.performance(db,actor,date(2026,1,1),date(2026,12,31),[project],courses)
    assert row[3]==performance[4]==expected and project.data==before
    performance=social_report.performance(db,actor,date(2026,1,1),date(2026,12,31),[project,project],courses)
    assert performance[4]==expected*2


def test_incomplete_course_mode_does_not_publish_a_partial_percentage():
    courses={'C1':{'project_code':'P1','minutes':Decimal(0),'hours':Decimal(8),'mode':'practice','created_at':datetime(2026,1,1)},
             'C2':{'project_code':'P1','minutes':Decimal(0),'hours':Decimal(8),'mode':None,'created_at':datetime(2026,1,1)}}
    db=SimpleNamespace(scalars=lambda query:SimpleNamespace(all=lambda:[]))
    actor=SimpleNamespace(org_ids=['school'],role=SimpleNamespace(kind='school'))
    assert '待补充' in social_report.performance(db,actor,date(2026,1,1),date(2026,12,31),[],courses)[-1]
