"""Current-source audit regressions; synthetic isolated fixtures only."""
from datetime import date
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
import pytest
from sqlalchemy import select
from conftest import auth, post, put
from test_social_training import setup_social, recognized, row, import_rows, close_project, proof, detail, action, payload
from test_social_rules_extra import closeout
from test_ledgers import sheet, upload, checked
from app.models import Role, Qualification, User, PointEntry, Attachment
from app.social_models import SocialProject, SocialResource, SocialClaim, SocialUnit, SocialAnnual
from app import social_training as st


def submit_sheet(client, rows, year=2026):
    auth(client, 'collegehead')
    r = upload(client, '/api/social/ledgers', sheet(rows=rows, year=year), org_id='college1', year=str(year), reason='隔离审计样本')
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.parametrize('state,label',[('SUBMITTED','候选等级'),('ARCHIVED','认定等级')])
def test_export_grade_label_matches_frozen_recognition(state,label):
    from app.social_export import workbook
    data=workbook({'item':{'state':state,'version':1,'data':{}},'state_label':'合成状态','trial':{'grade':'合格'}})
    book=load_workbook(BytesIO(data),read_only=True)
    assert (label,'合格') in list(book['业务概要'].values)


def test_revoked_qualification_permission_must_also_apply_to_mine(app2, client2):
    with app2.state.sessions.begin() as db:
        db.add(Qualification(teacher_id='teacher', category='双师型初级', name='隔离资质', certificate_no='SYNTHETIC-1', issuer='隔离机构', obtained=date(2026,1,1), state='APPROVED'))
    auth(client2, 'teacher')
    with app2.state.sessions.begin() as db:
        role = db.get(Role, 'role-teacher')
        role.permissions = [p for p in role.permissions if p != 'qualification.read']
    assert client2.get('/api/qualifications').status_code == 403
    response = client2.get('/api/social/mine?year=2026')
    assert response.status_code == 403 or response.json()['qualifications'] == [], '撤销资质读取权限后，个人培训接口仍返回完整资质条目'


def test_participating_teacher_can_open_project_link(setup_social, client2, app2):
    pid=setup_social
    unit=row();unit['teachers'].append({'teacher_id':'depthead','staff_type':'external','role':'instructor'})
    assert import_rows(client2, pid, [unit]).status_code == 200
    with app2.state.sessions.begin() as db:
        p=db.get(SocialProject,pid); p.owner_id='depthead'; p.data={**p.data,'owner_id':'depthead'}
    auth(client2,'teacher')
    assert any(p['id']==pid for p in client2.get('/api/social/mine?year=2026').json()['projects'])
    response=client2.get('/api/social/objects/social_project/'+pid)
    assert response.status_code==200, '本人授课项目显示了查看按钮，点击返回 '+str(response.status_code)
    body=response.json();assert body['readonly'] is True
    assert not body.get('attachments') and not body.get('events')
    assert not body['item'].get('closeout')
    assert 'expected_income' not in body['item']['data']
    text=response.text
    assert 'L0' not in text
    assert all(not unit['data'].get('learners') for unit in body['facts']['units'])
    assert '"teacher_id":"depthead"' not in text
    denied=put(client2,'/api/social/projects/'+pid,{**payload(),'expected_version':body['item']['version']})
    assert denied.status_code==403,denied.text
    auth(client2,'otherteacher')
    assert client2.get('/api/social/objects/social_project/'+pid).status_code==403


def test_quick_annual_income_uses_confirmed_project_receipt(setup_social, client2):
    pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200
    close_project(client2,pid)
    auth(client2,'schooladmin')
    actual=next(x for x in client2.get('/api/social/targets?year=2026').json()['items'] if x['org_id']=='college1')['actual']['income']
    assert Decimal(actual)==Decimal('10000'), '已结项项目到账10000元，速算却使用无关旧教师台账100元：'+actual


@pytest.mark.parametrize('count,minutes',[(0,'0'),(12,'95.5')])
def test_video_summary_matches_imported_video_count(setup_social, client2,count,minutes):
    data={'reason':'隔离审计导入','rows':[{'code':'AUDIT-CX','name':'隔离视频课程','project_code':'SOC-P1','source':'chaoxing','teacher':'测试教师一','teacher_id':'teacher','video_count':count,'video_minutes':minutes,'reason':'隔离审计导入'}]}
    preview=post(client2,'/api/social/courses/import/preview',data)
    assert preview.status_code==200,preview.text
    result=post(client2,'/api/social/courses/import/confirm',{**data,'preview_hash':preview.json()['preview_hash']})
    assert result.status_code==200,result.text
    course=client2.get('/api/social/workspace?section=courses').json()['course_stats'].get('AUDIT-CX',{'videos':0})
    assert course['videos']==count
    summary=next(x for x in client2.get('/api/social/resource-stats').json()['items'] if x['org_id']=='college1')
    assert summary['videos']==course['videos'], '课程管理显示12个视频，资源统计只算1个'
    assert client2.get('/api/social/overview?year=2026').json()['resources']['videos']==count
    auth(client2,'teacher')
    assert client2.get('/api/social/mine?year=2026').json()['courses']['videos']==count


def test_zeroed_project_does_not_show_original_personal_count_as_approved(client2):
    ledger=submit_sheet(client2,[[1,'隔离无效项目',100,8,100,'测试教师一：100','签到表','']])
    auth(client2,'schooladmin')
    r=post(client2,f'/api/social/ledgers/{ledger["id"]}/review',{'expected_version':ledger['version'],'rows':[{'row_no':4,'approved':'0','reason':'材料不符整项剔除'}],'reason':'隔离核定'})
    assert r.status_code==200,r.text
    auth(client2,'teacher')
    allocation=client2.get('/api/social/mine?year=2026').json()['allocations'][0]
    assert not allocation['reviewed'] or allocation['value']=='0', '项目核定为0，个人仍显示100且标记已核定：'+str(allocation)
    assert allocation['original_value']=='100' and allocation['approved_value'] is None
    assert allocation['reviewed'] is False and allocation['review_status']
    assert '项目核定 0 人次' in allocation['review_status'] and '不能作为个人认定依据' in allocation['review_status']


def test_annual_achievements_do_not_reuse_other_year(setup_social,client2):
    assert import_rows(client2,setup_social,[row()]).status_code==200
    close_project(client2,setup_social)
    auth(client2,'schooladmin')
    levels=next(x for x in client2.get('/api/social/targets?year=2025').json()['items'] if x['org_id']=='college1')['achievements']
    assert levels['school']==0, '2026结项自动成果计入2025年度：'+str(levels)


def test_new_ledger_version_does_not_confuse_row_position_with_project_identity(client2):
    def make(ledger):
        result=post(client2,f'/api/social/ledgers/{ledger["id"]}/projects',{'expected_version':ledger['version'],'reason':'隔离生成草稿'})
        assert result.status_code==200,result.text
        return result.json()
    old=submit_sheet(client2,[[1,'项目甲',30,8,30,'测试教师一：30','签到表',''],[2,'项目乙',30,8,30,'测试教师一：30','签到表','']])
    assert len(make(old)['made'])==2
    newer=submit_sheet(client2,[[1,'项目丙',30,8,30,'测试教师一：30','签到表',''],[2,'项目甲',30,8,30,'测试教师一：30','签到表',''],[3,'项目乙',30,8,30,'测试教师一：30','签到表','']])
    make(newer)
    names=[p['data']['name'] for p in client2.get('/api/social/workspace?section=projects').json()['projects']]
    assert sorted(names)==['项目丙','项目乙','项目甲'], '新版本插入一行后漏掉项目丙并重复项目乙：'+str(names)


def test_identical_names_do_not_share_personal_allocation(app2,client2):
    with app2.state.sessions.begin() as db:
        u=db.get(User,'otherteacher');u.name='测试教师一';u.org_id='dept1'
    ledger=submit_sheet(client2,[[1,'隔离项目',30,8,30,'测试教师一：30','签到表','']])
    auth(client2,'teacher')
    assert client2.get('/api/social/mine?year=2026').json()['allocations']==[]
    auth(client2,'otherteacher')
    allocations=client2.get('/api/social/mine?year=2026').json()['allocations']
    assert allocations==[], '另一个同名工号继承了同一条个人分配；应先按工号确认匹配'
    auth(client2,'collegehead')
    linked=post(client2,f'/api/social/ledgers/{ledger["id"]}/rows/4/person',{'expected_version':ledger['version'],'allocation_index':0,'teacher_id':'teacher','reason':'核验工号J001明确归属'})
    assert linked.status_code==200,linked.text
    auth(client2,'teacher')
    assert client2.get('/api/social/mine?year=2026').json()['allocations'][0]['original_value']=='30'
    auth(client2,'otherteacher')
    assert client2.get('/api/social/mine?year=2026').json()['allocations']==[]


def test_declared_college_conflict_blocks_ledger_review(client2):
    ledger=submit_sheet(client2,[[1,'隔离项目',30,8,30,'测试教师一：30','签到表','']])
    auth(client2,'schooladmin')
    declared=post(client2,'/api/social/conflicts',{'target_kind':'org','target_id':'college1','reason':'隔离利益冲突申报'})
    assert declared.status_code==200,declared.text
    result=post(client2,f'/api/social/ledgers/{ledger["id"]}/review',{'expected_version':ledger['version'],'rows':[{'row_no':4,'approved':'30'}],'reason':'隔离核定'})
    assert result.status_code==403, '已声明与学院存在利益冲突，仍可核定该学院台账：'+str(result.status_code)


def test_positive_income_cannot_close_without_distribution(setup_social,client2,app2):
    pid=setup_social
    assert import_rows(client2,pid,[row()]).status_code==200
    result=action(client2,'projects','social_project',pid,'close_submit',closeout=closeout())
    assert result.status_code in (409,422),result.text
    with app2.state.sessions() as db:
        project=db.get(SocialProject,pid)
        assert project.state=='FILED', '不完整分配不得改变原项目状态'


@pytest.mark.parametrize('permission,section',[('records.read','records'),('points.read','points'),('qualification.read','qualifications')])
def test_mine_revocation_applies_to_each_section(recognized,client2,app2,permission,section):
    with app2.state.sessions.begin() as db:
        db.add(Qualification(teacher_id='teacher',category='双师型初级',name='隔离撤权资质',certificate_no='REVOKE',issuer='隔离机构',obtained=date(2026,1,1),state='APPROVED'))
    auth(client2,'teacher')
    before=client2.get('/api/social/mine?year=2026').json()
    assert before['projects'] and before['claims'] and before['qualifications']
    with app2.state.sessions.begin() as db:
        role=db.get(Role,'role-teacher');role.permissions=[p for p in role.permissions if p!=permission]
    response=client2.get('/api/social/mine?year=2026');assert response.status_code==200,response.text
    body=response.json();assert body['sections'][section] is False
    if section=='records':assert body['projects']==[] and body['allocations']==[] and body['achievements']==[]
    elif section=='points':assert body['claims']==[]
    else:assert body['qualifications']==[] and not any('隔离撤权资质' in r['text'] for r in body['reminders'])
    overview=client2.get('/api/social/overview?year=2026');assert overview.status_code==200,overview.text
    if section=='points':assert overview.json()['points'] is None


@pytest.mark.parametrize('kind,target',[('org','college1'),('teacher','teacher')])
def test_declared_conflict_blocks_correction_approval(recognized,client2,kind,target):
    claim_id,_=recognized
    auth(client2,'teacher')
    result=post(client2,'/api/social/corrections',{'claim_id':claim_id,'period_id':'period2026','teaching':'-0.56','auxiliary':'0','decision_reference':'隔离更正决定','reason':'更正回避边界'})
    assert result.status_code==200,result.text
    item=result.json()['item'];ident=item['id'];proof(client2,'social_resource',ident,item['version'])
    assert action(client2,'corrections','social_resource',ident,'submit').status_code==200
    auth(client2,'collegehead');assert action(client2,'corrections','social_resource',ident,'approve').status_code==200
    auth(client2,'reviewer')
    declared=post(client2,'/api/social/conflicts',{'target_kind':kind,'target_id':target,'reason':'本人有潜在利益冲突'})
    assert declared.status_code==200,declared.text
    denied=action(client2,'corrections','social_resource',ident,'approve')
    assert denied.status_code==403 and '利益冲突' in denied.text
    assert detail(client2,'social_resource',ident)['item']['state']=='DEAN'


@pytest.mark.parametrize('values,allowed',[
    ({'actual_income':'0','funding_source':None},True),
    ({'funding_source':None,'retained':'2000','personnel':'6000'},False),
    ({'funding_source':'government_special','funding_basis':'','retained':'0','personnel':'10000'},False),
    ({'funding_source':'government_special','funding_basis':'合成主管部门专项资金办法及财务确认','retained':'0','personnel':'10000'},True),
])
def test_zero_income_and_government_funding_require_explicit_basis(setup_social,client2,values,allowed):
    pid=setup_social;assert import_rows(client2,pid,[row()]).status_code==200
    response=action(client2,'projects','social_project',pid,'close_submit',closeout=closeout(**values))
    if not allowed:
        assert response.status_code in (409,422),response.text
        return
    assert response.status_code==200,response.text
    response=action(client2,'projects','social_project',pid,'close')
    assert response.status_code==200,response.text
    assert response.json()['item']['state']=='CLOSED'


def test_same_named_project_rows_need_explicit_link_or_new(client2):
    old=submit_sheet(client2,[[1,'同名培训',30,8,30,'测试教师一：30','签到表','第一期'],[2,'同名培训',40,8,40,'测试教师一：40','签到表','第二期']])
    assert all(r['project_link_status']=='pending' for r in old['data']['rows'])
    for row_no in (4,5):
        explicit=post(client2,f'/api/social/ledgers/{old["id"]}/rows/{row_no}/project-link',{'expected_version':old['version'],'action':'new','reason':'核对原件确认本行为独立新期次'})
        assert explicit.status_code==200,explicit.text
        old=explicit.json()
    made=post(client2,f'/api/social/ledgers/{old["id"]}/projects',{'expected_version':old['version'],'reason':'首次生成两个不同期次'})
    assert made.status_code==200,made.text
    initial=made.json()['made'];assert len(initial)==2
    newer=submit_sheet(client2,[[1,'同名培训',30,8,30,'测试教师一：30','签到表','换版后期次未明确']])
    assert newer['data']['rows'][0]['project_link_status']=='pending'
    result=post(client2,f'/api/social/ledgers/{newer["id"]}/projects',{'expected_version':newer['version'],'reason':'尝试自动猜测项目'})
    assert result.status_code==200,result.text
    assert result.json()['made']==[] and len(client2.get('/api/social/workspace?section=projects').json()['projects'])==2
    linked=post(client2,f'/api/social/ledgers/{newer["id"]}/rows/4/project-link',{'expected_version':result.json()['version'],'action':'link','project_id':initial[0]['project_id'],'reason':'按原件明确关联第一期'})
    assert linked.status_code==200,linked.text
    assert linked.json()['data']['rows'][0]['project_id']==initial[0]['project_id']


def test_undated_achievement_is_pending_not_counted_in_every_year(client2,app2):
    with app2.state.sessions.begin() as db:
        db.add(SocialResource(kind='achievement',org_id='college1',code='UNDATED',name='缺确认日期的历史成果',data={'level':'national','recognition':'历史依据待补日期'}))
    auth(client2,'schooladmin')
    for year in (2025,2026):
        item=next(x for x in client2.get(f'/api/social/targets?year={year}').json()['items'] if x['org_id']=='college1')
        assert item['achievements']['national']==0 and item['achievements_pending']==1
        overview=client2.get(f'/api/social/overview?year={year}').json()
        assert overview['achievements']['levels']['national']==0


def test_large_numeric_headcount_is_not_proof_of_online_views():
    item=checked(sheet('2025',[[1,'多期线下培训汇总',10000,8,1,10000,'','2,3,4','']]))['rows'][0]
    assert 'ONLINE_VIEWS' not in {x['code'] for x in item['issues']}
    assert Decimal(item['recommended'])>0


def test_hours_matching_module_count_does_not_prove_wrong_counting():
    item=checked(sheet(rows=[[1,'两个模块每模块1课时',100,2,200,'','2,3,4','已列两个模块，需核验备案']]))['rows'][0]
    assert not any('把课时数当成了模块数' in x['text'] for x in item['issues'])
    assert item.get('approved') is None and item['needs_mode_confirmation']
    assert item['recommended']==item['claimed']=='200'


def test_original_workbook_retention_and_authorized_download(client2,app2):
    data=sheet(rows=[[1,'原件留存样本',30,8,30,'','2,3,4','']])
    auth(client2,'collegehead')
    response=upload(client2,'/api/social/ledgers',data,org_id='college1',year='2026',reason='隔离原件留存验证')
    assert response.status_code==200,response.text
    metadata=response.json()['data']['file'];assert metadata['id']
    files=[p for p in Path(app2.state.settings.storage_dir).rglob('*') if p.is_file()]
    assert any(sha256(p.read_bytes()).digest()==sha256(data).digest() for p in files)
    downloaded=client2.get(metadata['download_url']);assert downloaded.status_code==200,downloaded.text
    assert downloaded.content==data
    auth(client2,'otherteacher')
    assert client2.get(metadata['download_url']).status_code in (403,404)
    # 学院台账原件含全院人次分配，教师不能下载（部署隐患排查第 5 条）
    auth(client2,'teacher')
    assert client2.get(metadata['download_url']).status_code==403


def test_generated_unknown_dates_cannot_be_filed_without_confirmation(setup_social,client2,app2):
    ledger=submit_sheet(client2,[[1,'待补培训日期',30,8,30,'','2,3,4','']])
    made=post(client2,f'/api/social/ledgers/{ledger["id"]}/projects',{'expected_version':ledger['version'],'reason':'生成隔离待补草稿'})
    assert made.status_code==200,made.text
    pid=made.json()['made'][0]['project_id'];p=detail(client2,'social_project',pid)['item']
    assert p['data']['source']['dates_pending'] is True
    proof(client2,'social_project',pid,p['version'])
    submitted=action(client2,'projects','social_project',pid,'submit')
    assert submitted.status_code in (409,422),submitted.text
    with app2.state.sessions.begin() as db:db.get(SocialProject,pid).state='SUBMITTED'
    auth(client2,'reviewer')
    filed=action(client2,'projects','social_project',pid,'file')
    assert filed.status_code in (409,422),filed.text


def test_person_binding_rejects_foreign_teacher_and_stale_version_and_figure_recusal(client2):
    ledger=submit_sheet(client2,[[1,'隔离人员绑定',30,8,30,'测试教师一：30','签到表','']])
    path=f'/api/social/ledgers/{ledger["id"]}/rows/4/person'
    data={'expected_version':ledger['version'],'allocation_index':0,'teacher_id':'otherteacher','reason':'跨学院工号不能关联'}
    assert post(client2,path,data).status_code==422
    result=post(client2,path,{**data,'teacher_id':'teacher','reason':'核验工号J001'})
    assert result.status_code==200,result.text
    assert post(client2,path,{**data,'teacher_id':'teacher'}).status_code==409
    auth(client2,'schooladmin')
    assert post(client2,'/api/social/conflicts',{'target_kind':'teacher','target_id':'teacher','reason':'亲属申报回避'}).status_code==200
    denied=post(client2,f'/api/social/ledgers/{ledger["id"]}/school-figure',{'expected_version':result.json()['version'],'value':'30','reason':'有冲突不得录入正式核定数'})
    assert denied.status_code==403,denied.text


def test_existing_close_review_cannot_bypass_nonzero_funding_validation(setup_social,client2,app2):
    pid=setup_social;assert import_rows(client2,pid,[row()]).status_code==200
    with app2.state.sessions.begin() as db:
        project=db.get(SocialProject,pid);project.state='CLOSE_REVIEW';project.closeout=closeout()
    rejected=action(client2,'projects','social_project',pid,'close')
    assert rejected.status_code in (409,422),rejected.text
    with app2.state.sessions() as db:assert db.get(SocialProject,pid).state=='CLOSE_REVIEW'


def test_nonowner_claim_snapshot_and_workspace_are_redacted_without_changing_archive(recognized,client2,app2):
    ident,_=recognized
    with app2.state.sessions.begin() as db:
        claim=db.get(SocialClaim,ident);original=claim.snapshot
        project=db.get(SocialProject,claim.project_id);project.owner_id='depthead';project.data={**project.data,'owner_id':'depthead'}
    auth(client2,'teacher')
    for path in ('/api/social/objects/social_claim/'+ident,'/api/social/workspace?section=recognition'):
        response=client2.get(path);assert response.status_code==200,response.text
        assert '"reference":"L0"' not in response.text
        assert 'actual_income' not in response.text and 'finance_reference' not in response.text
    with app2.state.sessions() as db:assert db.get(SocialClaim,ident).snapshot==original


@pytest.mark.parametrize('confirmed_on,score',[(None,'2.00'),('2025-09-02','2.00'),('2026-09-02','10.00')])
def test_formal_annual_trial_blocks_undated_or_other_year_achievement(recognized,client2,confirmed_on,score):
    _,policy_id=recognized;auth(client2,'collegehead')
    response=post(client2,'/api/social/annuals',{'org_id':'college1','period_id':'period2026','policy_id':policy_id,'reason':'成果年度隔离验证','data':{'self_report':'合成自评','income_target':'10000','participants_target':30,'priority_target':30,'target_reference':'合成年度任务','achievements':[{'reference':'MANUAL-ANNUAL','name':'人工登记国家级成果','level':'national','confirmed_on':confirmed_on,'recognition_reference':'合成成果确认依据'}]}})
    assert response.status_code==200,response.text
    trial=response.json()['trial'];assert trial['achievement_base']==score
    if confirmed_on!='2026-09-02':assert any('成果' in x for x in trial['blockers'])


@pytest.mark.parametrize('self_participation',[True,False])
def test_project_reviewer_recuses_for_internal_participants(setup_social,client2,app2,self_participation):
    pid=setup_social;unit=row()
    if self_participation:
        with app2.state.sessions.begin() as db:db.get(User,'reviewer').staff_no='AUDIT-INTERNAL'
        unit['teachers'].append({'teacher_id':'reviewer','staff_type':'internal','role':'instructor'})
    assert import_rows(client2,pid,[unit]).status_code==200
    with app2.state.sessions.begin() as db:
        project=db.get(SocialProject,pid);project.owner_id='depthead';project.data={**project.data,'owner_id':'depthead'}
    assert action(client2,'projects','social_project',pid,'close_submit',closeout=closeout(retained='2000',personnel='6000')).status_code==200
    if not self_participation:
        assert post(client2,'/api/social/conflicts',{'target_kind':'teacher','target_id':'teacher','reason':'与授课人有利益冲突，其不是负责人'}).status_code==200
    denied=action(client2,'projects','social_project',pid,'close')
    assert denied.status_code==403,denied.text
    with app2.state.sessions() as db:assert db.get(SocialProject,pid).state=='CLOSE_REVIEW'


def test_project_export_requires_read_permission_even_when_export_is_kept(setup_social,client2,app2):
    auth(client2,'schooladmin')
    with app2.state.sessions.begin() as db:
        role=db.get(Role,'role-school');assert 'export' in role.permissions
        role.permissions=[p for p in role.permissions if p!='records.read']
    assert client2.get('/api/social/objects/social_project/'+setup_social+'/export').status_code==403


def test_participant_project_export_contains_only_redacted_information(setup_social,client2,app2):
    pid=setup_social;unit=row();unit['teachers'].append({'teacher_id':'depthead','staff_type':'external','role':'instructor'})
    assert import_rows(client2,pid,[unit]).status_code==200
    with app2.state.sessions.begin() as db:
        project=db.get(SocialProject,pid);project.owner_id='depthead';project.data={**project.data,'owner_id':'depthead','expected_income':'987654321'}
        project.closeout={'actual_income':'987654321','finance_reference':'保密财务凭据'}
    auth(client2,'teacher')
    response=client2.get('/api/social/objects/social_project/'+pid+'/export');assert response.status_code==200,response.text
    book=load_workbook(BytesIO(response.content),read_only=True)
    values=[str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
    assert 'L0' not in values and 'depthead' not in values and '987654321' not in values and '保密财务凭据' not in values
    assert '隔离技能培训' in values


def test_claim_approval_recuses_another_internal_participant(recognized,client2,app2):
    ident,_=recognized
    with app2.state.sessions.begin() as db:
        claim=db.get(SocialClaim,ident);claim.state='SUBMITTED'
        unit=db.scalar(select(SocialUnit).where(SocialUnit.project_id==claim.project_id))
        db.get(User,'collegehead').staff_no='AUDIT-DEAN'
        unit.data={**unit.data,'teachers':[*unit.data['teachers'],{'teacher_id':'collegehead','staff_type':'internal','role':'instructor'}]}
        db.flush();claim.snapshot=st.freeze(st.point_trial(db,claim))
    auth(client2,'collegehead')
    denied=action(client2,'claims','social_claim',ident,'approve')
    assert denied.status_code==403,denied.text
    with app2.state.sessions() as db:assert db.get(SocialClaim,ident).state=='SUBMITTED'


def test_legacy_points_read_redacts_social_source_without_changing_point_entry(recognized,client2,app2):
    ident,_=recognized
    with app2.state.sessions.begin() as db:
        claim=db.get(SocialClaim,ident);project=db.get(SocialProject,claim.project_id)
        project.owner_id='depthead';project.data={**project.data,'owner_id':'depthead'}
        entry=db.scalar(select(PointEntry).where(PointEntry.kind=='SOCIAL'));entry_id=entry.id;original=entry.source_snapshot
    auth(client2,'teacher')
    for path in ('/api/points?period_id=period2026','/api/points/'+entry_id):
        response=client2.get(path)
        assert response.status_code==200,response.text
        assert 'L0' not in response.text and 'finance_reference' not in response.text and 'actual_income' not in response.text
    assert '1.56' in client2.get('/api/points?period_id=period2026').text
    with app2.state.sessions() as db:assert db.get(PointEntry,entry_id).source_snapshot==original


@pytest.mark.parametrize('actor_id',['teacher','collegehead','schooladmin'])
def test_legacy_points_revoked_records_permission_redacts_source_for_all_roles(recognized,client2,app2,actor_id):
    claim_id,_=recognized
    with app2.state.sessions.begin() as db:
        claim=db.get(SocialClaim,claim_id)
        assert db.get(SocialProject,claim.project_id).owner_id=='teacher'
        entry=db.scalar(select(PointEntry).where(PointEntry.kind=='SOCIAL',PointEntry.teacher_id=='teacher',PointEntry.item_key=='social_teaching_'+claim_id))
        source=deepcopy(entry.source_snapshot);trial=source['recognition']['trial']
        trial['basis']['facts']['units'][0]['data']['teachers'].append({'teacher_id':'OTHER-TEACHER-SECRET','staff_type':'external','role':'instructor'})
        source['recognition']=st.freeze(trial);entry.source_snapshot=source
        entry_id=entry.id;original=deepcopy(source);amount=entry.amount
    auth(client2,actor_id)
    before=client2.get('/api/points/'+entry_id);assert before.status_code==200,before.text
    assert 'L0' in before.text and 'OTHER-TEACHER-SECRET' in before.text and 'finance_reference' in before.text
    with app2.state.sessions.begin() as db:
        role=db.get(Role,db.get(User,actor_id).role_id)
        assert 'points.read' in role.permissions and 'records.read' in role.permissions
        role.permissions=[p for p in role.permissions if p!='records.read']
    for path in ('/api/points?period_id=period2026','/api/points/'+entry_id):
        response=client2.get(path);assert response.status_code==200,response.text
        item=next(x for x in response.json()['items'] if x['id']==entry_id) if 'items' in response.json() else response.json()
        assert Decimal(item['amount'])==amount==Decimal('1.56')
        recognition=item['source_snapshot']['recognition']
        assert recognition['sha256']==original['recognition']['sha256']
        units=recognition['trial']['basis']['facts']['units'];assert len(units)==1
        assert units[0]['data']['counts']=={'effective':30,'excluded':0,'priority':30}
        assert units[0]['data']['learners']==[]
        assert [t['teacher_id'] for t in units[0]['data']['teachers']]==['teacher']
        assert all(secret not in response.text for secret in ('L0','OTHER-TEACHER-SECRET','finance_reference','actual_income'))
    with app2.state.sessions() as db:
        saved=db.get(PointEntry,entry_id)
        assert saved.source_snapshot==original and saved.amount==amount
        assert st.digest(saved.source_snapshot['recognition']['trial'])==original['recognition']['sha256']


@pytest.mark.parametrize('state',['PUBLIC','ARCHIVED'])
def test_teacher_annual_public_summary_hides_evidence_and_attachment(recognized,client2,app2,state):
    claim_id,policy_id=recognized
    with app2.state.sessions.begin() as db:
        claim=db.get(SocialClaim,claim_id)
        unit=db.scalar(select(SocialUnit).where(SocialUnit.project_id==claim.project_id))
        unit.data={**unit.data,'learners':[{**unit.data['learners'][0],'reference':'OTHER-LEARNER-SECRET'},*unit.data['learners'][1:]]}
    auth(client2,'collegehead')
    response=post(client2,'/api/social/annuals',{'org_id':'college1','period_id':'period2026','policy_id':policy_id,'reason':'已公示年度读取权限隔离样本','data':{'self_report':'合成年度自评','income_target':'10000','participants_target':30,'priority_target':30,'target_reference':'合成年度任务'}})
    assert response.status_code==200,response.text
    item=response.json()['item'];ident=item['id'];proof(client2,'social_annual',ident,item['version'])
    with app2.state.sessions.begin() as db:
        annual=db.get(SocialAnnual,ident);annual.state=state
        annual.snapshot={**st.freeze(st.annual_trial(db,annual)),'public_start':'2026-09-03','public_end':'2026-09-10'}
        original=annual.snapshot
        attachment=db.scalar(select(Attachment).where(Attachment.object_type=='social_annual',Attachment.object_id==ident));file_id=attachment.id
    auth(client2,'teacher')
    detail_response=client2.get('/api/social/objects/social_annual/'+ident)
    assert detail_response.status_code==200,detail_response.text
    public=detail_response.json();assert public['public_summary'] and public['item']['data']=={}
    assert public['attachments']==[] and public['events']==[] and 'basis' not in public['trial']
    assert not {'points_total','points_average','cost_details','veto_basis'} & public['trial'].keys()
    workspace=client2.get('/api/social/workspace?section=annual');assert workspace.status_code==200,workspace.text
    shown=next(x for x in workspace.json()['annuals'] if x['id']==ident)
    assert shown['data']=={} and 'basis' not in shown['snapshot']['trial']
    assert 'OTHER-LEARNER-SECRET' not in detail_response.text
    assert client2.get('/api/attachments/'+file_id+'/download').status_code==403
    exported=client2.get('/api/social/objects/social_annual/'+ident+'/export');assert exported.status_code==200,exported.text
    book=load_workbook(BytesIO(exported.content),read_only=True)
    assert not any('OTHER-LEARNER-SECRET' in str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row)
    auth(client2,'schooladmin')
    assert 'OTHER-LEARNER-SECRET' in client2.get('/api/social/objects/social_annual/'+ident).text
    assert client2.get('/api/attachments/'+file_id+'/download').status_code==200
    with app2.state.sessions() as db:
        annual=db.get(SocialAnnual,ident);assert annual.snapshot==original and annual.state==state
