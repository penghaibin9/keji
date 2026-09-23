"""HTTP regression tests for the implementation center; synthetic identities only."""
import io
import json
import zipfile
from datetime import timedelta
from urllib.parse import urlsplit, parse_qs
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.models import Connection, Setting, AuthSession, User, utcnow
from app.integrations import ADAPTERS, ConnectionCheck, configuration_hash, credential_fingerprint
from app.inbound_sync import IntegrationSchedule
from app.adapters.contracts import AdapterContractError
from conftest import auth, post, put

REASON = {'reason': '独立合成实施验收', 'basis': '接口收口专项'}


def connection_id(client, code='sso'):
    return next(x['connection']['id'] for x in client.get('/api/connections/catalog').json()['items']
                if x['connection']['code'] == code)


def sso_config():
    return {'endpoint': 'https://synthetic.invalid', 'owner': '合成学校接口负责人',
        'secret_env': 'ND_CLOSEOUT_TRIAL', 'adapter': {'protocol': 'cas', 'login_path': '/login',
            'validate_path': '/serviceValidate', 'identity_fields': {'account': 'account',
            'name': 'name', 'active': 'active', 'org_code': 'org', 'stable_id': 'id'}}}


def make_draft(client, ident, config=None):
    active = client.get('/api/connections/'+ident+'/draft').json()
    response = put(client, '/api/connections/'+ident+'/draft', {
        **REASON, 'expected_version': active['active_version'], 'expected_draft_version': active['draft_version'],
        'config': config or sso_config(), 'mappings': []})
    assert response.status_code == 200, response.text
    return {'expected_version': active['active_version'], 'draft_version': response.json()['draft_version'], **REASON}


def test_draft_save_preserves_live_connection_and_schedules(client2, app2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2, 'middle')
    with app2.state.sessions.begin() as db:
        obj = db.get(Connection, ident)
        obj.state = 'CONNECTED'
        db.add(IntegrationSchedule(connection_id=ident, data_type='training_project',
            frequency_minutes=30, auto_commit=True, active=True, approved_by='schooladmin'))
    action = make_draft(client2, ident, {'endpoint': 'https://synthetic.invalid', 'owner': '学校资料'})
    with app2.state.sessions() as db:
        assert db.get(Connection, ident).state == 'CONNECTED'
        assert db.get(Connection, ident).config == {}
        assert db.scalar(select(IntegrationSchedule)).active
        assert db.get(Connection, ident).version == action['expected_version']
    conflict = put(client2, '/api/connections/'+ident+'/draft', {**REASON,
        'expected_version': action['expected_version'], 'expected_draft_version': 0, 'config': {}, 'mappings': []})
    assert conflict.status_code == 409


class TrialAdapter:
    adapter_id = 'synthetic.trial.adapter'
    def probe(self, **kwargs):
        return {'authenticated': True, 'tls_verified': True, 'protocol_verified': True}
    def build_login_url(self, *, state, **kwargs):
        return {'redirect_url': 'https://synthetic.invalid/login?state='+state}
    def complete_login(self, **kwargs):
        return {'stable_id': 'synthetic-subject', 'account': 'teacher', 'staff_no': 'J001',
                'name': '合成教师', 'org_code': 'dept1', 'external_roles': ['external_admin'], 'active': True}


def setup_trial(client, app, monkeypatch):
    monkeypatch.setenv('ND_CLOSEOUT_TRIAL', 'synthetic-test-credential-only')
    monkeypatch.setattr(ADAPTERS, '_items', {'sso': TrialAdapter()})
    monkeypatch.setattr(ADAPTERS, '_test_codes', {'sso'})
    monkeypatch.setattr(ADAPTERS, 'production', False)
    auth(client, 'schooladmin')
    ident = connection_id(client)
    action = make_draft(client, ident)
    response = post(client, '/api/connections/'+ident+'/draft/login-test', action)
    assert response.status_code == 200, response.text
    state = parse_qs(urlsplit(response.json()['redirect_url']).query)['state'][0]
    return ident, action, state


def test_browser_trial_does_not_switch_session_or_bind_account(client2, app2, monkeypatch):
    ident, action, state = setup_trial(client2, app2, monkeypatch)
    before = client2.cookies.get(app2.state.settings.cookie_name)
    response = client2.get('/api/auth/sso/implementation-callback', params={'state': state, 'ticket': 'ST-synthetic'}, follow_redirects=False)
    assert response.status_code == 303, response.text
    assert client2.cookies.get(app2.state.settings.cookie_name) == before
    assert client2.get('/api/auth/me').json()['user']['id'] == 'schooladmin'
    with app2.state.sessions() as db:
        assert db.scalar(select(func.count()).select_from(AuthSession)) == 1
        assert not db.scalar(select(Setting).where(Setting.key.like('ssoid:%')))
        row = db.get(Setting, 'integration.draft:'+ident)
        proof = db.get(ConnectionCheck, row.value['check_id'])
        assert proof.passed and proof.mode == 'TEST_ONLY'
        assert proof.receipt['preview']['binding']['role_kind'] == 'teacher'
    published = post(client2, '/api/connections/'+ident+'/draft/publish', {**action, 'confirm_schedule_pause': True})
    assert published.status_code == 409
    replay = client2.get('/api/auth/sso/implementation-callback', params={'state': state, 'ticket': 'ST-synthetic'})
    assert replay.status_code == 401


@pytest.mark.parametrize('change_kind', ['state', 'draft', 'credential', 'revoked'])
def test_trial_rejects_changed_context(client2, app2, monkeypatch, change_kind):
    ident, action, state = setup_trial(client2, app2, monkeypatch)
    if change_kind == 'state':
        state = 'unmatched-state'
    elif change_kind == 'draft':
        make_draft(client2, ident)
    elif change_kind == 'credential':
        monkeypatch.setenv('ND_CLOSEOUT_TRIAL', 'different-synthetic-credential')
    else:
        with app2.state.sessions.begin() as db:
            db.get(User, 'schooladmin').is_admin = False
    response = client2.get('/api/auth/sso/implementation-callback', params={'state': state, 'ticket': 'ST-synthetic'}, follow_redirects=False)
    assert response.status_code in {401, 403, 409}
    with app2.state.sessions() as db:
        assert not db.scalar(select(Setting).where(Setting.key.like('ssoid:%')))
        assert db.scalar(select(func.count()).select_from(AuthSession)) == 1


def test_official_callbacks_are_server_generated_and_health_is_non_destructive(client2, app2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    action = make_draft(client2, ident)
    w = client2.get('/api/connections/'+ident+'/implementation').json()
    assert w['callbacks']['login'] == 'http://testserver/api/auth/sso/callback'
    assert w['callbacks']['test_login'] == 'http://testserver/api/auth/sso/implementation-callback'
    assert w['statuses']['acceptance'] == '待学校签认'
    response = post(client2, '/api/connections/'+ident+'/health', {**REASON, 'expected_version': action['expected_version']})
    assert response.status_code == 200, response.text
    assert response.json()['active_unchanged']
    assert response.json()['connection_version'] == action['expected_version']
    with app2.state.sessions() as db:
        assert db.get(Connection, ident).version == action['expected_version']


def test_materials_and_handover_do_not_imply_acceptance(client2, app2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    make_draft(client2, ident)
    w = client2.get('/api/connections/'+ident+'/implementation').json()
    items = [{'name': name, 'owner': '合成责任人', 'reference': 'SYNTHETIC-DOC-001',
              'confirmed_on': '2026-09-23'} for name in w['materials']['required']]
    r = put(client2, '/api/connections/'+ident+'/materials', {**REASON, 'expected_version': 0, 'items': items})
    assert r.status_code == 200, r.text
    w = client2.get('/api/connections/'+ident+'/implementation').json()
    assert w['statuses']['materials'] == '资料登记齐全'
    assert w['statuses']['acceptance'] == '待学校签认'
    archive = client2.get('/api/connections/'+ident+'/handover.zip')
    assert archive.status_code == 200
    handover = zipfile.ZipFile(io.BytesIO(archive.content))
    assert '接口实施交接工作簿.xlsx' in handover.namelist()
    data = handover.read('implementation.json').decode()
    assert '_credential_sha256' not in data and 'token_hash' not in data
    workbook = load_workbook(io.BytesIO(handover.read('接口实施交接工作簿.xlsx')), read_only=True)
    assert workbook.sheetnames == ['状态总览','资料清单','字段映射','凭据证书元数据','统一认证角色确认','权威来源台账','学校验收台账']
    payload = json.loads(data)
    assert payload['format'] == 'YUEKE_INTEGRATION_HANDOVER_R2'
    assert payload['school_acceptance']['summary']['state'] == '待学校签认'
    assert payload['credential_metadata']['values_returned'] is False


@pytest.mark.parametrize('username', ['teacher', 'depthead', 'collegehead'])
def test_non_admin_cannot_access_implementation_or_drafts(client2, app2, username):
    with app2.state.sessions() as db:
        ident = db.scalar(select(Connection.id).where(Connection.code == 'sso'))
    auth(client2, username)
    for endpoint in ('implementation', 'draft', 'handover.zip'):
        assert client2.get('/api/connections/'+ident+'/'+endpoint).status_code == 403
    r = post(client2, '/api/connections/'+ident+'/health', {**REASON, 'expected_version': 1})
    assert r.status_code == 403

def _acceptance_items(workspace, status='PASSED'):
    return [{'code': item['code'], 'status': status,
             'evidence_ref': 'SYNTHETIC-EVIDENCE-'+str(index+1),
             'school_contact': '合成学校联系人', 'checked_on': '2026-09-23'}
            for index, item in enumerate(workspace['acceptance']['summary']['requirements'])]


def test_credential_metadata_records_only_safe_lifecycle_facts(client2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    action = make_draft(client2, ident)
    payload = {**REASON, 'expected_version': 0, 'expected_connection_version': action['expected_version'],
        'expected_draft_version': action['draft_version'], 'credential_type': 'certificate',
        'managed_by': 'file_mount', 'owner': '学校服务器管理员',
        'certificate_fingerprint_sha256': 'a'*64, 'certificate_subject': 'CN=synthetic.example',
        'certificate_issuer': 'Synthetic CA', 'valid_from': '2026-09-01',
        'valid_until': '2099-12-31', 'last_rotated_on': '2026-09-01',
        'next_rotation_on': '2099-11-30'}
    result = put(client2, '/api/connections/'+ident+'/credential-metadata', payload)
    assert result.status_code == 200, result.text
    assert result.json()['secret_values_stored'] is False
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    assert workspace['credential']['metadata']['certificate_fingerprint_sha256'] == 'a'*64
    assert workspace['credential']['asset_status']['state'] == '安全资产有效'
    assert workspace['credential']['values_returned'] is False
    assert 'synthetic-test-credential-only' not in json.dumps(workspace, ensure_ascii=False)

def test_credential_metadata_rejects_bad_dates_and_secret_shaped_extra_fields(client2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    action = make_draft(client2, ident)
    base = {**REASON, 'expected_version': 0, 'expected_connection_version': action['expected_version'],
        'expected_draft_version': action['draft_version'], 'credential_type': 'certificate',
        'managed_by': 'server_env', 'owner': '合成管理员',
        'certificate_fingerprint_sha256': 'b'*64, 'certificate_subject': '',
        'certificate_issuer': '', 'valid_from': '2026-10-01',
        'valid_until': '2026-09-01', 'last_rotated_on': None, 'next_rotation_on': None}
    bad = put(client2, '/api/connections/'+ident+'/credential-metadata', base)
    assert bad.status_code == 422
    extra = put(client2, '/api/connections/'+ident+'/credential-metadata',
        {**base, 'valid_until': '2099-01-01', 'private_key': 'SHOULD-NOT-BE-ACCEPTED'})
    assert extra.status_code == 422


def test_acceptance_evidence_does_not_auto_sign_and_stales_after_config_change(client2, app2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    items = _acceptance_items(workspace)
    saved = put(client2, '/api/connections/'+ident+'/acceptance',
        {**REASON, 'expected_version': 0, 'expected_connection_version': workspace['connection_version'], 'items': items,
         'signed_by': '', 'signed_on': None, 'signed_reference': ''})
    assert saved.status_code == 200, saved.text
    assert saved.json()['summary']['passed'] == 6
    assert saved.json()['summary']['signed'] is False
    blocked = put(client2, '/api/connections/'+ident+'/acceptance',
        {**REASON, 'expected_version': 1, 'expected_connection_version': workspace['connection_version'], 'items': items,
         'signed_by': '学校验收人', 'signed_on': '2026-09-23',
         'signed_reference': 'SYNTHETIC-SIGNOFF'})
    assert blocked.status_code == 409
    with app2.state.sessions.begin() as db:
        db.get(Connection, ident).version += 1
    current = client2.get('/api/connections/'+ident+'/implementation').json()
    assert current['statuses']['acceptance'] == '配置已变化，验收需重新确认'
    assert current['acceptance']['summary']['stale'] is True


def test_province_source_gaps_block_final_signoff(client2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2, 'province')
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    items = _acceptance_items(workspace)
    result = put(client2, '/api/connections/'+ident+'/acceptance',
        {**REASON, 'expected_version': 0, 'expected_connection_version': workspace['connection_version'], 'items': items,
         'signed_by': '学校验收人', 'signed_on': '2026-09-23',
         'signed_reference': 'SYNTHETIC-SIGNOFF'})
    assert result.status_code == 409
    assert '权威来源' in result.text


@pytest.mark.parametrize('username', ['teacher', 'depthead', 'collegehead'])
def test_non_admin_cannot_write_credential_or_acceptance(client2, app2, username):
    with app2.state.sessions() as db:
        ident = db.scalar(select(Connection.id).where(Connection.code == 'sso'))
    auth(client2, username)
    credential = put(client2, '/api/connections/'+ident+'/credential-metadata', {
        **REASON, 'expected_version': 0, 'expected_connection_version': 1, 'expected_draft_version': 0, 'credential_type': 'token',
        'managed_by': 'server_env', 'owner': '无权限账号',
        'certificate_fingerprint_sha256': '', 'certificate_subject': '',
        'certificate_issuer': '', 'valid_from': None, 'valid_until': None,
        'last_rotated_on': None, 'next_rotation_on': None})
    assert credential.status_code == 403
    acceptance = put(client2, '/api/connections/'+ident+'/acceptance', {
        **REASON, 'expected_version': 0, 'expected_connection_version': 1, 'items': [],
        'signed_by': '', 'signed_on': None, 'signed_reference': ''})
    assert acceptance.status_code == 403

def test_current_real_health_and_complete_evidence_can_record_school_signoff(client2, app2, monkeypatch):
    auth(client2, 'schooladmin')
    ident = connection_id(client2, 'sso')
    monkeypatch.setenv('ND_ACCEPTANCE_SECRET', 'synthetic-acceptance-secret')
    with app2.state.sessions.begin() as db:
        obj = db.get(Connection, ident)
        obj.config = {'endpoint': 'https://synthetic.invalid', 'secret_env': 'ND_ACCEPTANCE_SECRET'}
        obj.state = 'CONNECTED'
        db.flush()
        db.add(Setting(key='integration.health:'+ident, value={
            'passed': True, 'mode': 'REAL', 'configuration_hash': configuration_hash(obj),
            '_credential_sha256': credential_fingerprint(obj)}, version=1))
        db.add(Setting(key='integration.credential:'+ident, value={
            'credential_type': 'token', 'managed_by': 'server_env', 'owner': 'synthetic-admin',
            'credential_reference': 'ND_ACCEPTANCE_SECRET'}, version=1))
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    materials = [{'name': name, 'owner': '学校信息中心', 'reference': 'DOC-'+str(index+1),
                  'confirmed_on': '2026-09-23'}
                 for index, name in enumerate(workspace['materials']['required'])]
    saved = put(client2, '/api/connections/'+ident+'/materials',
                {**REASON, 'expected_version': 0, 'items': materials})
    assert saved.status_code == 200, saved.text
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    items = _acceptance_items(workspace)
    signed = put(client2, '/api/connections/'+ident+'/acceptance', {
        **REASON, 'expected_version': 0, 'expected_connection_version': workspace['connection_version'], 'items': items,
        'signed_by': '学校验收人', 'signed_on': '2026-09-23',
        'signed_reference': 'SYNTHETIC-SIGNOFF-001'})
    assert signed.status_code == 200, signed.text
    assert signed.json()['summary']['state'] == '学校已签认'
    assert signed.json()['summary']['signed'] is True
    current = client2.get('/api/connections/'+ident+'/implementation').json()
    assert current['statuses']['acceptance'] == '学校已签认'
    monkeypatch.setenv('ND_ACCEPTANCE_SECRET', 'synthetic-rotated-after-signoff')
    runtime_changed = client2.get('/api/connections/'+ident+'/implementation').json()
    assert runtime_changed['statuses']['acceptance'] == '运行凭据或健康状态已变化，验收需复核'
    assert runtime_changed['acceptance']['summary']['signed'] is False
    monkeypatch.setenv('ND_ACCEPTANCE_SECRET', 'synthetic-acceptance-secret')
    with app2.state.sessions.begin() as db:
        obj = db.get(Connection, ident)
        obj.config = {**obj.config, 'owner': '配置发生变化'}
        obj.version += 1
    stale = client2.get('/api/connections/'+ident+'/implementation').json()
    assert stale['statuses']['acceptance'] == '配置已变化，验收需重新确认'
    assert stale['acceptance']['summary']['signed'] is False


def test_security_metadata_and_acceptance_are_version_bound(client2):
    auth(client2, 'schooladmin')
    ident = connection_id(client2)
    action = make_draft(client2, ident)
    stale_credential = put(client2, '/api/connections/'+ident+'/credential-metadata', {
        **REASON, 'expected_version': 0, 'expected_connection_version': action['expected_version'],
        'expected_draft_version': action['draft_version']-1, 'credential_type': 'token',
        'managed_by': 'server_env', 'owner': 'synthetic-admin',
        'certificate_fingerprint_sha256': '', 'certificate_subject': '', 'certificate_issuer': '',
        'valid_from': None, 'valid_until': None, 'last_rotated_on': None, 'next_rotation_on': None})
    assert stale_credential.status_code == 409
    workspace = client2.get('/api/connections/'+ident+'/implementation').json()
    stale_acceptance = put(client2, '/api/connections/'+ident+'/acceptance', {
        **REASON, 'expected_version': 0, 'expected_connection_version': workspace['connection_version']+1,
        'items': _acceptance_items(workspace), 'signed_by': '', 'signed_on': None, 'signed_reference': ''})
    assert stale_acceptance.status_code == 409
