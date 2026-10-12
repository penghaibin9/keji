"""Error import batches terminate without deleting previews or bypassing access."""
import copy
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.models import Audit, ImportJob, ImportRow, Role, TrainingRecord, User
from conftest import auth, post
from test_r3_excel import fill_template, upload_preview, REASON


def preview(client, *, error=True, key='DISCARD-TEST'):
    auth(client, 'depthead')
    values=[key,'J001','2026-09-10','DP','项目','DC','课程','DCL','班级',2,3,45,
            '无效来源' if error else 'Excel导入']
    response=upload_preview(client,'records',fill_template(client,'records',values))
    assert response.status_code==200,response.text
    return response.json()


def discard(client, job, **kwargs):
    return post(client,f"/api/imports/{job['id']}/discard",
                {'expected_version':job['version'],**REASON},**kwargs)


def test_error_import_terminates_and_discard_keeps_rows_receipt_and_audit(client2,app2):
    job=preview(client2)
    assert job['state']=='INVALID' and job['errors']==1
    assert client2.get('/api/imports').json()['items'][0]['state']=='INVALID'
    assert client2.get('/api/imports/'+job['id']).json()['job']['state']=='INVALID'
    assert post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON}).status_code==409
    with app2.state.sessions() as db:
        saved=db.get(ImportJob,job['id']);receipt=copy.deepcopy(saved.receipt)
        rows=[(r.id,r.payload,r.errors) for r in db.scalars(select(ImportRow).where(ImportRow.job_id==job['id']))]
        assert saved.state=='INVALID'
    response=discard(client2,job);assert response.status_code==200,response.text
    done=response.json();assert done['state']=='VOID' and done['version']==job['version']+1
    assert done['receipt']['written']==0 and done['receipt']['discard_reason']==REASON['reason']
    assert done['receipt']['discard_basis']==REASON['basis'] and 'access_proof' not in done['receipt']
    workbook=client2.get(f"/api/imports/{job['id']}/errors");assert workbook.status_code==200
    wb=load_workbook(BytesIO(workbook.content),read_only=True)
    try:assert '错误原因' in next(wb.active.iter_rows(values_only=True))
    finally:wb.close()
    assert post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':done['version'],**REASON}).status_code==409
    with app2.state.sessions() as db:
        saved=db.get(ImportJob,job['id'])
        assert saved.state=='VOID' and saved.sha256==job['sha256'] and saved.filename==job['filename']
        assert saved.receipt['access_proof']==receipt['access_proof']
        assert saved.receipt['preview_rows_sha256']==receipt['preview_rows_sha256']
        assert rows==[(r.id,r.payload,r.errors) for r in db.scalars(select(ImportRow).where(ImportRow.job_id==job['id']))]
        events=db.scalars(select(Audit).where(Audit.object_id==job['id'],Audit.action=='放弃Excel导入批次')).all()
        assert len(events)==1 and events[0].reason==REASON['reason']
        assert events[0].before['state']=='INVALID' and events[0].after['state']=='VOID'
        assert db.scalar(select(func.count()).select_from(TrainingRecord).where(TrainingRecord.external_key=='DISCARD-TEST'))==0


def test_legacy_error_preview_can_be_discarded_without_data_migration(client2,app2):
    job=preview(client2)
    with app2.state.sessions.begin() as db:db.get(ImportJob,job['id']).state='PREVIEW'
    assert client2.get('/api/imports').json()['items'][0]['state']=='INVALID'
    assert client2.get('/api/imports/'+job['id']).json()['job']['state']=='INVALID'
    with app2.state.sessions() as db:assert db.get(ImportJob,job['id']).state=='PREVIEW'
    response=discard(client2,job);assert response.status_code==200,response.text
    assert response.json()['state']=='VOID'
    with app2.state.sessions() as db:
        event=db.scalar(select(Audit).where(Audit.object_id==job['id'],Audit.action=='放弃Excel导入批次'))
        assert event.before['state']=='PREVIEW'


def test_discard_same_idempotency_key_replays_once_and_other_requests_reject(client2,app2):
    job=preview(client2)
    first=discard(client2,job,key='discard-repeat-key');assert first.status_code==200,first.text
    again=discard(client2,job,key='discard-repeat-key');assert again.status_code==200,again.text
    assert again.json()==first.json()
    assert discard(client2,job).status_code==409
    assert discard(client2,first.json()).status_code==409
    with app2.state.sessions() as db:
        assert db.scalar(select(func.count()).select_from(Audit).where(Audit.object_id==job['id'],Audit.action=='放弃Excel导入批次'))==1


def test_valid_import_keeps_confirm_behavior_and_cannot_be_discarded(client2,app2):
    job=preview(client2,error=False)
    assert job['state']=='PREVIEW' and job['errors']==0
    assert discard(client2,job).status_code==409
    response=post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON})
    assert response.status_code==200,response.text
    done=response.json();assert done['state']=='COMMITTED' and done['receipt']['written']==1
    assert discard(client2,done).status_code==409


def test_confirm_revalidation_failure_is_terminal_and_can_be_discarded(client2,app2):
    first=preview(client2,error=False)
    second=preview(client2,error=False)
    response=post(client2,f"/api/imports/{first['id']}/confirm",{'expected_version':first['version'],**REASON})
    assert response.status_code==200,response.text
    response=post(client2,f"/api/imports/{second['id']}/confirm",{'expected_version':second['version'],**REASON})
    assert response.status_code==200,response.text
    failed=response.json();assert failed['state']=='INVALID' and failed['errors']==1
    assert failed['receipt']['written']==0 and failed['receipt']['stale'] is True
    assert discard(client2,failed).status_code==200
    with app2.state.sessions() as db:
        assert db.scalar(select(func.count()).select_from(TrainingRecord).where(TrainingRecord.external_key=='DISCARD-TEST'))==1


@pytest.mark.parametrize('mode,status',[('owner',404),('permission',403),('scope',403),('proof',409)])
def test_discard_retains_owner_permission_scope_and_proof_checks(client2,app2,mode,status):
    job=preview(client2)
    if mode=='owner':auth(client2,'otherdept')
    else:
        with app2.state.sessions.begin() as db:
            if mode=='permission':
                role=db.get(Role,'role-dept');role.permissions=[p for p in role.permissions if p!='records.read']
            elif mode=='scope':db.get(User,'teacher').org_id='dept2'
            else:
                saved=db.get(ImportJob,job['id']);receipt=dict(saved.receipt);receipt.pop('access_proof');saved.receipt=receipt
    response=discard(client2,job);assert response.status_code==status,response.text
    with app2.state.sessions() as db:
        assert db.get(ImportJob,job['id']).state=='INVALID'
        assert db.scalar(select(func.count()).select_from(Audit).where(Audit.object_id==job['id'],Audit.action=='放弃Excel导入批次'))==0


def test_discard_requires_reason_current_version_and_csrf(client2,app2):
    job=preview(client2);url=f"/api/imports/{job['id']}/discard"
    assert post(client2,url,{'expected_version':job['version'],'reason':''}).status_code==422
    assert post(client2,url,{'expected_version':job['version']+1,**REASON}).status_code==409
    assert post(client2,url,{'expected_version':job['version'],**REASON},headers={'x-csrf-token':'invalid'}).status_code==403
    with app2.state.sessions() as db:assert db.get(ImportJob,job['id']).state=='INVALID'
