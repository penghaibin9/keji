"""超星建课人自己把登录名改成工号后，下一次全量同步自动对上；管理员能下载还没对上的名单。"""
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.reference_data import TeacherCourseTerm,TeacherHrProfile
from conftest import auth,post
from test_procurement_inbound_rollback import REASON
from test_teacher_course_inbound import ScheduleHub,typed

CX_MAP={'creater_login_name':'staff_no','teachername':'teacher_name','courseid':'course_code','coursename':'course_name','xn':'school_year','xq':'term_no','credit':'credit'}

class CxHub(ScheduleHub):
    def __init__(self,rows):
        super().__init__(rows)
        self.config['adapter']['data_types']={'online_course':self.config['adapter']['data_types']['teacher_course']}

def cx_full_run(app,rows):
    from app.integrations import AdapterRegistry
    from app.inbound_sync import IntegrationSchedule,create_pull_task,process_pull_task_once
    from app.models import Connection,Task
    from test_integration_closeout import actor
    hub=CxHub(rows);registry=AdapterRegistry();registry.register_for_test('middle',hub.adapter)
    with app.state.sessions.begin() as db:
        connection=db.scalar(select(Connection).where(Connection.code=='middle'))
        connection.config=hub.config;connection.mappings=typed('online_course',CX_MAP,{'staff_no','course_code','course_name'});connection.state='CONNECTED'
        schedule=db.scalar(select(IntegrationSchedule).where(IntegrationSchedule.data_type=='online_course'))
        if not schedule:
            schedule=IntegrationSchedule(connection_id=connection.id,data_type='online_course',frequency_minutes=10080,auto_commit=True,active=True,approved_by='schooladmin')
            db.add(schedule);db.flush()
        ident=create_pull_task(db,actor(db),connection,'online_course','FULL','合成全量',schedule=schedule).id
    for _ in range(50):
        with app.state.sessions.begin() as db:
            task=db.get(Task,ident);process_pull_task_once(db,task,registry)
            if task.state in {'DONE','FAILED','BLOCKED'}:return task.state,task.error,task.result
    raise AssertionError('run did not finish')

def course(login,cid,name='测试教师一'):
    return {'creater_login_name':login,'teachername':name,'courseid':cid,'coursename':'线上课'+cid,'xn':'2025','xq':'1','credit':'2'}

def cx_rows(app):
    with Session(app.state.engine) as db:
        return {(x.staff_no,x.course_code):(x.active,x.user_id) for x in db.scalars(select(TeacherCourseTerm).where(TeacherCourseTerm.origin=='CX'))}

def test_teacher_who_fixes_login_is_matched_on_next_full_sync_and_listed_until_then(client2,app2):
    c=client2;auth(c,'schooladmin')
    with Session(app2.state.engine) as db:
        db.add(TeacherHrProfile(staff_no='J001',user_id='teacher',name='测试教师一',org_code='dept1'));db.commit()
    r=post(c,'/api/settings',{'key':'dict:cx_login_alias','value':{'name':'超星登录名对照','items':['gone01=J001','J001=J777','wrong01=J001']},**REASON})
    assert r.status_code==200,r.text
    others=[course('wp9',f'X{i}','外聘老师') for i in range(2,6)]
    state,error,_=cx_full_run(app2,[course('phone138','X1')]+[course('wrong01','X9','别的老师')]+others)
    assert state=='DONE',error
    assert cx_rows(app2)[('phone138','X1')]==(True,None)
    book=c.get('/api/reference/cx-unmatched.xlsx');assert book.status_code==200,book.text
    wb=load_workbook(BytesIO(book.content))
    unmatched={row[0]:row for row in wb['还没对上工号'].iter_rows(min_row=2,values_only=True)}
    assert set(unmatched)=={'phone138','wp9'} and unmatched['wp9'][2]==4
    checks={row[0]:row[4] for row in wb['登录名对照核对'].iter_rows(min_row=2,values_only=True)}
    assert '可以删除' in checks['gone01'] and '请删除' in checks['J001'] and '不一致' in checks['wrong01']
    # 老师在超星把登录名改成工号 J001：下一次全量同步按工号入库并对上账号，旧登录名那条自动停用。
    state,error,result=cx_full_run(app2,[course('J001','X1')]+[course('wrong01','X9','别的老师')]+others)
    assert state=='DONE',error
    got=cx_rows(app2)
    assert got[('J001','X1')]==(True,'teacher') and got[('phone138','X1')][0] is False
    assert result['reconcile']['deactivated']==1
    wb=load_workbook(BytesIO(c.get('/api/reference/cx-unmatched.xlsx').content))
    assert {row[0] for row in wb['还没对上工号'].iter_rows(min_row=2,values_only=True)}=={'wp9'}

def test_cx_unmatched_download_is_school_admin_only(client2):
    c=client2;auth(c,'teacher')
    assert c.get('/api/reference/cx-unmatched.xlsx').status_code==403
