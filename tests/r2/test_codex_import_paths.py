"""Isolated import regressions; no school endpoints or production data."""
from hashlib import sha256
from io import BytesIO
import zipfile
import pytest
from openpyxl import Workbook,load_workbook
from sqlalchemy import select
from app.adapters.datahub import ConfigurableDataHubAdapter
from app.adapters.contracts import AdapterContractError
from app.adapters.http_transport import PinnedHttpsTransport,TransportError
from app.excel_io import TEMPLATES,inspect_xlsx
from app.models import Task,User,Audit
from app.reference_data import TeacherHrProfile,TeacherCourseTerm
from conftest import auth,post
from test_r3_excel import upload_preview,REASON


@pytest.fixture
def cursor_hub(monkeypatch):
    adapter=ConfigurableDataHubAdapter()
    monkeypatch.setattr(adapter,'_request_json',lambda *,context,**kw:{
        'rows':[{'code':str(context['page'])}], 'cursor':f"cursor-{context['page']}",
        'more':context['page']<301,'total':301})
    config={'adapter':{'data_types':{'teacher_course':{
        'request':{'method':'GET','path':'courses'},
        'pagination':{'mode':'cursor','page_size':1,'max_pages':301},
        'response':{'records_path':'rows','next_cursor_path':'cursor','has_more_path':'more','total_path':'total'}}}}}
    return adapter,dict(data_type='teacher_course',mode='FULL',cursor=None,limit=1000,
                        config=config,credential='TEST_ONLY',timeout_seconds=1)


def test_cursor_continues_after_300_pages_and_sqlite_checkpoint_roundtrip(cursor_hub,app2):
    adapter,kwargs=cursor_hub;checkpoint=None
    for page in range(1,301):
        result=adapter.pull_chunk(checkpoint=checkpoint,**kwargs)
        assert result['page']==page and result['records']==[{'code':str(page)}] and result['has_more']
        checkpoint=result['checkpoint']
    assert len(checkpoint['seen'])==300 and all(len(x)==64 for x in checkpoint['seen'])
    assert app2.state.engine.dialect.name in {'sqlite','mysql'}
    with app2.state.sessions.begin() as db:
        task=Task(kind='integration_pull',owner_id='schooladmin',key='import-paths-checkpoint',payload={},result={'checkpoint':checkpoint})
        db.add(task);db.flush();ident=task.id
    with app2.state.sessions() as db:
        saved=db.get(Task,ident).result['checkpoint']
        assert saved==checkpoint
    last=adapter.pull_chunk(checkpoint=saved,**kwargs)
    assert last['page']==301 and not last['has_more'] and last['checkpoint'] is None


@pytest.mark.parametrize('next_cursor',['cursor-7','cursor-160'])
def test_cursor_loop_still_rejected_after_160_pages(cursor_hub,monkeypatch,next_cursor):
    adapter,kwargs=cursor_hub
    checkpoint={'page':161,'cursor':'cursor-160','seen':[sha256(f'cursor-{n}'.encode()).hexdigest() for n in range(1,161)]}
    monkeypatch.setattr(adapter,'_request_json',lambda **kw:{'rows':[{}],'cursor':next_cursor,'more':True,'total':301})
    with pytest.raises(AdapterContractError,match='PULL_CURSOR_NOT_ADVANCING'):
        adapter.pull_chunk(checkpoint=checkpoint,**kwargs)


@pytest.mark.parametrize('page,count,error',[(302,301,'PULL_PAGE_LIMIT_EXCEEDED'),(161,302,'PULL_PAGING_INVALID')])
def test_cursor_limits_still_enforced(cursor_hub,page,count,error):
    adapter,kwargs=cursor_hub
    checkpoint={'page':page,'cursor':'current','seen':[sha256(str(n).encode()).hexdigest() for n in range(count)]}
    with pytest.raises(AdapterContractError,match=error):adapter.pull_chunk(checkpoint=checkpoint,**kwargs)


@pytest.mark.parametrize('method,args,kwargs',[
    ('request_json',('GET','response'),{}),('request_form_json',('response',),{'form':{'key':'fixture'}}),
    ('request_text',('GET','response'),{})])
@pytest.mark.parametrize('encoding',['utf-8','utf-8-sig'])
def test_transport_accepts_json_and_text_with_or_without_bom(monkeypatch,method,args,kwargs,encoding):
    transport=PinnedHttpsTransport('https://school.invalid');body='{"name":"教师"}'
    monkeypatch.setattr(transport,'request_bytes',lambda *a,**kw:(200,'application/json',body.encode(encoding)))
    assert getattr(transport,method)(*args,**kwargs)==(body if method=='request_text' else {'name':'教师'})


@pytest.mark.parametrize('method,args,kwargs,error',[
    ('request_json',('GET','response'),{},'UPSTREAM_JSON_INVALID'),
    ('request_form_json',('response',),{'form':{}},'UPSTREAM_JSON_INVALID'),
    ('request_text',('GET','response'),{},'UPSTREAM_TEXT_INVALID')])
def test_transport_rejects_invalid_utf8_after_bom(monkeypatch,method,args,kwargs,error):
    transport=PinnedHttpsTransport('https://school.invalid')
    monkeypatch.setattr(transport,'request_bytes',lambda *a,**kw:(200,'application/json',b'\xef\xbb\xbf{"name":"\xff"}'))
    with pytest.raises(TransportError,match=error):getattr(transport,method)(*args,**kwargs)


def staff_workbook(kind,values,number_format='General'):
    wb=Workbook();ws=wb.active;ws.append([x[0] for x in TEMPLATES[kind]])
    defaults={'username':'import_teacher','name':'导入教师','org_code':'dept1','title':'讲师',
        'external_key':'import-paths-row','event_date':'2026-09-10','project_code':123,'project':'项目',
        'course_code':'C1','course':'课程','class_code':'CL1','class_name':'班级','hours':2.5,
        'participants':12,'income':321.5,'source':'Excel导入','category':'职业资格',
        'certificate_no':123,'issuer':'机构','obtained':'2026-01-01','expires':''}
    for value in values:
        ws.append([value if field in {'staff_no','teacher_staff_no'} else defaults[field] for _,field in TEMPLATES[kind]])
        for i,(_,field) in enumerate(TEMPLATES[kind],1):
            if field in {'staff_no','teacher_staff_no','project_code','certificate_no'}:ws.cell(ws.max_row,i).number_format=number_format
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()


@pytest.mark.parametrize('kind',list(TEMPLATES))
@pytest.mark.parametrize('value,number_format,expected',[
    (123,'00000','00123'),(123.0,'General','123'),(123.0,'00000','00123'),
    (0,'00000','00000'),(123456,'00000','123456'),('00123','General','00123'),(123,'0.00','123')])
def test_staff_number_preserves_zero_padding_and_integer_float(kind,value,number_format,expected):
    data=staff_workbook(kind,[value],number_format)
    if isinstance(value,float):
        # openpyxl writes 123.0 as 123; preserve a genuine float in the input XML.
        out=BytesIO()
        with zipfile.ZipFile(BytesIO(data)) as src,zipfile.ZipFile(out,'w') as dest:
            for entry in src.infolist():
                content=src.read(entry.filename)
                if entry.filename=='xl/worksheets/sheet1.xml':content=content.replace(b'<v>123</v>',b'<v>123.0</v>')
                dest.writestr(entry,content)
        data=out.getvalue()
        wb=load_workbook(BytesIO(data),read_only=True)
        column=next(i for i,(_,field) in enumerate(TEMPLATES[kind],1) if field in {'staff_no','teacher_staff_no'})
        assert type(wb.active.cell(2,column).value) is float;wb.close()
    row=inspect_xlsx(data,'staff.xlsx',kind)[0]
    field='staff_no' if kind=='teachers' else 'teacher_staff_no'
    assert not row['errors'] and row['raw']['教师工号']==expected and row['normalized'][field]==expected
    if kind in {'records','participations'}:
        assert row['raw']['项目编码']==123 and row['normalized']['project_code'] in {'123','123.0'}
        assert row['normalized']['hours']=='2.5'
    if kind=='qualifications':assert row['normalized']['certificate_no'] in {'123','123.0'}


@pytest.mark.parametrize('kind',list(TEMPLATES))
@pytest.mark.parametrize('number_format',['General','00000'])
def test_fractional_staff_number_is_a_row_error(kind,number_format):
    rows=inspect_xlsx(staff_workbook(kind,[123.5,124],number_format),'staff.xlsx',kind)
    assert rows[0]['row_no']==2 and rows[0]['errors']==['工号必须是文本']
    assert rows[1]['row_no']==3 and not rows[1]['errors']


def test_teacher_excel_links_existing_hr_and_all_course_origins(client2,app2):
    with app2.state.sessions.begin() as db:
        db.add(TeacherHrProfile(id='hr-import',staff_no='00123',name='导入教师',org_code='dept1'))
        db.add(TeacherHrProfile(id='hr-unrelated',staff_no='OTHER',name='其他教师',org_code='dept1'))
        for ident,origin,staff,user in [('jw','JW','00123',None),('cx','CX','00123',None),
                                      ('linked','JW','00123','teacher'),('unrelated','JW','OTHER',None)]:
            db.add(TeacherCourseTerm(id=ident,external_key=ident,origin=origin,staff_no=staff,user_id=user,
                course_code=ident,course_name='课程',term_name='2025-2026-2'))
    auth(client2,'schooladmin')
    response=upload_preview(client2,'teachers',staff_workbook('teachers',[123],'00000'))
    assert response.status_code==200,response.text
    job=response.json();assert job['errors']==0
    with app2.state.sessions() as db:
        assert db.scalar(select(User).where(User.username=='import_teacher')) is None
        assert db.get(TeacherHrProfile,'hr-import').user_id is None
    response=post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON})
    assert response.status_code==200,response.text
    assert response.json()['state']=='COMMITTED' and response.json()['receipt']['written']==1
    with app2.state.sessions() as db:
        user=db.scalar(select(User).where(User.username=='import_teacher'))
        assert user.staff_no=='00123' and not user.active and user.must_change_password
        profile=db.get(TeacherHrProfile,'hr-import');assert profile.user_id==user.id and profile.version==2
        for ident in ('jw','cx'):
            term=db.get(TeacherCourseTerm,ident);assert term.user_id==user.id and term.version==2
        assert db.get(TeacherHrProfile,'hr-unrelated').user_id is None
        assert db.get(TeacherCourseTerm,'unrelated').user_id is None
        assert db.get(TeacherCourseTerm,'linked').user_id=='teacher' and db.get(TeacherCourseTerm,'linked').version==1
        audits=db.scalars(select(Audit).where(Audit.module.in_(['teacher_hr','teacher_course']))).all()
        assert {a.action for a in audits}=={'新账号对应人事记录','新账号对应授课记录'}


def test_fractional_staff_number_cannot_create_account(client2,app2):
    auth(client2,'schooladmin')
    response=upload_preview(client2,'teachers',staff_workbook('teachers',[123.5],'00000'))
    assert response.status_code==200,response.text
    job=response.json();assert job['errors']==1
    response=post(client2,f"/api/imports/{job['id']}/confirm",{'expected_version':job['version'],**REASON})
    assert response.status_code==409,response.text
    with app2.state.sessions() as db:assert db.scalar(select(User).where(User.username=='import_teacher')) is None
