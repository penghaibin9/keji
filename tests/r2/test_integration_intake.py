"""School-facing XLSX intake: synthetic data only, never contacts external systems."""
from io import BytesIO
import json
import uuid
from urllib.parse import quote

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Connection, Setting
from app.integration_intake import build_intake_workbook, parse_intake_workbook, overlay_intake
from conftest import auth

MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def template_bytes(client):
    response=client.get('/api/connections/intake-template.xlsx')
    assert response.status_code==200,response.text
    assert response.content.startswith(b'PK')
    return response.content


def save_workbook(workbook):
    stream=BytesIO();workbook.save(stream);workbook.close();return stream.getvalue()


def prepare_school_workbook(data):
    workbook=load_workbook(BytesIO(data))
    basic=workbook['接口基础']
    for row in range(2,basic.max_row+1):
        code=basic.cell(row,1).value
        basic.cell(row,3).value='https://'+code+'.synthetic.invalid'
        basic.cell(row,4).value='合成学校信息中心 / 实施联系人'
        basic.cell(row,5).value='cas' if code=='sso' else ('token' if code=='middle' else 'signature')
        basic.cell(row,9).value='SYNTHETIC-SCHOOL'

    sso=workbook['统一认证参数']
    values={'protocol':'cas','login_path':'/login','validate_path':'/serviceValidate',
            'probe.kind':'cas_ticket','active_true_values':'ACTIVE,1,true',
            'active_false_values':'DISABLED,0,false'}
    for row in range(2,sso.max_row+1):
        key=sso.cell(row,2).value
        if key in values:sso.cell(row,3).value=values[key]

    identity=workbook['统一认证身份']
    identity_values={'stable_id':'id','account':'account','staff_no':'staff_no',
                     'name':'name','org_code':'org','roles':'roles','active':'status'}
    for row in range(2,identity.max_row+1):
        identity.cell(row,3).value=identity_values[identity.cell(row,1).value]

    materials=workbook['资料登记']
    for row in range(2,materials.max_row+1):
        materials.cell(row,4).value='合成资料责任人'
        materials.cell(row,5).value='SYNTHETIC-DOC-'+str(row)
        materials.cell(row,6).value='2026-09-23'

    credential=workbook['凭据证书元数据']
    for row in range(2,credential.max_row+1):
        credential.cell(row,3).value='token'
        credential.cell(row,4).value='server_env'
        credential.cell(row,5).value='合成服务器管理员'

    roles=workbook['统一认证角色确认']
    for row in range(2,roles.max_row+1):
        kind=roles.cell(row,1).value
        roles.cell(row,3).value='external_'+str(kind)
        roles.cell(row,6).value='合成学校信息中心'
        roles.cell(row,7).value='2026-09-23'

    sources=workbook['权威来源']
    columns={cell.value:index+1 for index,cell in enumerate(sources[1])}
    for row in range(2,sources.max_row+1):
        sources.cell(row,columns['权威来源系统']).value='合成权威业务系统'
        sources.cell(row,columns['数据提供方式']).value='REST API'
        sources.cell(row,columns['同步/调用频率']).value='每30分钟或按正式规范'
        sources.cell(row,columns['历史数据范围']).value='2025-01-01至上线日'
        sources.cell(row,columns['责任人']).value='合成数据责任人'
        sources.cell(row,columns['责任边界说明']).value='对方保证源数据完整；本平台负责映射校验和处理'
        sources.cell(row,columns['确认依据/文档编号']).value='SOURCE-CONFIRM-'+str(row)
        sources.cell(row,columns['双方确认日期']).value='2026-09-23'
    return save_workbook(workbook)


def upload_headers(filename='school-intake.xlsx', **extra):
    return {'Content-Type':MIME, 'X-Intake-Filename':quote(filename), **extra}


def preview(client,data):
    return client.post('/api/connections/intake-preview', content=data,
        headers=upload_headers())


def apply(client,data,reason='导入学校正式联调资料'):
    return client.post('/api/connections/intake-apply', content=data,
        headers=upload_headers(**{'X-Intake-Reason':quote(reason),
            'X-Intake-Basis':quote('合成专项测试'),
            'Idempotency-Key':str(uuid.uuid4())}))


def connection_snapshot(app):
    with app.state.sessions() as db:
        rows=db.scalars(select(Connection).order_by(Connection.code)).all()
        return {x.code:{'version':x.version,'state':x.state,'config':json.loads(json.dumps(x.config))}
                for x in rows}

def test_school_intake_template_is_structured_and_secret_free(client2):
    auth(client2,'schooladmin')
    data=template_bytes(client2)
    workbook=load_workbook(BytesIO(data),read_only=False)
    assert workbook.sheetnames==[
        '填写说明','接口基础','统一认证参数','统一认证身份','统一认证组织映射','统一认证角色确认',
        '数据中台接口','湘培接口','字段映射','资料登记','凭据证书元数据','权威来源','_系统']
    assert workbook['_系统'].sheet_state=='hidden'
    assert workbook['_系统']['B1'].value=='YUEKE_INTEGRATION_INTAKE_V1'
    basic=workbook['接口基础']
    refs={basic.cell(r,1).value:basic.cell(r,6).value for r in range(2,basic.max_row+1)}
    assert refs=={'sso':'ND_SSO_SECRET','middle':'ND_MIDDLE_SECRET','province':'ND_PROVINCE_SECRET'}
    headers=[cell.value for cell in workbook['凭据证书元数据'][1]]
    assert '私钥' not in headers and '密码' not in headers and '令牌值' not in headers
    workbook.close()


def test_preview_is_read_only_and_reports_three_connectors(client2,app2):
    auth(client2,'schooladmin')
    data=prepare_school_workbook(template_bytes(client2))
    before=connection_snapshot(app2)
    response=preview(client2,data)
    assert response.status_code==200,response.text
    body=response.json()
    assert [x['code'] for x in body['connectors']]==['sso','middle','province']
    assert body['effects']['active_connections_changed'] is False
    assert body['effects']['external_called'] is False
    assert body['effects']['drafts_changed'] is False
    assert connection_snapshot(app2)==before

    with app2.state.sessions() as db:
        assert not db.scalar(select(Setting).where(Setting.key.like('integration.draft:%')))


def test_apply_stages_drafts_materials_credentials_and_sources_only(client2,app2):
    auth(client2,'schooladmin')
    data=prepare_school_workbook(template_bytes(client2))
    before=connection_snapshot(app2)
    response=apply(client2,data)
    assert response.status_code==200,response.text
    body=response.json()
    assert body['applied'] is True and body['external_called'] is False
    assert body['published'] is False and body['school_acceptance'] is False
    assert connection_snapshot(app2)==before
    with app2.state.sessions() as db:
        connections={x.code:x for x in db.scalars(select(Connection)).all()}
        for code,obj in connections.items():
            draft=db.get(Setting,'integration.draft:'+obj.id)
            materials=db.get(Setting,'integration.materials:'+obj.id)
            credential=db.get(Setting,'integration.credential:'+obj.id)
            assert draft and draft.value['base_version']==obj.version
            assert materials and len(materials.value['items'])==6
            assert credential and credential.value['credential_reference']=={
                'sso':'ND_SSO_SECRET','middle':'ND_MIDDLE_SECRET','province':'ND_PROVINCE_SECRET'}[code]
        assert db.get(Setting,'integration.sources:'+connections['middle'].id)
        assert db.get(Setting,'integration.sources:'+connections['province'].id)
        policy=db.get(Setting,'integration.identity_policy:'+connections['sso'].id)
        assert policy and policy.value['authorization_mode']=='LOCAL_ROLE_AUTHORITATIVE'
        assert {x['kind'] for x in policy.value['items']}=={'teacher','dept','college','school'}

    catalog=client2.get('/api/connections/catalog').json()['items']
    province=next(x for x in catalog if x['connection']['code']=='province')
    plan=next(x for x in province['data_types'] if x['data_type']=='plan')
    assert plan['source_confirmed'] is True
    assert plan['source_system']=='合成权威业务系统'
    workspace=client2.get('/api/connections/'+province['connection']['id']+'/implementation').json()
    assert workspace['source_gaps']==[]
    assert any('培训计划请求路径' in issue for issue in workspace['issues'])
    assert workspace['intake']['record']['file_sha256']==body['file_sha256']

    stale=apply(client2,data,'重复导入旧学校资料表')
    assert stale.status_code==409
    assert '重新下载学校资料表' in stale.text


def test_intake_rejects_formula_and_credential_reference_tampering(client2):
    auth(client2,'schooladmin')
    raw=template_bytes(client2)
    workbook=load_workbook(BytesIO(raw))
    workbook['接口基础']['C2']='=1+1'
    formula=save_workbook(workbook)
    response=preview(client2,formula)
    assert response.status_code==422
    assert '不允许公式' in response.text

    workbook=load_workbook(BytesIO(raw))
    workbook['接口基础']['F2']='REAL_SECRET_VALUE'
    tampered=save_workbook(workbook)
    response=preview(client2,tampered)
    assert response.status_code==422
    assert '凭据引用不可由Excel修改' in response.text

@pytest.mark.parametrize('username',['teacher','depthead','collegehead'])
def test_non_admin_cannot_use_school_intake(client2,username):
    auth(client2,username)
    assert client2.get('/api/connections/intake-template.xlsx').status_code==403
    response=client2.post('/api/connections/intake-preview',content=b'PK-not-real',
        headers=upload_headers())
    assert response.status_code==403


def test_intake_stale_material_version_refuses_overwrite(client2,app2):
    auth(client2,'schooladmin')
    raw=prepare_school_workbook(template_bytes(client2))
    with app2.state.sessions.begin() as db:
        sso=db.scalar(select(Connection).where(Connection.code=='sso'))
        db.add(Setting(key='integration.materials:'+sso.id,
            value={'items':[{'name':'正式接口规范','owner':'另一实施人员',
                             'reference':'NEWER-DOC','confirmed_on':'2026-09-23'}]},version=1))
    response=apply(client2,raw,'尝试导入已过期资料表')
    assert response.status_code==409
    assert '重新下载学校资料表' in response.text
    with app2.state.sessions() as db:
        assert not db.scalar(select(Setting).where(Setting.key.like('integration.draft:%')))

def test_intake_invalidates_prior_school_signoff(client2,app2):
    auth(client2,'schooladmin')
    with app2.state.sessions.begin() as db:
        sso=db.scalar(select(Connection).where(Connection.code=='sso'))
        db.add(Setting(key='integration.acceptance:'+sso.id,value={
            'signed_by':'学校验收人','signed_on':'2026-09-23','signed_reference':'OLD-SIGNOFF',
            'items':[],'connection_version':sso.version,'configuration_hash':'synthetic'},version=1))
    raw=prepare_school_workbook(template_bytes(client2))
    response=apply(client2,raw,'学校提供了新一版接口资料')
    assert response.status_code==200,response.text
    with app2.state.sessions() as db:
        sso=db.scalar(select(Connection).where(Connection.code=='sso'))
        acceptance=db.get(Setting,'integration.acceptance:'+sso.id)
        assert acceptance.value['invalidated_at']
        assert acceptance.value['invalidation_reason']=='导入新的学校联调资料'
        sso_id=sso.id
    workspace=client2.get('/api/connections/'+sso_id+'/implementation').json()
    assert workspace['statuses']['acceptance'] in {
        '配置已变化，验收需重新确认','实施资料已变化，验收需重新确认'}


def test_middle_paging_parameter_location_roundtrips_through_excel(app2):
    with app2.state.sessions.begin() as db:
        middle=db.scalar(select(Connection).where(Connection.code=='middle'))
        middle.config={'endpoint':'https://middle.synthetic.invalid','owner':'合成中台',
            'secret_env':'ND_MIDDLE_SECRET','retry_count':3,'frequency_minutes':30,'auth_mode':'token',
            'adapter':{'data_types':{'training_project':{
                'request':{'method':'GET','path':'/projects','headers':{
                    'X-Page':'$'+'{page}','X-Limit':'$'+'{limit}'},'auth':{'mode':'none'}},
                'pagination':{'mode':'page','page_size':200,'max_pages':20,'start_page':1},
                'response':{'records_path':'items','has_more_path':'more'}}}}}
        db.flush()
        raw=build_intake_workbook(db)
        parsed=parse_intake_workbook(raw)
        item,_=__import__('app.integration_center',fromlist=['candidate']).candidate(db,middle)
        staged=overlay_intake(item,parsed,'middle')

        request=staged['config']['adapter']['data_types']['training_project']['request']
        assert request['headers']['X-Page']=='$'+'{page}'
        assert request['headers']['X-Limit']=='$'+'{limit}'
        assert '$'+'{page}' not in request.get('query',{}).values()
        assert '$'+'{limit}' not in request.get('query',{}).values()
