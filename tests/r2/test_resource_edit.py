"""资源补正走正常权限、版本与审计；只用隔离合成数据。"""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select,func
from conftest import auth,post,put
from test_social_training import setup_social,project_candidates
from app.social_models import SocialResource,SocialPolicy,SocialProject,SocialClaim
from app.models import Org,User,WorkflowEvent,Setting,PointEntry
from app.reference_data import TrainingCourseMaster,TrainingProjectMaster
from app.core import serial,digest
from app import social_training as st


def test_resource_edit_scope_version_identity_and_history(setup_social,client2,app2):
    body={'kind':'achievement','org_id':'college1','code':'EDIT-A1','name':'隔离成果','recognition':'隔离认定依据','reason':'隔离登记'}
    auth(client2,'collegehead');r=post(client2,'/api/social/resources',body);assert r.status_code==200,r.text
    x=r.json()['item'];assert r.json()['state_label']=='待补认定日期'
    edit={**body,'expected_version':x['version'],'confirmed_on':'2026-10-11','name':'隔离补正成果','reason':'补齐日期'}
    auth(client2,'otherdept');assert put(client2,'/api/social/resources/'+x['id'],edit).status_code in (403,404)
    auth(client2,'depthead');assert put(client2,'/api/social/resources/'+x['id'],edit).status_code==403
    auth(client2,'teacher');assert put(client2,'/api/social/resources/'+x['id'],edit).status_code==403
    auth(client2,'collegehead');assert put(client2,'/api/social/resources/'+x['id'],{**edit,'org_id':'college2'}).status_code==422
    r=put(client2,'/api/social/resources/'+x['id'],edit);assert r.status_code==200,r.text
    assert r.json()['item']['data']['confirmed_on']=='2026-10-11' and r.json()['state_label']=='已登记'
    assert r.json()['events'][-1]['action']=='补正培训资源资料'
    assert put(client2,'/api/social/resources/'+x['id'],edit).status_code==409
    with app2.state.sessions.begin() as db:
        obj=db.get(SocialResource,x['id']);obj.data={**obj.data,'auto':True}
    edit['expected_version']=r.json()['item']['version']
    assert put(client2,'/api/social/resources/'+x['id'],edit).status_code==409


def test_course_resource_edit_preserves_import_metadata_and_rechecks_course(setup_social,client2,app2):
    auth(client2,'collegehead')
    body={'kind':'course_resource','org_id':'college1','code':'EDIT-C1','name':'隔离视频','recognition':'隔离课程依据','course_code':'SOC-C1','resource_type':'video','duration_minutes':'10','teacher_id':'teacher','reason':'登记课程资源'}
    r=post(client2,'/api/social/resources',body);assert r.status_code==200,r.text
    x=r.json()['item']
    with app2.state.sessions.begin() as db:
        obj=db.get(SocialResource,x['id']);obj.data={**obj.data,'source':'chaoxing','video_count':12}
    edit={**body,'expected_version':x['version'],'duration_minutes':'12.5'}
    assert put(client2,'/api/social/resources/'+x['id'],{**edit,'course_code':'not-a-course'}).status_code==422
    r=put(client2,'/api/social/resources/'+x['id'],edit);assert r.status_code==200,r.text
    assert r.json()['item']['data']['video_count']==12 and r.json()['item']['data']['source']=='chaoxing'
    with app2.state.sessions.begin() as db:db.get(User,'teacher').active=False
    inactive={**body,'code':'INACTIVE-TEACHER'}
    refused=post(client2,'/api/social/resources',inactive)
    assert refused.status_code==422 and '课程教师与学院或姓名不匹配' in refused.text
    unassigned=post(client2,'/api/social/resources',{**inactive,'teacher_id':None})
    assert unassigned.status_code==200,unassigned.text
    assert unassigned.json()['item']['data']['teacher_id'] is None


def test_resource_course_candidates_use_full_college_project_scope_and_keep_old_reference_queries(project_candidates,client2,app2):
    with app2.state.sessions.begin() as db:
        for index in range(207):
            db.add(TrainingCourseMaster(code=f'OPTION-C-{index:03d}',name=f'隔离候选课程{index}',project_code=f'OPT-{index:03d}'))
        db.add(TrainingCourseMaster(code='NO-PROJECT',name='未关联项目课程'))
        db.add(TrainingCourseMaster(code='INACTIVE',name='已停用课程',project_code='OPT-000',active=False))
        db.add(Org(id='inactive-college',code='inactive-college',name='停用学院',kind='college',parent_id='school',active=False))
    params={'resource_org_id':'college1','active':'true','size':200}
    assert client2.get('/api/reference/courses',params=params).status_code==401
    auth(client2,'collegehead')
    pages=[client2.get('/api/reference/courses',params={**params,'page':page}) for page in (1,2)]
    assert all(page.status_code==200 for page in pages)
    first,second=[page.json() for page in pages]
    assert first['total']==second['total']==206 and [len(first['items']),len(second['items'])]==[200,6]
    assert len({row['code'] for page in (first,second) for row in page['items']})==206
    assert second['items'][-1]['code']=='OPTION-C-205'
    assert client2.get('/api/reference/courses',params={**params,'page':3}).json()['items']==[]
    found=client2.get('/api/reference/courses',params={**params,'q':'OPTION-C-205'}).json()
    assert found['total']==1 and found['items'][0]['project_code']=='OPT-205'
    assert client2.get('/api/reference/courses',params={**params,'q':'%'}).json()['total']==0
    for org in ('college2','dept1','school','inactive-college','missing'):
        assert client2.get('/api/reference/courses',params={**params,'resource_org_id':org}).status_code==403
    # 普通主数据检索仍按原项目筛选，不套用课程资源用途的学院限制。
    legacy=client2.get('/api/reference/courses',params={'project_code':'OPT-206','active':'true'}).json()
    assert [row['code'] for row in legacy['items']]==['OPTION-C-206']
    auth(client2,'schooladmin')
    other=client2.get('/api/reference/courses',params={**params,'resource_org_id':'college2'}).json()
    assert other['total']==1 and other['items'][0]['code']=='OPTION-C-206'
    auth(client2,'depthead')
    assert client2.get('/api/reference/courses',params=params).status_code==403


def test_school_public_reports_stay_in_recognition_and_obey_pagination(app2,client2):
    with app2.state.sessions.begin() as db:
        db.add(SocialResource(kind='base',org_id='college1',code='BASE-SCOPE',name='隔离基地',data={'onsite':True}))
        db.add(SocialResource(kind='point_report',org_id='school',code='REPORT-SCOPE',name='全校积分公示',state='ARCHIVED',data={'rows':[]}))
    auth(client2,'teacher')
    for section in ['bases','courses','outcomes']:
        r=client2.get('/api/social/workspace',params={'section':section});assert r.status_code==200,r.text
        assert all(x['kind']!='point_report' for x in r.json()['resources'])
    r=client2.get('/api/social/workspace',params={'section':'recognition'}).json()
    assert [x['kind'] for x in r['resources']]==['point_report'] and r['resource_pagination']['total']==1
    r=client2.get('/api/social/workspace',params={'section':'recognition','resource_page':2}).json()
    assert r['resources']==[] and r['resource_pagination']['total']==1


@pytest.mark.parametrize('state',['PUBLIC','ARCHIVED'])
def test_point_report_public_feedback_is_personal_without_hiding_published_rows(app2,client2,state):
    rows=[{'teacher_name':'测试教师一','org_name':'测试学院一','project':'公开项目一','teaching':'1.00','auxiliary':'0.00','total':'1.00'},
          {'teacher_name':'测试教师二','org_name':'测试学院二','project':'公开项目二','teaching':'2.00','auxiliary':'0.00','total':'2.00'}]
    disputes=[{'actor':'teacher','reason':'TEACHER-ONE-PRIVATE-FEEDBACK','resolved':True,'resolution':'TEACHER-ONE-PRIVATE-RESOLUTION','resolved_by':'reviewer'},
              {'actor':'otherteacher','reason':'TEACHER-TWO-PRIVATE-FEEDBACK','resolved':False}]
    with app2.state.sessions.begin() as db:
        admin=db.get(User,'schooladmin')
        for account in ('reviewer','secretary'):
            db.add(User(id=account,username=account,name='正常学校办理人'+account,org_id='school',role_id='role-school',password_hash=admin.password_hash,must_change_password=False))
        db.add(Setting(key='social_training_duties',value={'coordinator':['reviewer'],'secretary':['secretary']}));db.flush()
        obj=SocialResource(kind='point_report',org_id='school',code='REPORT-PRIVATE',name='全校年度积分公示',state=state,
            data={'year':2025,'start_month':1,'end_month':12,'created_by':'reviewer','rows':rows,'disputes':disputes,'public_start':'2026-10-18','public_end':'2026-10-23'})
        db.add(obj);db.flush();ident=obj.id
        event=WorkflowEvent(object_type='social_resource',object_id=ident,actor_id='reviewer',actor_name='授权归口人员',org_id='school',action='记录公示异议',
            before={'data':{'disputes':disputes[:1]}},after={'data':{'disputes':disputes}},reason='INTERNAL-REPORT-EVENT-REASON',basis='INTERNAL-REPORT-EVENT-BASIS')
        db.add(event);db.flush();event_id=event.id
        original=deepcopy(serial(obj));original_event=deepcopy(serial(event))
    for index,account in enumerate(('teacher','otherteacher')):
        auth(client2,account)
        response=client2.get('/api/social/objects/social_resource/'+ident);assert response.status_code==200,response.text
        detail=response.json();own={k:v for k,v in disputes[index].items() if k in {'reason','resolved','resolution'}}
        assert detail['item']['data']['disputes']==[own]
        assert detail['item']['data']['rows']==rows and detail['events']==[]
        assert disputes[1-index]['reason'] not in response.text and 'INTERNAL-REPORT-EVENT' not in response.text
        workspace=client2.get('/api/social/workspace?section=recognition');assert workspace.status_code==200,workspace.text
        shown=next(x for x in workspace.json()['resources'] if x['id']==ident)
        assert shown['data']['disputes']==[own] and shown['data']['rows']==rows
        assert disputes[1-index]['reason'] not in workspace.text
        exported=client2.get('/api/social/objects/social_resource/'+ident+'/export');assert exported.status_code==200,exported.text
        book=load_workbook(BytesIO(exported.content),read_only=True)
        values=[str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row]
        assert not any(disputes[1-index]['reason'] in value or 'INTERNAL-REPORT-EVENT' in value for value in values)
        assert list(book['已认定积分汇总'].values)[1:]==[tuple(float(Decimal(row[k])) if k in ('teaching','auxiliary','total') else row[k] for k in ('teacher_name','org_name','project','teaching','auxiliary','total')) for row in rows]
        book.close()
    for account in ('reviewer','secretary','schooladmin'):
        auth(client2,account)
        managed=client2.get('/api/social/objects/social_resource/'+ident).json()
        assert managed['item']['data']['disputes']==disputes and managed['item']['data']['rows']==rows
        assert managed['events']==[original_event]
        shown=next(x for x in client2.get('/api/social/workspace?section=recognition').json()['resources'] if x['id']==ident)
        assert shown['data']['disputes']==disputes
    with app2.state.sessions() as db:
        assert serial(db.get(SocialResource,ident))==original
        assert serial(db.get(WorkflowEvent,event_id))==original_event


@pytest.fixture
def refreshable_report(app2,client2,monkeypatch):
    monkeypatch.setattr(st,'today',lambda:date(2026,10,18))
    with app2.state.sessions.begin() as db:
        admin=db.get(User,'schooladmin')
        for account in ('refresh-secretary','other-secretary','refresh-reviewer'):
            db.add(User(id=account,username=account,name='合成汇总办理人',org_id='school',role_id='role-school',password_hash=admin.password_hash,must_change_password=False))
        db.add(Setting(key='social_training_duties',value={'secretary':['refresh-secretary','other-secretary'],'coordinator':['refresh-reviewer']}))
        db.add(SocialPolicy(id='refresh-policy',period_id='period2026',name='合成汇总政策',state='ACTIVE',created_by='schooladmin',data={'calendar_start':'2026-01-01','calendar_end':'2027-12-31','calendar_reference':'合成校历','holidays':[],'workdays':[]}));db.flush()
        for index,teacher,org,teaching,auxiliary in ((1,'teacher','college1','0.07','0.00'),(2,'otherteacher','college2','0.07','0.50')):
            code=f'REFRESH-P{index}';db.add(TrainingProjectMaster(id=code,code=code,name='合成独立项目'+str(index)));db.flush()
            db.add(SocialProject(id=code,master_id=code,org_id=org,owner_id=teacher,cohort='一期',state='CLOSED',data={'code':code,'name':'合成独立项目'+str(index)},closeout={'closed_on':'2026-10-16'}));db.flush()
            trial={'teaching':teaching,'auxiliary':auxiliary,'basis':{'facts':{'project':{'closeout':{'closed_on':'2026-10-16'}}}}}
            db.add(SocialClaim(id='refresh-claim'+str(index),project_id=code,teacher_id=teacher,period_id='period2026',policy_id='refresh-policy',state='ARCHIVED' if index==1 else 'DRAFT',data={},snapshot={'trial':trial,'sha256':digest(trial)}))
        db.add(PointEntry(teacher_id='teacher',period_id='period2026',item_key='refresh-fix',kind='SOCIAL',amount='0.01',raw_value='0.01',formula='合成追加更正',source_snapshot={'social_claim':'refresh-claim1','correction_id':'refresh-fix','category':'teaching'},actor_id='schooladmin',reason='合成追加更正'))
    auth(client2,'refresh-secretary')
    response=post(client2,'/api/social/reports',{'org_id':'school','policy_id':'refresh-policy','year':2026,'start_month':9,'end_month':10,'reason':'合成原双月汇总'})
    assert response.status_code==200,response.text
    return response.json()['item']


def test_report_refresh_keeps_original_id_snapshot_and_requires_resubmission(refreshable_report,app2,client2):
    original=refreshable_report;ident=original['id'];path='/api/social/reports/'+ident+'/actions'
    assert original['data']['rows'][0]['total']=='0.08'
    submitted=post(client2,path,{'action':'submit','expected_version':original['version'],'reason':'明确提交原双月汇总'})
    assert submitted.status_code==200,submitted.text
    before=submitted.json()['item']
    with app2.state.sessions.begin() as db:db.get(SocialClaim,'refresh-claim2').state='ARCHIVED'
    result=post(client2,path,{'action':'refresh','expected_version':before['version'],'reason':'第二位教师归档后重新取数','basis':'合成新增0.57，原0.08留在历史'})
    assert result.status_code==200,result.text
    after=result.json()['item'];assert after['id']==ident and after['state']=='DRAFT' and after['version']==before['version']+1
    assert (after['org_id'],after['data']['year'],after['data']['start_month'],after['data']['end_month'],after['data']['created_by'],after['data']['policy_id'])==('school',2026,9,10,'refresh-secretary','refresh-policy')
    assert sum(Decimal(row['total']) for row in after['data']['rows'])==Decimal('0.65')
    assert after['data']['sha256']==digest(after['data']['rows'])!=before['data']['sha256']
    with app2.state.sessions() as db:
        assert db.scalar(select(func.count()).select_from(SocialResource).where(SocialResource.kind=='point_report'))==1
        change=db.scalar(select(WorkflowEvent).where(WorkflowEvent.object_id==ident,WorkflowEvent.action=='积分汇总：重新取数并保存草稿'))
        assert change.before['data']['rows']==before['data']['rows'] and change.after['data']['rows']==after['data']['rows']
        assert change.before['state']=='SUBMITTED' and change.after['state']=='DRAFT' and change.basis=='合成新增0.57，原0.08留在历史'
    auth(client2,'refresh-reviewer')
    refused=post(client2,path,{'action':'approve','expected_version':after['version'],'reason':'不能替生成者提交草稿'})
    assert refused.status_code==409 and '尚未提交' in refused.text
    auth(client2,'refresh-secretary')
    draft=post(client2,path,{'action':'refresh','expected_version':after['version'],'reason':'草稿阶段再次核对最新来源'})
    assert draft.status_code==200 and draft.json()['item']['state']=='DRAFT'
    assert post(client2,path,{'action':'refresh','expected_version':after['version'],'reason':'过期窗口不能覆盖'}).status_code==409
    submitted=post(client2,path,{'action':'submit','expected_version':draft.json()['item']['version'],'reason':'重新核对后明确提交'})
    assert submitted.status_code==200 and submitted.json()['item']['state']=='SUBMITTED'
    auth(client2,'refresh-reviewer')
    waiting=post(client2,path,{'action':'approve','expected_version':submitted.json()['item']['version'],'reason':'真实期间限制继续生效'})
    assert waiting.status_code==409 and '汇总期间尚未结束' in waiting.text


def test_report_refresh_rejects_noncreator_revoked_duty_closed_states_and_invalid_policy(refreshable_report,app2,client2):
    original=refreshable_report;ident=original['id'];path='/api/social/reports/'+ident+'/actions'
    body={'action':'refresh','expected_version':original['version'],'reason':'合成原单重新取数'}
    auth(client2,'other-secretary');assert post(client2,path,body).status_code==403
    auth(client2,'refresh-secretary')
    with app2.state.sessions.begin() as db:
        setting=db.get(Setting,'social_training_duties');duties=deepcopy(setting.value);setting.value={**duties,'secretary':['other-secretary']}
    assert post(client2,path,body).status_code==403
    with app2.state.sessions.begin() as db:db.get(Setting,'social_training_duties').value=duties
    for state in ('PUBLIC','ARCHIVED','VOID'):
        with app2.state.sessions.begin() as db:db.get(SocialResource,ident).state=state
        assert post(client2,path,body).status_code==409
    with app2.state.sessions.begin() as db:
        db.get(SocialResource,ident).state='DRAFT';db.get(SocialPolicy,'refresh-policy').state='VOID'
    assert post(client2,path,body).status_code==409
    with app2.state.sessions.begin() as db:
        db.get(SocialPolicy,'refresh-policy').state='ACTIVE';obj=db.get(SocialResource,ident);obj.data={**obj.data,'sha256':'invalid'}
    assert post(client2,path,body).status_code==409
    with app2.state.sessions.begin() as db:
        obj=db.get(SocialResource,ident);obj.data=original['data']
        db.add(Setting(key='social_conflicts',value=[{'user_id':'refresh-secretary','target_kind':'teacher','target_id':'teacher','target_name':'合成教师','active':True}]))
    assert post(client2,path,body).status_code==403
    with app2.state.sessions() as db:
        obj=db.get(SocialResource,ident);assert obj.data==original['data'] and obj.version==original['version'] and obj.state=='DRAFT'


def test_base_edit_keeps_fields_not_submitted_and_requires_external_partner(setup_social,client2,app2):
    body={'kind':'base','org_id':'college1','code':'EDIT-B1','name':'隔离校内基地','recognition':'隔离认定依据','reason':'隔离登记',
          'onsite':True,'usages':['practice','ideology'],'established_on':'2026-03-01','partner':'隔离共建单位','level':'city'}
    auth(client2,'collegehead');r=post(client2,'/api/social/resources',body);assert r.status_code==200,r.text
    x=r.json()['item']
    # 旧表单只提交名称、编号、地点、依据和校内外：其余资料必须保持原值。
    edit={k:body[k] for k in ('kind','org_id','code','recognition','onsite')}|{'name':'隔离校内基地（改名）','reason':'改名','expected_version':x['version']}
    r=put(client2,'/api/social/resources/'+x['id'],edit);assert r.status_code==200,r.text
    d=r.json()['item']['data']
    assert d['usages']==['practice','ideology'] and d['established_on']=='2026-03-01' and d['partner']=='隔离共建单位' and d['level']=='city'
    # 改为校外时共建单位不能清空；带共建单位则可以保存。
    v=r.json()['item']['version']
    assert put(client2,'/api/social/resources/'+x['id'],{**edit,'onsite':False,'partner':'','expected_version':v}).status_code==422
    r=put(client2,'/api/social/resources/'+x['id'],{**edit,'onsite':False,'expected_version':v});assert r.status_code==200,r.text
    assert r.json()['item']['data']['onsite'] is False and r.json()['item']['data']['partner']=='隔离共建单位'
    assert post(client2,'/api/social/resources',{**body,'code':'EDIT-B2','level':'national'}).status_code==422
