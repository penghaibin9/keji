import io,json,uuid,zipfile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from conftest import auth,post,put
from app.models import Connection,User
from app.reference_data import (TrainingProjectMaster,TrainingCourseMaster,TrainingClassMaster,
    TeacherParticipation,TeacherHrProfile,InboundBatch)
from app.worker import run_once

REASON={'reason':'R4内部收口测试','basis':'合成隔离数据，不是学校正式样本'}

def middle(client):
    return next(x for x in client.get('/api/connections').json()['items'] if x['code']=='middle')

def configure_middle(client,mappings):
    c=middle(client)
    r=put(client,'/api/connections/'+c['id'],{'mappings':mappings,'config':{'endpoint':'https://school.example/api','owner':'学校数据中心','retry_count':3,'secret_env':'ND_MIDDLE_TOKEN','frequency_minutes':30,'auth_mode':'token'},'expected_version':c['version'],**REASON})
    assert r.status_code==200,r.text
    return r.json()

def fill_template(client,kind,values):
    r=client.get('/api/imports/templates/'+kind);assert r.status_code==200,r.text
    wb=load_workbook(io.BytesIO(r.content));ws=wb.active;ws.append(values);out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

def upload_preview(client,kind,data):
    return client.post('/api/imports/'+kind+'/preview',files={'file':('import.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},headers={'Idempotency-Key':str(uuid.uuid4())})

def test_training_record_builds_reference_masters_without_changing_fact(client2,app2):
    auth(client2,'depthead')
    r=post(client2,'/api/records',{'external_key':'MASTER-REC-1','teacher_id':'teacher','event_date':'2026-09-18','project_code':'P-MASTER','project':'主数据项目','course_code':'C-MASTER','course':'主数据课程','class_code':'CL-MASTER','class_name':'主数据班级','hours':'2.5','participants':10,'income':'500.00','source':'经核对人工录入',**REASON})
    assert r.status_code==200,r.text
    with Session(app2.state.engine) as db:
        p=db.scalar(select(TrainingProjectMaster).where(TrainingProjectMaster.code=='P-MASTER'))
        c=db.scalar(select(TrainingCourseMaster).where(TrainingCourseMaster.code=='C-MASTER'))
        cl=db.scalar(select(TrainingClassMaster).where(TrainingClassMaster.code=='CL-MASTER'))
        assert p and p.name=='主数据项目' and c and c.project_code=='P-MASTER' and cl and cl.course_code=='C-MASTER'


def test_teacher_participation_is_separate_from_teaching_and_excel_roundtrip(client2,app2):
    auth(client2,'depthead')
    data=fill_template(client2,'participations',['LEARN-001','J001','2026-09-12','PX','外出培训项目','CX','学习课程',8,'Excel导入'])
    r=upload_preview(client2,'participations',data);assert r.status_code==200,r.text
    job=r.json();assert job['errors']==0
    r=post(client2,'/api/imports/'+job['id']+'/confirm',{'expected_version':job['version'],**REASON});assert r.status_code==200,r.text
    auth(client2,'teacher')
    summary=client2.get('/api/teacher-activity-summary');assert summary.status_code==200,summary.text
    body=summary.json();assert body['teaching']['records']==1 and body['participation']['records']==1
    assert str(body['participation']['hours']) in {'8','8.0','8.00'}
    rows=client2.get('/api/participations');assert rows.status_code==200 and rows.json()['total']==1
    # Teaching ledger remains the original one record; participation never adds income/person-times there.
    assert client2.get('/api/records').json()['total']==1


def test_middle_inbound_preview_confirm_project_course_and_teacher_hr_revoke_scope(client2,app2):
    # Establish a real teacher session before HR scope change.
    auth(client2,'teacher');assert client2.get('/api/auth/me').status_code==200
    teacher_cookie=dict(client2.cookies)
    auth(client2,'schooladmin')
    c=configure_middle(client2,[
      {'source':'staff','target':'staff_no','transform':'text','required':True,'default':None},
      {'source':'name','target':'name','transform':'text','required':True,'default':None},
      {'source':'title','target':'title','transform':'text','required':False,'default':''},
      {'source':'org','target':'org_code','transform':'text','required':True,'default':None},
      {'source':'active','target':'active','transform':'bool','required':True,'default':True},
    ])
    preview=post(client2,'/api/connections/'+c['id']+'/inbound/preview',{'data_type':'teacher_hr','records':[{'staff':'J001','name':'测试教师一（中台）','title':'副教授','org':'dept2','active':True}],'source_label':'教师人事脱敏样本',**REASON})
    assert preview.status_code==200,preview.text;batch=preview.json();assert batch['errors']==0 and batch['state']=='PREVIEW'
    r=post(client2,'/api/inbound-batches/'+batch['id']+'/confirm',{'expected_version':batch['version'],**REASON});assert r.status_code==200,r.text
    with Session(app2.state.engine) as db:
        profile=db.scalar(select(TeacherHrProfile).where(TeacherHrProfile.staff_no=='J001'));u=db.scalar(select(User).where(User.staff_no=='J001'))
        assert profile and profile.org_code=='dept2' and profile.user_id==u.id and u.org_id=='dept2' and u.name.endswith('（中台）')
    # Old teacher cookie is now stale because HR-driven authorization scope changed.
    from fastapi.testclient import TestClient
    stale=TestClient(app2,raise_server_exceptions=False);stale.headers['origin']='http://testserver'
    for k,v in teacher_cookie.items():stale.cookies.set(k,v)
    assert stale.get('/api/auth/me').status_code==401

    # Project and course use the same preview -> explicit confirm chain.
    c=configure_middle(client2,[
      {'source':'projectCode','target':'code','transform':'text','required':True,'default':None},
      {'source':'projectName','target':'name','transform':'text','required':True,'default':None},
      {'source':'enabled','target':'active','transform':'bool','required':True,'default':True},
    ])
    p=post(client2,'/api/connections/'+c['id']+'/inbound/preview',{'data_type':'training_project','records':[{'projectCode':'MID-P1','projectName':'中台项目','enabled':True}],'source_label':'项目脱敏样本',**REASON}).json()
    assert post(client2,'/api/inbound-batches/'+p['id']+'/confirm',{'expected_version':p['version'],**REASON}).status_code==200
    c=configure_middle(client2,[
      {'source':'courseCode','target':'code','transform':'text','required':True,'default':None},
      {'source':'courseName','target':'name','transform':'text','required':True,'default':None},
      {'source':'projectCode','target':'project_code','transform':'text','required':True,'default':None},
      {'source':'enabled','target':'active','transform':'bool','required':True,'default':True},
    ])
    co=post(client2,'/api/connections/'+c['id']+'/inbound/preview',{'data_type':'course','records':[{'courseCode':'MID-C1','courseName':'中台课程','projectCode':'MID-P1','enabled':True}],'source_label':'课程脱敏样本',**REASON}).json()
    assert post(client2,'/api/inbound-batches/'+co['id']+'/confirm',{'expected_version':co['version'],**REASON}).status_code==200
    with Session(app2.state.engine) as db:
        assert db.scalar(select(TrainingProjectMaster).where(TrainingProjectMaster.code=='MID-P1'))
        assert db.scalar(select(TrainingCourseMaster).where(TrainingCourseMaster.code=='MID-C1')).project_code=='MID-P1'


def test_inbound_preview_with_errors_never_writes_master(client2,app2):
    auth(client2,'schooladmin')
    c=configure_middle(client2,[{'source':'projectCode','target':'code','transform':'text','required':True,'default':None}])
    p=post(client2,'/api/connections/'+c['id']+'/inbound/preview',{'data_type':'training_project','records':[{'projectCode':'BAD-P'}],'source_label':'缺字段样本',**REASON}).json()
    assert p['errors']==1
    assert post(client2,'/api/inbound-batches/'+p['id']+'/confirm',{'expected_version':p['version'],**REASON}).status_code==409
    with Session(app2.state.engine) as db:assert db.scalar(select(TrainingProjectMaster).where(TrainingProjectMaster.code=='BAD-P')) is None


def test_delivery_configuration_xlsx_and_handover_contain_new_open_metadata(client2,app2):
    auth(client2,'schooladmin')
    x=client2.get('/api/delivery/configuration.xlsx');assert x.status_code==200,x.text
    wb=load_workbook(io.BytesIO(x.content),read_only=True);assert {'组织机构','角色权限','基础参数与字典','接口配置','接口数据类型','数据库字典'}<=set(wb.sheetnames);wb.close()
    r=post(client2,'/api/exports',{'kind':'handover','filters':{},'confirmed':True});assert r.status_code==200,r.text
    task=r.json();assert run_once(app2.state.settings)['status']=='DONE'
    zresp=client2.get('/api/tasks/'+task['id']+'/download');assert zresp.status_code==200
    with zipfile.ZipFile(io.BytesIO(zresp.content)) as z:
        names=set(z.namelist())
        assert {'data/teacher_participations.csv','data/training_projects.csv','data/training_courses.csv','data/teacher_hr_profiles.csv','delivery/initialization_configuration.json','delivery/database_dictionary.json','delivery/interface_data_catalog.json'}<=names
        dictionary=json.loads(z.read('delivery/database_dictionary.json'))
        assert any(x['table']=='teacher_participations' for x in dictionary)

def test_data_type_specific_mappings_do_not_bleed_between_payloads(client2):
    auth(client2,'schooladmin')
    c=configure_middle(client2,[
      {'data_type':'teacher_hr','source':'teacherNo','target':'staff_no','transform':'text','required':True,'default':None},
      {'data_type':'teacher_hr','source':'teacherName','target':'name','transform':'text','required':True,'default':None},
      {'data_type':'teacher_hr','source':'orgCode','target':'org_code','transform':'text','required':True,'default':None},
      {'data_type':'training_project','source':'projectCode','target':'code','transform':'text','required':True,'default':None},
      {'data_type':'training_project','source':'projectName','target':'name','transform':'text','required':True,'default':None},
    ])
    a=post(client2,'/api/connections/validate',{'connection_id':c['id'],'data_type':'teacher_hr','records':[{'teacherNo':'J001','teacherName':'张老师','orgCode':'dept1'}],**REASON})
    assert a.status_code==200,a.text;assert a.json()['items'][0]['mapped']=={'staff_no':'J001','name':'张老师','org_code':'dept1'}
    b=post(client2,'/api/connections/validate',{'connection_id':c['id'],'data_type':'training_project','records':[{'projectCode':'P9','projectName':'项目九'}],**REASON})
    assert b.status_code==200,b.text;assert b.json()['items'][0]['mapped']=={'code':'P9','name':'项目九'}

def test_unverified_legacy_helper_only_not_a_supported_upgrade(tmp_path):
    import importlib.util
    import sqlalchemy as sa
    from pathlib import Path
    migration=Path(__file__).resolve().parents[2]/'backend/alembic_legacy_unverified/20260917_0002_org_assessment_and_sync.py'
    spec=importlib.util.spec_from_file_location('nd_migration_0002',migration);mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    engine=sa.create_engine('sqlite:///'+str(tmp_path/'upgrade.db'))
    md=sa.MetaData()
    sa.Table('organizations',md,sa.Column('id',sa.String(36),primary_key=True),sa.Column('code',sa.String(64),nullable=False))
    sa.Table('users',md,sa.Column('id',sa.String(36),primary_key=True),sa.Column('staff_no',sa.String(64)),sa.Column('name',sa.String(100)),sa.Column('title',sa.String(80)),sa.Column('org_id',sa.String(36)),sa.Column('active',sa.Boolean()))
    sa.Table('training_records',md,sa.Column('id',sa.String(36),primary_key=True),sa.Column('created_at',sa.DateTime()),sa.Column('project_code',sa.String(64)),sa.Column('project',sa.String(160)),sa.Column('course_code',sa.String(64)),sa.Column('course',sa.String(160)),sa.Column('class_code',sa.String(64)),sa.Column('class_name',sa.String(160)))
    sa.Table('teacher_hr_profiles',md,sa.Column('staff_no',sa.String(64),primary_key=True),sa.Column('user_id',sa.String(36)),sa.Column('name',sa.String(100)),sa.Column('title',sa.String(80)),sa.Column('org_code',sa.String(64)),sa.Column('active',sa.Boolean()),sa.Column('source',sa.String(100)),sa.Column('extra',sa.JSON()),sa.Column('source_updated_at',sa.DateTime()),sa.Column('version',sa.Integer()),sa.Column('id',sa.String(36)),sa.Column('created_at',sa.DateTime()))
    sa.Table('training_projects',md,sa.Column('code',sa.String(64),primary_key=True),sa.Column('name',sa.String(160)),sa.Column('external_key',sa.String(128)),sa.Column('start_date',sa.Date()),sa.Column('end_date',sa.Date()),sa.Column('active',sa.Boolean()),sa.Column('source',sa.String(100)),sa.Column('version',sa.Integer()),sa.Column('id',sa.String(36)),sa.Column('created_at',sa.DateTime()))
    sa.Table('training_courses',md,sa.Column('code',sa.String(64),primary_key=True),sa.Column('name',sa.String(160)),sa.Column('project_code',sa.String(64)),sa.Column('external_key',sa.String(128)),sa.Column('active',sa.Boolean()),sa.Column('source',sa.String(100)),sa.Column('version',sa.Integer()),sa.Column('id',sa.String(36)),sa.Column('created_at',sa.DateTime()))
    sa.Table('training_classes',md,sa.Column('code',sa.String(64),primary_key=True),sa.Column('name',sa.String(160)),sa.Column('project_code',sa.String(64)),sa.Column('course_code',sa.String(64)),sa.Column('start_date',sa.Date()),sa.Column('end_date',sa.Date()),sa.Column('active',sa.Boolean()),sa.Column('source',sa.String(100)),sa.Column('version',sa.Integer()),sa.Column('id',sa.String(36)),sa.Column('created_at',sa.DateTime()))
    sa.Table('teacher_participations',md,sa.Column('id',sa.String(36),primary_key=True))
    md.create_all(engine)
    now=mod._utcnow()
    with engine.begin() as conn:
        conn.execute(md.tables['organizations'].insert(),[{'id':'dept1','code':'D01'}])
        conn.execute(md.tables['users'].insert(),[{'id':'u1','staff_no':'J001','name':'旧库教师','title':'讲师','org_id':'dept1','active':True},{'id':'admin','staff_no':None,'name':'管理员','title':'','org_id':'dept1','active':True}])
        conn.execute(md.tables['training_records'].insert(),[
          {'id':'r1','created_at':now,'project_code':'P1','project':'项目旧名','course_code':'C1','course':'课程旧名','class_code':'CL1','class_name':'班级旧名'},
          {'id':'r2','created_at':now,'project_code':'P1','project':'项目最终名','course_code':'C1','course':'课程最终名','class_code':'CL1','class_name':'班级最终名'}])
        mod._backfill_reference_data(conn)
        assert conn.execute(sa.select(sa.func.count()).select_from(md.tables['teacher_participations'])).scalar_one()==0
        hr=conn.execute(sa.select(md.tables['teacher_hr_profiles'])).mappings().one();assert hr['staff_no']=='J001' and hr['org_code']=='D01' and hr['source']=='升级回填：本地账号'
        p=conn.execute(sa.select(md.tables['training_projects'])).mappings().one();assert p['code']=='P1' and p['name']=='项目最终名'
        c=conn.execute(sa.select(md.tables['training_courses'])).mappings().one();assert c['code']=='C1' and c['name']=='课程最终名' and c['project_code']=='P1'
        cl=conn.execute(sa.select(md.tables['training_classes'])).mappings().one();assert cl['code']=='CL1' and cl['name']=='班级最终名' and cl['course_code']=='C1'
