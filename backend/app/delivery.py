"""School-owned configuration/data-dictionary delivery helpers."""
from __future__ import annotations
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from .models import Base,Org,Role,Period,Setting,Connection
from .core import serial
from .integrations import catalog_for
from .inbound_sync import IntegrationSchedule

SECRETISH={'password_hash','token_hash','csrf','auth_version','storage_key','failures','locked_until'}

def configuration_snapshot(db):
    return {
      'organizations':serial(db.scalars(select(Org).order_by(Org.code)).all()),
      'roles':serial(db.scalars(select(Role).order_by(Role.kind,Role.code)).all()),
      'periods':serial(db.scalars(select(Period).order_by(Period.start_date,Period.name)).all()),
      'settings':serial(db.scalars(select(Setting).where(~Setting.key.like('ssoid:%'),~Setting.key.like('ssostate:%')).order_by(Setting.key)).all()),
      'connections':[{'id':c.id,'code':c.code,'name':c.name,'state':c.state,'mappings':serial(c.mappings),'config':serial(c.config),'version':c.version,
                      'data_types':catalog_for(c)} for c in db.scalars(select(Connection).order_by(Connection.code))],
      'integration_schedules':serial(db.scalars(select(IntegrationSchedule).order_by(IntegrationSchedule.connection_id,IntegrationSchedule.data_type)).all())
    }

def database_dictionary():
    rows=[]
    for table in Base.metadata.sorted_tables:
        for col in table.columns:
            fks=','.join(sorted(str(fk.target_fullname) for fk in col.foreign_keys))
            rows.append({'table':table.name,'column':col.name,'type':str(col.type),'nullable':bool(col.nullable),'primary_key':bool(col.primary_key),'foreign_key':fks,
                         'sensitive':col.name in SECRETISH})
    return rows

def _sheet(ws,headers,rows):
    ws.append(headers);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}1'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78');cell.alignment=Alignment(horizontal='center')
    for row in rows:ws.append(row)
    for idx,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(idx)].width=max(12,min(42,len(str(h))*2+8))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)

def delivery_workbook(db):
    snap=configuration_snapshot(db);wb=Workbook();wb.remove(wb.active)
    ws=wb.create_sheet('组织机构');_sheet(ws,['编码','名称','层级','上级ID','启用','版本'],[[x['code'],x['name'],x['kind'],x.get('parent_id') or '',x['active'],x['version']] for x in snap['organizations']])
    ws=wb.create_sheet('角色权限');_sheet(ws,['编码','名称','层级','启用','权限'],[[x['code'],x['name'],x['kind'],x['active'],'\n'.join(x['permissions'])] for x in snap['roles']])
    ws=wb.create_sheet('基础参数与字典');_sheet(ws,['配置键','版本','当前值(JSON)'],[[x['key'],x['version'],json.dumps(x['value'],ensure_ascii=False,indent=2)] for x in snap['settings']])
    ws=wb.create_sheet('考核周期');_sheet(ws,['名称','开始日期','结束日期','状态','版本'],[[x['name'],x['start_date'],x['end_date'],x['state'],x['version']] for x in snap['periods']])
    conrows=[];catrows=[]
    for c in snap['connections']:
        conrows.append([c['code'],c['name'],c['state'],c['version'],json.dumps(c['config'],ensure_ascii=False,indent=2),json.dumps(c['mappings'],ensure_ascii=False,indent=2)])
        for d in c['data_types']:catrows.append([c['code'],c['name'],d['data_type'],d['name'],d['local_status'],d['connection_state']])
    ws=wb.create_sheet('接口配置');_sheet(ws,['接口编码','接口名称','连接状态','版本','参数(JSON)','字段映射(JSON)'],conrows)
    ws=wb.create_sheet('接口数据类型');_sheet(ws,['接口编码','接口名称','数据类型','业务名称','本地承载','连接状态'],catrows)
    ws=wb.create_sheet('定时同步计划');_sheet(ws,['接口ID','数据类型','频率(分钟)','自动确认','启用','批准人','下次执行','上次执行','上次游标','上次错误','版本'],[[x['connection_id'],x['data_type'],x['frequency_minutes'],x['auto_commit'],x['active'],x['approved_by'],x['next_at'],x.get('last_run_at') or '',x.get('last_cursor') or '',x.get('last_error') or '',x['version']] for x in snap['integration_schedules']])
    ws=wb.create_sheet('数据库字典');drows=database_dictionary();_sheet(ws,['表','字段','类型','可空','主键','外键','敏感字段'],[[r['table'],r['column'],r['type'],r['nullable'],r['primary_key'],r['foreign_key'],r['sensitive']] for r in drows])
    wb.properties.title='一站式非学历教育智慧前台-初始化配置与数据字典';wb.properties.creator='湖南跃科信息工程有限公司'
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()
