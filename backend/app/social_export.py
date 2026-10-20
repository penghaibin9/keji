"""Chinese business exports using the installed workbook library."""
from io import BytesIO
from openpyxl import Workbook
from .social_training import STATES


def workbook(payload):
    book=Workbook();sheet=book.active;sheet.title='业务概要'
    obj=payload['item'];sheet.append(['项目','内容'])
    sheet.append(['办理状态',payload['state_label']]);sheet.append(['记录版本',obj['version']])
    nature='已归档结果' if obj['state']=='ARCHIVED' else '过程材料或候选试算，非正式结果'
    if payload.get('facts') is not None:
        nature='已结项项目材料' if obj['state']=='CLOSED' else '已备案项目材料（结项尚未完成）' if obj['state'] in {'FILED','CLOSE_COLLEGE','CLOSE_REVIEW'} else '项目办理过程材料（尚未完成备案或结项）'
        if obj['state']=='VOID' or payload['facts'].get('invalidated'):nature='项目作废后的历史材料（不得作为有效结项依据）'
    elif obj['state']=='PUBLIC':nature='公示中的积分汇总（尚未归档）' if obj.get('kind')=='point_report' else '公示中的考核材料（尚未归档）'
    elif obj['state']=='ACTIVE':nature='已生效业务资料（当前启用）'
    if obj['data'].get('superseded_by'):nature='已被新公示替代的历史材料'
    sheet.append(['正式性',nature])
    if payload.get('readonly'):sheet.append(['查看范围',payload.get('readonly_reason','仅本人参与摘要，不含学员名册、财务和他人工作资料')])
    if payload.get('public_summary'):sheet.append(['查看范围','仅公开年度评分摘要，不含项目原始证据及个人积分明细'])
    for reason in obj.get('closeout_pending',[]):sheet.append(['历史资料待补',reason])
    trial=payload.get('trial') or {}
    for key,label in [('teaching','授课积分'),('auxiliary','辅助积分'),('base','考核基础分'),('bonus_awarded','认定加分'),('total','合计'),('grade','认定等级' if obj['state']=='ARCHIVED' else '候选等级'),('deadline','申报期限')]:
        if key in trial:sheet.append([label,'待核定' if trial[key] is None else str(trial[key])])
    for reason in trial.get('blockers',[]):sheet.append(['待确认事项',reason])
    files=book.create_sheet('佐证材料');files.append(['名称','字节数','完整性校验值'])
    for file in payload.get('attachments',[]):files.append([file['name'],file['size'],file['sha256']])
    flow=book.create_sheet('办理记录');flow.append(['办理人','操作','时间','说明','依据'])
    for row in payload.get('events',[]):flow.append([row['actor_name'],row['action'],row['created_at'],row['reason'],row['basis']])
    detail=book.create_sheet('计算明细');detail.append(['核算单元','日期','有效人次','标准人数','层级权重','场景权重','岗位权重','积分'])
    for row in trial.get('details',[]):detail.append([row.get(k) for k in ('unit','date','participants','standard','level','scene','role','score')])
    facts=payload.get('facts')
    if facts:
        project=obj['data']
        for key,label in [('name','项目名称'),('code','项目编号'),('cohort','期次'),('expected_income','预计收入'),('start_date','培训开始'),('end_date','培训结束')]:sheet.append([label,project.get(key)])
        if payload.get('readonly'):
            summary=book.create_sheet('参与摘要');summary.append(['核算单元','课程编号','班级编号','培训日期','模块','教师','岗位'])
            for unit in facts['units']:
                u=unit['data']
                for teacher in u['teachers']:
                    summary.append([unit['external_key'],u['course_code'],u['class_code'],u['event_date'],u['module'],teacher.get('teacher_name','授权教师'),{'lecturer':'主讲','instructor':'实训指导','manager':'管理辅助'}.get(teacher['role'],'待确认')])
        participants=book.create_sheet('有效人次依据');participants.append(['单元编号','课程编号','班级编号','模块','培训日期','学员内部编号','实名签到','学时达标','全程参训','结业通过','全日制生','准毕业生例外','重点群体','剔除说明'])
        from .social_training import eligible
        for unit in facts['units']:
            u=unit['data']
            for r in u['learners']:
                participants.append([unit['external_key'],u['course_code'],u['class_code'],u['module'],u['event_date'],r['reference'],*['是' if r[k] else '否' for k in ('signed','hours_met','full_attendance','passed','full_time_student','graduate_exception','priority_group')],r['exclusion_reason'] or ('有效' if eligible(r) else '未满足全部参训条件')])
    if trial.get('metrics'):
        annual=book.create_sheet('年度指标');annual.append(['指标','实际值','任务目标','完成率','基础分','加分'])
        for row in trial['metrics']:annual.append([row.get(k) for k in ('label','actual','target','completion','base','bonus')])
        for k,l in [('achievement_base','成果基础分'),('cost_base','资金管理基础分'),('bonus_raw','封顶前加分'),('bonus_awarded','封顶后加分'),('veto','一票否决')]:annual.append([l,str(trial.get(k))])
        for row in trial.get('cost_details',[])+trial.get('outcome_details',[]):annual.append([row['label'],row['value']])
    data=obj['data']
    if obj.get('kind')=='point_report':
        if data.get('superseded_by'):
            sheet.append(['替代说明','已被新公示替代，旧快照保留，不作为当前公示或归档依据']);sheet.append(['替代日期',data.get('superseded_on')])
        elif data.get('replaces'):sheet.append(['重新公示','按最新归档积分重新公示，期限从本次发布日重新计算'])
        report=book.create_sheet('已认定积分汇总');report.append(['教师','学院','项目','授课积分','辅助积分','合计'])
        for row in data['rows']:report.append([row[k] for k in ('teacher_name','org_name','project','teaching','auxiliary','total')])
        sheet.append(['公示开始',data.get('public_start')]);sheet.append(['公示截止',data.get('public_end')])
    if obj.get('kind')=='correction':
        sheet.append(['授课积分更正',data['teaching']]);sheet.append(['辅助积分更正',data['auxiliary']]);sheet.append(['正式更正依据',data['decision_reference']])
    for page in book:
        page.freeze_panes='A2'
        for col in page.columns:page.column_dimensions[col[0].column_letter].width=26
        for row in page:
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.value="'"+cell.value
    stream=BytesIO();book.save(stream);return stream.getvalue()
