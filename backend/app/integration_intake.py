"""School-facing integration intake workbook and safe draft staging.

The workbook contains configuration facts and evidence references only.
It never carries real passwords, tokens, private keys or certificate bodies.
"""
from __future__ import annotations
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
import json
import re
import zipfile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select

from .models import Connection, utcnow
from .integrations import DATA_CATALOG
from .integration_center import MATERIALS, candidate, source_records, stored
from .integration_profile_templates import templates, merge_template

FORMAT = 'YUEKE_INTEGRATION_INTAKE_V1'
CODES = ('sso', 'middle', 'province')
NAMES = {'sso':'学校统一认证', 'middle':'学校数据中台', 'province':'湘培直通车'}
SECRET_REFS = {'sso':'ND_SSO_SECRET', 'middle':'ND_MIDDLE_SECRET', 'province':'ND_PROVINCE_SECRET'}

SSO_PARAMS = [
 ('认证','protocol','认证协议','cas / oidc / oauth2 / json_ticket'),
 ('认证','login_path','登录入口相对路径','CAS/票据登录入口'),
 ('认证','validate_path','CAS票据校验路径','CAS使用'),
 ('认证','authorize_path','授权入口相对路径','OIDC/OAuth2使用'),
 ('认证','token_path','令牌交换相对路径','OIDC/OAuth2使用'),
 ('认证','userinfo_path','用户信息相对路径','OIDC/OAuth2使用'),
 ('认证','logout_path','退出登录相对路径','学校支持时填写'),
 ('认证','issuer','签发方标识','OIDC使用'),
 ('认证','jwks_path','签名公钥相对路径','OIDC使用'),
 ('认证','scope','授权信息范围','如 openid profile'),
 ('认证','client_auth_style','应用凭据提交位置','client_secret_post / basic'),
 ('身份','identity_root','身份信息根路径','返回报文有外层对象时填写'),
 ('身份','roles_delimiter','多角色分隔符','默认逗号'),
 ('身份','callback_param','授权码字段名','默认 code'),
 ('身份','state_param','防伪状态字段名','默认 state'),
 ('身份','ticket_param','CAS票据字段名','默认 ticket'),
 ('身份','active_true_values','启用取值','多个值用逗号分隔'),
 ('身份','active_false_values','停用取值','多个值用逗号分隔'),
 ('身份','active_default','无启停字段时默认','留空=必须提供；true/false'),
 ('验证','probe.kind','真实验证方式','cas_ticket / authorization_code / json / oidc_metadata'),
 ('验证','probe.callback_url','验证回调地址','一次性验证需要时填写'),
 ('验证','probe.method','身份检查请求方式','GET / POST'),
 ('验证','probe.path','身份检查相对路径','JSON身份接口使用'),
 ('验证','probe.success_path','身份检查成功标志路径','如 data.success'),
 ('验证','probe.success_values','身份检查成功允许值','多个值用逗号分隔'),
 ('验证','probe.credential_header','验证凭据请求头名称','只填名称'),
 ('验证','probe.credential_key','服务器凭据字段名','只填字段名'),
 ('验证','probe.state_key','验证状态字段名','默认 probe_state'),
 ('票据','verify.method','票据校验请求方式','POST / GET'),
 ('票据','verify.path','票据校验相对路径','json_ticket使用'),
 ('票据','verify.body_type','票据提交格式','form / json'),
 ('票据','verify.ticket_field','票据字段名','按正式文档'),
 ('票据','verify.redirect_field','回调地址字段名','默认 redirect_uri'),
 ('票据','verify.success_path','票据校验成功路径','按正式文档'),
 ('票据','verify.success_values','票据校验成功值','多个值用逗号分隔'),
 ('票据','verify.auth.mode','票据接口鉴权方式','none / bearer / api_key / basic / signature'),
 ('票据','verify.auth.header','票据接口令牌头','只填头名称'),
 ('票据','verify.auth.token_path','票据接口取令牌路径','需要时填写'),
 ('票据','verify.auth.scope','票据接口授权范围','需要时填写'),
 ('票据','verify.auth.client_auth_style','应用凭据提交位置','client_secret_post / basic'),
]
IDENTITY_FIELDS = [
 ('stable_id','稳定身份编号'),('account','登录账号'),('staff_no','教师工号'),
 ('name','姓名'),('org_code','组织编码'),('roles','角色/身份标识'),('active','账号启停状态')
]

AUTH_HEADERS = [
 '认证方式','令牌请求头','获取令牌路径','授权范围','应用凭据位置',
 '时间戳字段','时间戳位置','随机数字段','随机数位置','签名字段','签名位置',
 '幂等字段','幂等位置','应用标识字段','应用标识位置','签名原文组成','签名算法','签名编码'
]
MIDDLE_HEADERS = [
 '数据类型','业务名称','请求方式','接口相对路径',*AUTH_HEADERS,
 '机构编码位置','机构编码字段','翻页方式','每页数量','最大页数','起始页','结束判定',
 '游标位置','游标字段','页码位置','页码字段','偏移位置','偏移字段','每页数量位置','每页数量字段','同步模式位置','同步模式字段',
 '成功标志路径','成功允许值','数据列表路径','下一游标路径','是否还有更多路径',
 '总数路径','增量水位路径','完整规则JSON'
]
PROVINCE_HEADERS = [
 '数据类型','业务名称','请求方式','接口相对路径',*AUTH_HEADERS,
 '机构编码位置','机构编码字段','业务数据字段',
 '接收成功路径','接收成功值','返回代码路径','回执编号路径','失败代码','完整规则JSON'
]
MAPPING_HEADERS = ['接口编码','数据类型','业务名称','源字段','目标字段','转换方式','必填','默认值']
MATERIAL_HEADERS = ['接口编码','对接项目','资料项目','责任人','文档或证据编号','双方确认日期']
CREDENTIAL_HEADERS = [
 '接口编码','对接项目','凭据类型','安全保管方式','轮换责任人','证书SHA-256指纹',
 '证书主题','证书签发方','证书生效日','证书到期日','上次轮换日期','下次计划轮换日期'
]
SOURCE_HEADERS = ['接口编码','数据类型','业务名称','本地承载状态','权威来源系统','数据提供方式','同步/调用频率','历史数据范围','责任人','责任边界说明','确认依据/文档编号','双方确认日期']
ROLE_HEADERS = ['本地角色层级','中文层级','学校外部角色/身份标识','本地授权策略','数据范围说明','学校确认人','确认日期']
ROLE_SCOPE = {
    'teacher':('普通教师','外部角色仅作身份参考；实际授权以本地账号角色和组织为准','仅本人数据'),
    'dept':('教研室','外部角色仅作身份参考；实际授权以本地账号角色和组织为准','本教研室数据'),
    'college':('二级学院','外部角色仅作身份参考；实际授权以本地账号角色和组织为准','本学院数据'),
    'school':('学校管理部门','外部角色仅作身份参考；实际授权以本地账号角色和组织为准','按授权访问全校数据')
}


def _text(value, limit=20000):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    value = str(value).strip()
    if len(value) > limit:
        raise HTTPException(422, '学校资料表单元格内容过长')
    if value.startswith("'") and len(value) > 1 and value[1] in '=+-@':
        value = value[1:]
    return value


def _excel_value(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(',', ':'))
    if isinstance(value, bool):
        value = 'true' if value else 'false'
    value = str(value)
    return "'" + value if value and value[:1] in '=+-@' else value


def _list_text(values):
    return '，'.join('true' if x is True else 'false' if x is False else str(x)
                     for x in (values or []))


def _parse_values(value):
    out = []
    for raw in re.split(r'[，,\n]', _text(value)):
        item = raw.strip()
        if not item:
            continue
        low = item.lower()
        if low == 'true':
            out.append(True)
        elif low == 'false':
            out.append(False)
        elif re.fullmatch(r'-?\d+', item):
            out.append(int(item))
        else:
            out.append(item)
    return out


def _int(value, default, low, high):
    raw = _text(value)
    if raw == '':
        return default
    try:
        result = int(raw)
    except ValueError:
        raise HTTPException(422, f'学校资料表数字格式无效：{raw}') from None
    if not low <= result <= high:
        raise HTTPException(422, f'学校资料表数字超出允许范围：{raw}')
    return result


def _bool(value):
    raw = _text(value).lower()
    if raw in {'是','true','1','必填'}:
        return True
    if raw in {'否','false','0','非必填',''}:
        return False
    raise HTTPException(422, '字段映射“必填”只能填写是或否')


def _date_text(value):
    if value in (None, ''):
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _text(value, 32)
    try:
        return date.fromisoformat(raw).isoformat()
    except ValueError:
        raise HTTPException(422, f'日期格式无效：{raw}') from None


def _nested_get(data, path):
    current = data
    for key in path.split('.'):
        if not isinstance(current, dict):
            return ''
        current = current.get(key)
    return current if current is not None else ''


def _nested_set(data, path, value):
    keys = path.split('.')
    current = data
    for key in keys[:-1]:
        child = current.get(key)
        if not isinstance(child, dict):
            child = {}
            current[key] = child
        current = child
    current[keys[-1]] = value


def _nested_delete(data, path):
    keys = path.split('.')
    current = data
    parents = []
    for key in keys[:-1]:
        child = current.get(key)
        if not isinstance(child, dict):
            return
        parents.append((current, key))
        current = child
    current.pop(keys[-1], None)
    for parent, key in reversed(parents):
        if isinstance(parent.get(key), dict) and not parent[key]:
            parent.pop(key, None)


def _binding(request, token, preferred=None):
    order = ['query','headers','body']
    if preferred in {'query','header','body'}:
        key = 'headers' if preferred == 'header' else preferred
        order = [key] + [x for x in order if x != key]
    for target in order:
        data = request.get(target)
        if not isinstance(data, dict):
            continue
        for name, value in data.items():
            if value == token:
                return ('header' if target == 'headers' else target), name
    return 'query', ''


def _clear_binding(request, token):
    for target in ('query','headers','body'):
        data = request.get(target)
        if not isinstance(data, dict):
            continue
        for name in [name for name, value in data.items() if value == token]:
            data.pop(name, None)


def _set_binding(request, target, name, token):
    _clear_binding(request, token)
    name = _text(name, 120)
    if not name:
        return
    key = 'headers' if target == 'header' else target
    if key not in {'query','headers','body'}:
        raise HTTPException(422, '参数放置位置只能是查询参数、请求头或请求正文')
    if not isinstance(request.get(key), dict):
        request[key] = {}
    request[key][name] = token


def _auth_row(auth):
    auth = auth or {}
    parts = auth.get('canonical_parts') or []
    return [
        auth.get('mode','none'), auth.get('header',''), auth.get('token_path',''), auth.get('scope',''),
        auth.get('client_auth_style','client_secret_post'),
        auth.get('timestamp_name',auth.get('timestamp_header','X-Timestamp')), auth.get('timestamp_target','header'),
        auth.get('nonce_name',auth.get('nonce_header','X-Nonce')), auth.get('nonce_target','header'),
        auth.get('signature_name',auth.get('signature_header','X-Signature')), auth.get('signature_target','header'),
        auth.get('idempotency_name',auth.get('idempotency_header','')), auth.get('idempotency_target','header'),
        auth.get('app_id_name',auth.get('app_id_header','')), auth.get('app_id_target','header'),
        '，'.join(parts), auth.get('algorithm','hmac_sha256'), auth.get('encoding','hex')
    ]


def _apply_auth(current, values):
    auth = deepcopy(current or {})
    keys = ['mode','header','token_path','scope','client_auth_style',
            'timestamp_name','timestamp_target','nonce_name','nonce_target',
            'signature_name','signature_target','idempotency_name','idempotency_target',
            'app_id_name','app_id_target','canonical_parts','algorithm','encoding']
    for key, value in zip(keys, values):
        if key == 'canonical_parts':
            auth[key] = [x.strip() for x in re.split(r'[，,\n]', _text(value)) if x.strip()]
        else:
            auth[key] = _text(value)
    auth['mode'] = auth.get('mode') or 'none'
    return auth


def _sheet(workbook, title, headers):
    ws = workbook.create_sheet(title)
    ws.append([_excel_value(x) for x in headers])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    return ws


def _finish(ws, widths=None):
    widths = widths or {}
    for idx in range(1, ws.max_column + 1):
        header = _text(ws.cell(1, idx).value)
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, 18)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[1].height = 32


def _append(ws, values):
    ws.append([_excel_value(value) for value in values])


def _dropdown(ws, column, values, start=2, end=300):
    validator = DataValidation(type='list', formula1='"'+','.join(values)+'"', allow_blank=True)
    ws.add_data_validation(validator)
    validator.add(f'{column}{start}:{column}{end}')


TOKEN_INSTITUTION = '$' + '{institution_code}'
TOKEN_CURSOR = '$' + '{cursor}'
TOKEN_PAGE = '$' + '{page}'
TOKEN_OFFSET = '$' + '{offset}'
TOKEN_LIMIT = '$' + '{limit}'
TOKEN_MODE = '$' + '{mode}'
TOKEN_PAYLOAD = '$' + '{payload}'


def _rule_to_middle(code, name, rule):
    rule = rule or {}
    request = rule.get('request') or {}
    response = rule.get('response') or {}
    paging = rule.get('pagination') or {}
    target, institution = _binding(request, TOKEN_INSTITUTION)
    paging_mode = paging.get('mode','cursor')
    if paging_mode == 'page':
        page_target, page = _binding(request, TOKEN_PAGE)
        offset_target, offset = _binding(request, TOKEN_OFFSET, page_target if page else None)
        primary_target = page_target if page else offset_target if offset else None
        limit_target, limit = _binding(request, TOKEN_LIMIT, primary_target)
        cursor_target, cursor = _binding(request, TOKEN_CURSOR)
        mode_target, mode = _binding(request, TOKEN_MODE, primary_target)
    else:
        cursor_target, cursor = _binding(request, TOKEN_CURSOR)
        primary_target = cursor_target if cursor else None
        limit_target, limit = _binding(request, TOKEN_LIMIT, primary_target)
        mode_target, mode = _binding(request, TOKEN_MODE, primary_target)
        page_target, page = _binding(request, TOKEN_PAGE)
        offset_target, offset = _binding(request, TOKEN_OFFSET, page_target if page else None)
    return [code,name,request.get('method','GET'),request.get('path',''),*_auth_row(request.get('auth')),
        target,institution,paging.get('mode','cursor'),paging.get('page_size',1000),paging.get('max_pages',20),
        paging.get('start_page',1),paging.get('end_condition',''),cursor_target,cursor,page_target,page,offset_target,offset,limit_target,limit,mode_target,mode,
        response.get('success_path',''),_list_text(response.get('success_values')),
        response.get('records_path',''),response.get('next_cursor_path',''),response.get('has_more_path',''),
        response.get('total_path',''),response.get('watermark_path',''),
        json.dumps(rule,ensure_ascii=False,separators=(',',':'))]


def _rule_to_province(code, name, rule):
    rule = rule or {}
    request = rule.get('request') or {}
    response = rule.get('response') or {}
    target, institution = _binding(request, TOKEN_INSTITUTION)
    _, payload = _binding(request, TOKEN_PAYLOAD)
    return [code,name,request.get('method','POST'),request.get('path',''),*_auth_row(request.get('auth')),
        target,institution,payload,response.get('accepted_path',''),_list_text(response.get('accepted_values')),
        response.get('code_path',''),response.get('receipt_id_path',''),
        '，'.join(response.get('error_codes') or []),
        json.dumps(rule,ensure_ascii=False,separators=(',',':'))]


def _connector_state(db, code):
    connection = db.scalar(select(Connection).where(Connection.code == code))
    if not connection:
        raise HTTPException(503, '三大接口初始化记录不完整')
    item, draft = candidate(db, connection)
    return connection, item, draft


def build_intake_workbook(db):
    states = {code:_connector_state(db, code) for code in CODES}
    wb = Workbook()
    wb.remove(wb.active)

    guide = wb.create_sheet('填写说明')
    guide.append(['跃科 · 三大接口学校资料采集表'])
    guide.append(['格式版本', FORMAT])
    guide.append(['使用方式', '学校信息处/实施人员只填写正式资料；导回系统后先预检，只写草稿，不会自动发布或连接外部系统。'])
    guide.append(['安全边界', '密码、令牌、密钥、私钥、证书正文不得填写到本表；只填写服务器凭据引用和证书指纹/有效期等安全元数据。'])
    guide.append(['实施流程', '接入资料 → 协议与环境 → 回调及凭据 → 字段与身份范围 → 测试与样本预览 → 发布与同步计划 → 日志诊断及交付'])
    guide.append(['采购对应', '统一认证=采购第2项；学校数据中台=第4项；湘培直通车=第11项。'])
    guide.column_dimensions['A'].width = 18
    guide.column_dimensions['B'].width = 110
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    guide['A1'].font = Font(bold=True, size=16)


    basic_headers = ['接口编码','对接项目','正式HTTPS总地址','责任单位/联系人','认证/鉴权模式',
                     '服务器凭据引用','失败重试次数','同步频率(分钟)','学校/机构编码','备注']
    basic = _sheet(wb, '接口基础', basic_headers)
    for code in CODES:
        _, item, _ = states[code]
        config = item.config or {}
        adapter = config.get('adapter') or {}
        institution = (adapter.get('public_parameters') or {}).get('institution_code','')
        _append(basic, [code,NAMES[code],config.get('endpoint',''),config.get('owner',''),
                        config.get('auth_mode',''),SECRET_REFS[code],config.get('retry_count',3),
                        config.get('frequency_minutes',30),institution,''])
    _finish(basic, {'正式HTTPS总地址':42,'责任单位/联系人':28,'备注':35})
    _dropdown(basic, 'E', ['school_sso','cas','oidc','oauth2','token','signature','custom'])

    sso = _sheet(wb, '统一认证参数', ['分类','参数编码','学校/实施填写','填写说明'])
    _, sso_item, _ = states['sso']
    sso_adapter = merge_template(templates(DATA_CATALOG)['sso'], (sso_item.config or {}).get('adapter') or {})
    for group, code, label, help_text in SSO_PARAMS:
        value = _nested_get(sso_adapter, code)
        if code in {'active_true_values','active_false_values','probe.success_values','verify.success_values'}:
            value = _list_text(value)
        elif code == 'active_default':
            value = '' if type(value) is not bool else ('true' if value else 'false')
        _append(sso, [group,code,value,label+'；'+help_text])
    _finish(sso, {'参数编码':28,'学校/实施填写':42,'填写说明':60})


    identity = _sheet(wb, '统一认证身份', ['本平台身份字段','中文含义','学校返回字段路径','要求'])
    identity_map = sso_adapter.get('identity_fields') or {}
    for key, label in IDENTITY_FIELDS:
        requirement = '必填' if key == 'name' else '稳定身份编号/账号/工号至少一项' if key in {'stable_id','account','staff_no'} else '按学校实际'
        _append(identity, [key,label,identity_map.get(key,''),requirement])
    _finish(identity, {'学校返回字段路径':36,'要求':38})

    org_map = _sheet(wb, '统一认证组织映射', ['学校组织编码','本系统组织编码','说明'])
    for upstream, local in sorted((sso_adapter.get('org_map') or {}).items()):
        _append(org_map, [upstream, local, '学校组织编码映射到本系统学院/教研室编码'])
    _finish(org_map, {'学校组织编码':30,'本系统组织编码':30,'说明':50})

    role_ws = _sheet(wb, '统一认证角色确认', ROLE_HEADERS)
    sso_connection, _, _ = states['sso']
    role_row = stored(db, 'identity_policy', sso_connection.id)
    current_roles = {x.get('kind'):x for x in ((role_row.value or {}).get('items',[]) if role_row else []) if isinstance(x,dict)}
    for kind, (label, strategy, scope) in ROLE_SCOPE.items():
        current = current_roles.get(kind,{})
        _append(role_ws, [kind,label,current.get('external_roles',''),strategy,scope,
                          current.get('confirmed_by',''),current.get('confirmed_on','')])
    _finish(role_ws, {'学校外部角色/身份标识':36,'本地授权策略':55,'数据范围说明':28,'学校确认人':24})

    middle = _sheet(wb, '数据中台接口', MIDDLE_HEADERS)
    _, middle_item, _ = states['middle']
    middle_adapter = merge_template(templates(DATA_CATALOG)['middle'], (middle_item.config or {}).get('adapter') or {})
    for code, name, _ in DATA_CATALOG['middle']:
        _append(middle, _rule_to_middle(code,name,(middle_adapter.get('data_types') or {}).get(code)))
    _finish(middle, {'接口相对路径':34,'完整规则JSON':70,'业务名称':24})
    _dropdown(middle, 'C', ['GET','POST'])
    _dropdown(middle, 'E', ['none','bearer','api_key','basic','oauth2_client_credentials','signature'])

    province = _sheet(wb, '湘培接口', PROVINCE_HEADERS)
    _, province_item, _ = states['province']
    province_adapter = merge_template(templates(DATA_CATALOG)['province'], (province_item.config or {}).get('adapter') or {})
    for code, name, _ in DATA_CATALOG['province']:
        _append(province, _rule_to_province(code,name,(province_adapter.get('data_types') or {}).get(code)))
    _finish(province, {'接口相对路径':34,'完整规则JSON':70,'业务名称':22})
    _dropdown(province, 'C', ['POST','PUT','PATCH'])
    _dropdown(province, 'E', ['none','bearer','api_key','basic','oauth2_client_credentials','signature'])

    mappings = _sheet(wb, '字段映射', MAPPING_HEADERS)
    names_by_code = {code:{x[0]:x[1] for x in DATA_CATALOG[code]} for code in CODES}
    for code in ('middle','province'):
        _, item, _ = states[code]
        for mapping in item.mappings or []:
            dtype = mapping.get('data_type','')
            _append(mappings, [code,dtype,names_by_code[code].get(dtype,''),
                mapping.get('source',''),mapping.get('target',''),mapping.get('transform','text'),
                '是' if mapping.get('required') else '否',
                '' if mapping.get('default') is None else mapping.get('default')])
    _finish(mappings, {'业务名称':24,'源字段':30,'目标字段':30,'默认值':24})
    _dropdown(mappings, 'F', ['text','int','decimal','date','bool'])
    _dropdown(mappings, 'G', ['是','否'])

    material_ws = _sheet(wb, '资料登记', MATERIAL_HEADERS)
    for code in CODES:
        connection, _, _ = states[code]
        row = stored(db, 'materials', connection.id)
        current = {x.get('name'):x for x in ((row.value or {}).get('items',[]) if row else []) if isinstance(x,dict)}
        for name in MATERIALS:
            value = current.get(name,{})
            _append(material_ws, [code,NAMES[code],name,value.get('owner',''),value.get('reference',''),value.get('confirmed_on','')])
    _finish(material_ws, {'资料项目':30,'责任人':24,'文档或证据编号':40,'双方确认日期':18})

    credential_ws = _sheet(wb, '凭据证书元数据', CREDENTIAL_HEADERS)
    for code in CODES:
        connection, _, _ = states[code]
        row = stored(db, 'credential', connection.id)
        value = deepcopy(row.value or {}) if row else {}
        _append(credential_ws, [code,NAMES[code],value.get('credential_type','token'),
            value.get('managed_by','server_env'),value.get('owner',''),
            value.get('certificate_fingerprint_sha256',''),value.get('certificate_subject',''),
            value.get('certificate_issuer',''),value.get('valid_from',''),value.get('valid_until',''),
            value.get('last_rotated_on',''),value.get('next_rotation_on','')])
    _finish(credential_ws, {'轮换责任人':25,'证书SHA-256指纹':68,'证书主题':36,'证书签发方':36})
    _dropdown(credential_ws, 'C', ['token','oauth_client','hmac','certificate','mtls','custom'])
    _dropdown(credential_ws, 'D', ['server_env','secret_manager','file_mount','hardware','other'])

    source_ws = _sheet(wb, '权威来源', SOURCE_HEADERS)
    for code in ('middle','province'):
        connection, _, _ = states[code]
        _, records = source_records(db, connection.id)
        current = {x.get('data_type'):x for x in records if isinstance(x,dict)}
        for dtype, name, status in DATA_CATALOG[code]:
            value = current.get(dtype,{})
            _append(source_ws, [code,dtype,name,status,value.get('source_system',''),
                                value.get('data_method',''),value.get('frequency',''),
                                value.get('history_range',''),value.get('owner',''),
                                value.get('responsibility',''),value.get('reference',''),
                                value.get('confirmed_on','')])
    _finish(source_ws, {'业务名称':24,'权威来源系统':32,'数据提供方式':24,'同步/调用频率':24,
                        '历史数据范围':26,'责任人':24,'责任边界说明':40,
                        '确认依据/文档编号':40,'双方确认日期':18})

    meta = wb.create_sheet('_系统')
    meta.append(['format', FORMAT])
    meta.append(['generated_at', utcnow().isoformat()])
    meta.append(['接口编码','connection_id','connection_version','draft_version',
                 'materials_version','credential_version','source_version','acceptance_version',
                 'identity_policy_version'])
    for code in CODES:
        connection, _, draft = states[code]
        materials = stored(db, 'materials', connection.id)
        credential = stored(db, 'credential', connection.id)
        sources = stored(db, 'sources', connection.id)
        acceptance = stored(db, 'acceptance', connection.id)
        identity_policy = stored(db, 'identity_policy', connection.id)
        meta.append([code,connection.id,connection.version,draft.version if draft else 0,
                     materials.version if materials else 0, credential.version if credential else 0,
                     sources.version if sources else 0, acceptance.version if acceptance else 0,
                     identity_policy.version if identity_policy else 0])
    meta.sheet_state = 'hidden'

    wb.properties.title = '跃科三大接口学校资料采集表'
    wb.properties.creator = '湖南跃科信息工程有限公司'
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def _inspect_xlsx(data):
    if not isinstance(data, (bytes, bytearray)) or not 1 <= len(data) <= 5*1024*1024:
        raise HTTPException(413, '学校资料Excel超过5MB或为空')
    if not data.startswith(b'PK'):
        raise HTTPException(422, '只接收系统生成的XLSX学校资料表')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 250 or sum(x.file_size for x in infos) > 25*1024*1024:
                raise HTTPException(422, 'Excel内部结构过大')
            lowered = [x.filename.lower() for x in infos]
            if any('vbaproject' in name or name.endswith('.bin') for name in lowered):
                raise HTTPException(422, '学校资料表不允许包含宏或二进制嵌入')
    except zipfile.BadZipFile:
        raise HTTPException(422, 'Excel文件损坏') from None


def _table(ws, headers):
    actual = [_text(cell.value, 200) for cell in ws[1]]
    missing = [name for name in headers if name not in actual]
    if missing:
        raise HTTPException(422, f'{ws.title}缺少列：'+'、'.join(missing))
    index = {name:actual.index(name)+1 for name in headers}
    rows = []
    for r in range(2, ws.max_row+1):
        row = {name:ws.cell(r,index[name]).value for name in headers}
        if any(_text(v) for v in row.values()):
            rows.append(row)
    return rows


def parse_intake_workbook(data):
    _inspect_xlsx(data)
    try:
        wb = load_workbook(BytesIO(data), data_only=False, read_only=False, keep_links=False)
    except Exception:
        raise HTTPException(422, 'Excel无法读取，请使用系统下载的学校资料模板') from None

    required = {'填写说明','接口基础','统一认证参数','统一认证身份','统一认证组织映射','统一认证角色确认',
                '数据中台接口','湘培接口','字段映射','资料登记','凭据证书元数据','权威来源','_系统'}
    if not required <= set(wb.sheetnames):
        wb.close()
        raise HTTPException(422, '学校资料表缺少系统要求的工作表')
    for ws in wb.worksheets:
        if ws.max_row > 600 or ws.max_column > 60:
            wb.close()
            raise HTTPException(422, f'{ws.title}行列数量超过模板范围')
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == 'f':
                    wb.close()
                    raise HTTPException(422, '学校资料表不允许公式，请填写固定值')
                _text(cell.value)

    meta = wb['_系统']
    if _text(meta['B1'].value) != FORMAT:
        wb.close()
        raise HTTPException(422, '学校资料表格式版本不匹配，请重新下载模板')
    expected = {}
    for r in range(4, meta.max_row+1):
        code = _text(meta.cell(r,1).value)
        if code:
            expected[code] = {
                'connection_id':_text(meta.cell(r,2).value),
                'connection_version':_int(meta.cell(r,3).value,0,1,10**9),
                'draft_version':_int(meta.cell(r,4).value,0,0,10**9),
                'materials_version':_int(meta.cell(r,5).value,0,0,10**9),
                'credential_version':_int(meta.cell(r,6).value,0,0,10**9),
                'source_version':_int(meta.cell(r,7).value,0,0,10**9),
                'acceptance_version':_int(meta.cell(r,8).value,0,0,10**9),
                'identity_policy_version':_int(meta.cell(r,9).value,0,0,10**9)
            }
    if set(expected) != set(CODES):
        wb.close()
        raise HTTPException(422, '学校资料表版本信息不完整')

    basic_rows = _table(wb['接口基础'], ['接口编码','对接项目','正式HTTPS总地址','责任单位/联系人',
        '认证/鉴权模式','服务器凭据引用','失败重试次数','同步频率(分钟)','学校/机构编码','备注'])
    basic = {}
    for row in basic_rows:
        code = _text(row['接口编码'])
        if code not in CODES or code in basic:
            wb.close()
            raise HTTPException(422, '接口基础必须且只能各有一行 sso / middle / province')
        if _text(row['服务器凭据引用']) != SECRET_REFS[code]:
            wb.close()
            raise HTTPException(422, f'{NAMES[code]}服务器凭据引用不可由Excel修改')
        basic[code] = row
    if set(basic) != set(CODES):
        wb.close()
        raise HTTPException(422, '接口基础三项不完整')

    sso_params = {}
    allowed_sso = {x[1] for x in SSO_PARAMS}
    for row in _table(wb['统一认证参数'], ['分类','参数编码','学校/实施填写','填写说明']):
        key = _text(row['参数编码'], 120)
        if not key or key in sso_params:
            wb.close()
            raise HTTPException(422, '统一认证参数编码为空或重复')
        if key not in allowed_sso:
            wb.close()
            raise HTTPException(422, '统一认证参数包含模板外字段')
        sso_params[key] = row['学校/实施填写']

    identities = {}
    for row in _table(wb['统一认证身份'], ['本平台身份字段','中文含义','学校返回字段路径','要求']):
        key = _text(row['本平台身份字段'], 64)
        if key not in dict(IDENTITY_FIELDS) or key in identities:
            wb.close()
            raise HTTPException(422, '统一认证身份字段无效或重复')
        identities[key] = _text(row['学校返回字段路径'], 120)

    org_map = {}
    for row in _table(wb['统一认证组织映射'], ['学校组织编码','本系统组织编码','说明']):
        upstream = _text(row['学校组织编码'], 120)
        local = _text(row['本系统组织编码'], 120)
        if not upstream and not local:
            continue
        if not upstream or not local or upstream in org_map:
            wb.close()
            raise HTTPException(422, '统一认证组织映射必须成对填写且学校编码不能重复')
        org_map[upstream] = local

    role_policy = []
    for row in _table(wb['统一认证角色确认'], ROLE_HEADERS):
        kind = _text(row['本地角色层级'], 20)
        if kind not in ROLE_SCOPE or any(x['kind'] == kind for x in role_policy):
            wb.close()
            raise HTTPException(422, '统一认证角色层级无效或重复')
        expected_label, expected_strategy, expected_scope = ROLE_SCOPE[kind]
        if (_text(row['中文层级']) != expected_label
                or _text(row['本地授权策略']) != expected_strategy
                or _text(row['数据范围说明']) != expected_scope):
            wb.close()
            raise HTTPException(422, '本地角色授权策略和数据范围为系统控制项，不允许由Excel修改')
        role_policy.append({
            'kind':kind,
            'external_roles':_text(row['学校外部角色/身份标识'],400),
            'authorization_strategy':expected_strategy,
            'scope':expected_scope,
            'confirmed_by':_text(row['学校确认人'],120),
            'confirmed_on':_date_text(row['确认日期']) or None
        })
    if {x['kind'] for x in role_policy} != set(ROLE_SCOPE):
        wb.close()
        raise HTTPException(422, '统一认证四级角色确认行不完整')

    middle_rows = _table(wb['数据中台接口'], MIDDLE_HEADERS)
    province_rows = _table(wb['湘培接口'], PROVINCE_HEADERS)

    def keyed(rows, code):
        allowed = {x[0] for x in DATA_CATALOG[code]}
        result = {}
        for row in rows:
            dtype = _text(row['数据类型'],64)
            if dtype not in allowed or dtype in result:
                raise HTTPException(422, f'{NAMES[code]}数据类型无效或重复：{dtype}')
            result[dtype] = row
        if set(result) != allowed:
            raise HTTPException(422, f'{NAMES[code]}接口规则行不完整')
        return result

    middle_rules = keyed(middle_rows,'middle')
    province_rules = keyed(province_rows,'province')

    mappings = {'middle':[], 'province':[]}
    allowed_transforms = {'text','int','decimal','date','bool'}
    for row in _table(wb['字段映射'], MAPPING_HEADERS):
        code = _text(row['接口编码'])
        dtype = _text(row['数据类型'],64)
        if code not in mappings or dtype not in {x[0] for x in DATA_CATALOG[code]}:
            wb.close()
            raise HTTPException(422, '字段映射接口编码或数据类型无效')
        transform = _text(row['转换方式']) or 'text'
        if transform not in allowed_transforms:
            wb.close()
            raise HTTPException(422, '字段映射转换方式无效')
        source, target = _text(row['源字段'],120), _text(row['目标字段'],120)
        if not source or not target:
            wb.close()
            raise HTTPException(422, '字段映射源字段和目标字段不能为空')
        default = row['默认值']
        mappings[code].append({'data_type':dtype,'source':source,'target':target,'transform':transform,
                               'required':_bool(row['必填']),'default':None if _text(default)=='' else default})

    materials = {code:[] for code in CODES}
    for row in _table(wb['资料登记'], MATERIAL_HEADERS):
        code, name = _text(row['接口编码']), _text(row['资料项目'])
        if code not in materials or name not in MATERIALS:
            wb.close()
            raise HTTPException(422, '资料登记接口或资料项目无效')
        materials[code].append({'name':name,'owner':_text(row['责任人'],120),
            'reference':_text(row['文档或证据编号'],400),'confirmed_on':_date_text(row['双方确认日期']) or None})
    for code in CODES:
        if {x['name'] for x in materials[code]} != set(MATERIALS) or len(materials[code]) != len(MATERIALS):
            wb.close()
            raise HTTPException(422, f'{NAMES[code]}资料登记必须保留六项固定清单')

    credentials = {}
    for row in _table(wb['凭据证书元数据'], CREDENTIAL_HEADERS):
        code = _text(row['接口编码'])
        if code not in CODES or code in credentials:
            wb.close()
            raise HTTPException(422, '凭据证书元数据接口编码无效或重复')
        credentials[code] = {
            'credential_type':_text(row['凭据类型']) or 'token',
            'managed_by':_text(row['安全保管方式']) or 'server_env',
            'owner':_text(row['轮换责任人'],120),
            'certificate_fingerprint_sha256':_text(row['证书SHA-256指纹'],64),
            'certificate_subject':_text(row['证书主题'],240),
            'certificate_issuer':_text(row['证书签发方'],240),
            'valid_from':_date_text(row['证书生效日']) or None,
            'valid_until':_date_text(row['证书到期日']) or None,
            'last_rotated_on':_date_text(row['上次轮换日期']) or None,
            'next_rotation_on':_date_text(row['下次计划轮换日期']) or None
        }
    if set(credentials) != set(CODES):
        wb.close()
        raise HTTPException(422, '三项凭据证书元数据不完整')

    sources = {code:[] for code in ('middle','province')}
    for row in _table(wb['权威来源'], SOURCE_HEADERS):
        code, dtype = _text(row['接口编码']), _text(row['数据类型'],64)
        if code not in sources or dtype not in {x[0] for x in DATA_CATALOG[code]}:
            wb.close()
            raise HTTPException(422, '权威来源接口编码或数据类型无效')
        sources[code].append({
            'data_type':dtype,'source_system':_text(row['权威来源系统'],160),
            'data_method':_text(row['数据提供方式'],160),
            'frequency':_text(row['同步/调用频率'],160),
            'history_range':_text(row['历史数据范围'],200),
            'owner':_text(row['责任人'],120),
            'responsibility':_text(row['责任边界说明'],400),
            'reference':_text(row['确认依据/文档编号'],400),
            'confirmed_on':_date_text(row['双方确认日期']) or None
        })
    for code in sources:
        if {x['data_type'] for x in sources[code]} != {x[0] for x in DATA_CATALOG[code]}:
            wb.close()
            raise HTTPException(422, f'{NAMES[code]}权威来源行不完整')

    wb.close()
    return {
        'format':FORMAT,'expected':expected,'basic':basic,'sso_params':sso_params,
        'identities':identities,'org_map':org_map,'role_policy':role_policy,
        'rules':{'middle':middle_rules,'province':province_rules},
        'mappings':mappings,'materials':materials,'credentials':credentials,'sources':sources
    }


def _raw_rule(value, current):
    raw = _text(value)
    if not raw:
        return deepcopy(current or {})
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        raise HTTPException(422, '完整规则JSON格式不正确') from None
    if not isinstance(parsed, dict):
        raise HTTPException(422, '完整规则JSON必须是对象')
    return parsed


def _overlay_basic(config, adapter, code, row):
    config = deepcopy(config or {})
    config['endpoint'] = _text(row['正式HTTPS总地址'], 500)
    config['owner'] = _text(row['责任单位/联系人'], 160)
    config['auth_mode'] = _text(row['认证/鉴权模式'], 40)
    config['secret_env'] = SECRET_REFS[code]
    config['retry_count'] = _int(row['失败重试次数'], 3, 3, 10)
    config['frequency_minutes'] = _int(row['同步频率(分钟)'], 30, 1, 10080)
    if code != 'sso':
        public = deepcopy(adapter.get('public_parameters') or {})
        public['institution_code'] = _text(row['学校/机构编码'], 160)
        adapter['public_parameters'] = public
    config['adapter'] = adapter
    return config


def _overlay_sso(adapter, parsed):
    adapter = deepcopy(adapter)
    params = parsed['sso_params']
    list_keys = {'active_true_values','active_false_values','probe.success_values','verify.success_values'}
    for _, key, _, _ in SSO_PARAMS:
        value = params.get(key, '')
        if key in list_keys:
            _nested_set(adapter,key,_parse_values(value))
        elif key == 'active_default':
            raw = _text(value).lower()
            if raw == '':
                adapter.pop('active_default', None)
            elif raw in {'true','1','是','启用'}:
                adapter['active_default'] = True
            elif raw in {'false','0','否','停用'}:
                adapter['active_default'] = False
            else:
                raise HTTPException(422, 'active_default只能填写true、false或留空')
        else:
            raw = _text(value,500)
            if raw == '':
                _nested_delete(adapter,key)
            else:
                _nested_set(adapter,key,raw)
    adapter['identity_fields'] = dict(parsed['identities'])
    adapter['org_map'] = dict(parsed['org_map'])
    return adapter


def _overlay_middle_rule(current, row):
    rule = _raw_rule(row['完整规则JSON'], current)
    request = deepcopy(rule.get('request') or {})
    request.setdefault('headers',{})
    request.setdefault('query',{})
    request['method'] = _text(row['请求方式']) or 'GET'
    request['path'] = _text(row['接口相对路径'],500)
    request['auth'] = _apply_auth(request.get('auth'), [row[name] for name in AUTH_HEADERS])
    _set_binding(request,_text(row['机构编码位置']) or 'query',row['机构编码字段'],TOKEN_INSTITUTION)
    for token, target_header, name_header in [
        (TOKEN_CURSOR,'游标位置','游标字段'),(TOKEN_PAGE,'页码位置','页码字段'),
        (TOKEN_OFFSET,'偏移位置','偏移字段'),(TOKEN_LIMIT,'每页数量位置','每页数量字段'),
        (TOKEN_MODE,'同步模式位置','同步模式字段')
    ]:
        name = _text(row[name_header],120)
        target = _text(row[target_header]) or 'query'
        if name:
            _set_binding(request,target,name,token)
        else:
            _clear_binding(request,token)
    paging = deepcopy(rule.get('pagination') or {})
    paging.update(
        mode=_text(row['翻页方式']) or 'cursor',
        page_size=_int(row['每页数量'],1000,1,5000),
        max_pages=_int(row['最大页数'],20,1,10000),
        start_page=_int(row['起始页'],1,0,10**6),
        end_condition=_text(row['结束判定'],64)
    )
    response = deepcopy(rule.get('response') or {})
    response.update(
        success_path=_text(row['成功标志路径'],240),
        success_values=_parse_values(row['成功允许值']),
        records_path=_text(row['数据列表路径'],240),
        next_cursor_path=_text(row['下一游标路径'],240),
        has_more_path=_text(row['是否还有更多路径'],240),
        total_path=_text(row['总数路径'],240),
        watermark_path=_text(row['增量水位路径'],240)
    )
    rule.update(request=request,pagination=paging,response=response)
    return rule


def _overlay_province_rule(current, row):
    rule = _raw_rule(row['完整规则JSON'], current)
    request = deepcopy(rule.get('request') or {})
    request.setdefault('headers',{})
    request.setdefault('query',{})
    request.setdefault('body',{})
    request['method'] = _text(row['请求方式']) or 'POST'
    request['path'] = _text(row['接口相对路径'],500)
    request['auth'] = _apply_auth(request.get('auth'), [row[name] for name in AUTH_HEADERS])
    _set_binding(request,_text(row['机构编码位置']) or 'query',row['机构编码字段'],TOKEN_INSTITUTION)
    _set_binding(request,'body',row['业务数据字段'],TOKEN_PAYLOAD)
    response = deepcopy(rule.get('response') or {})
    response.update(
        accepted_path=_text(row['接收成功路径'],240),
        accepted_values=_parse_values(row['接收成功值']),
        code_path=_text(row['返回代码路径'],240),
        receipt_id_path=_text(row['回执编号路径'],240),
        error_codes=[x.strip() for x in re.split(r'[，,\n]',_text(row['失败代码'])) if x.strip()]
    )
    rule.update(request=request,response=response)
    return rule


def overlay_intake(item, parsed, code):
    base = templates(DATA_CATALOG)[code]
    adapter = merge_template(base, (item.config or {}).get('adapter') or {})
    config = _overlay_basic(item.config, adapter, code, parsed['basic'][code])
    adapter = config['adapter']
    if code == 'sso':
        adapter = _overlay_sso(adapter, parsed)
        config['adapter'] = adapter
        identity = adapter.get('identity_fields') or {}
        mappings = [
            {'data_type':'identity','source':source,'target':key,
             'transform':'bool' if key == 'active' else 'text',
             'required':key in {'name','active'},'default':None}
            for key, source in identity.items() if source
        ]
    else:
        current_rules = adapter.get('data_types') or {}
        next_rules = {}
        for dtype, _, _ in DATA_CATALOG[code]:
            row = parsed['rules'][code][dtype]
            current = current_rules.get(dtype) or {}
            next_rules[dtype] = (
                _overlay_middle_rule(current,row) if code == 'middle'
                else _overlay_province_rule(current,row)
            )
        adapter['data_types'] = next_rules
        config['adapter'] = adapter
        mappings = deepcopy(parsed['mappings'][code])

    return {
        'config':config,
        'mappings':mappings,
        'materials':deepcopy(parsed['materials'][code]),
        'credential':deepcopy(parsed['credentials'][code]),
        'sources':deepcopy(parsed['sources'].get(code,[])),
        'role_policy':deepcopy(parsed['role_policy']) if code == 'sso' else []
    }
