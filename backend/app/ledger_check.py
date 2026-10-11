"""学院《社会培训人次统计表》读取与口径体检。

学院现在用 Excel 报人次，2025 版有“当量系数”列（人次=人数×课时×当量÷8），
2026 版按模块式或人天式填报。这里只读表、按制度复算并标出问题，不替学校核定：
推荐值按“系统推荐口径（待学校书面确认）”给出，核定值由继续教育学院填写。
"""
from __future__ import annotations
import io, re, zipfile
from decimal import Decimal, ROUND_HALF_UP
from fastapi import HTTPException

D = Decimal
MAX_BYTES = 32 * 1024 * 1024        # 学院表里常嵌照片，2026 年素质教育学院原表约 15 MB
MAX_UNCOMPRESSED = 300 * 1024 * 1024
MAX_ROWS = 500
ONLINE_WORDS = ('浏览', '观看', '在线', '点击', '阅读量', '播放')
ONLINE_HINTS = ('线上', '直播', '在线', '客户端', '云端', '网络课')
ROSTER_WORDS = ('名单', '签到', '花名册', '名册', '签名')
LEVELS = {'严重': 3, '注意': 2, '提示': 1}


def q2(value) -> str:
    return str(D(str(value)).quantize(D('0.01'), rounding=ROUND_HALF_UP)).rstrip('0').rstrip('.')


def _guard_xlsx(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(x.file_size for x in infos) > MAX_UNCOMPRESSED: raise ValueError
            for x in infos:
                name = x.filename.lower()
                if '..' in x.filename.split('/') or 'vbaproject' in name or '/externallinks/' in name: raise ValueError
    except (zipfile.BadZipFile, ValueError):
        raise HTTPException(422, 'Excel 文件无法安全读取，请另存为 .xlsx 后重新上传')


def _cells_xlsx(data: bytes):
    import openpyxl
    _guard_xlsx(data)
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True).worksheets[0]
    out = []
    for vrow, frow in zip(values.iter_rows(max_row=MAX_ROWS + 20), formulas.iter_rows(max_row=MAX_ROWS + 20)):
        out.append([(v.value, f.value) for v, f in zip(vrow, frow)])
    return out


def _cells_xls(data: bytes):
    try:
        import xlrd
    except ImportError:
        raise HTTPException(422, '服务器暂不能读取 .xls 旧格式，请在 Excel 里另存为 .xlsx 后上传')
    try:
        sheet = xlrd.open_workbook(file_contents=data, on_demand=True).sheet_by_index(0)
    except Exception:
        raise HTTPException(422, '.xls 文件无法读取，请另存为 .xlsx 后重新上传')
    return [[(sheet.cell_value(r, c), None) for c in range(sheet.ncols)] for r in range(min(sheet.nrows, MAX_ROWS + 20))]


def _text(value) -> str:
    if value is None: return ''
    if isinstance(value, float) and value.is_integer(): value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def _number(value):
    """返回 (数值或 None, 原文)。'204人次'、'40人天'、'在线浏览19000左右' 都取出数字。"""
    raw = _text(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool): return D(str(value)), raw
    match = re.search(r'\d+(?:\.\d+)?', raw.replace(',', ''))
    return (D(match.group()) if match else None), raw


_ROW_FORMULA = re.compile(r'^=\(?C(\d+)\*D\1\*E\1\)?(/8)?$')
_PRODUCT = re.compile(r'^=\d+(?:\.\d+)?(?:\*\d+(?:\.\d+)?)*$')


def _formula_value(formula, people, hours, coeff):
    """表格没保存计算结果时，按单元格公式自己算（只认学校模板里的乘法公式）。"""
    if not isinstance(formula, str): return None
    text = formula.replace(' ', '').upper()
    if _PRODUCT.match(text):
        result = D(1)
        for part in text[1:].split('*'): result *= D(part)
        return result
    match = _ROW_FORMULA.match(text)
    if match and None not in (people, hours, coeff):
        value = people * hours * coeff
        return value / 8 if match.group(2) else value
    return None


def parse_allocations(text: str):
    if '%' in text or '％' in text: return []
    return [(name, D(num)) for name, num in re.findall(r'([一-龥]{2,6})\s*[:：]?\s*(\d+(?:\.\d+)?)', text)]


def _header(rows):
    for index, row in enumerate(rows[:15]):
        names = [_text(v) for v, _ in row]
        if any(n in ('培训名称', '培训项目名称') for n in names) and any('人数' in n for n in names):
            return index, names
    raise HTTPException(422, '没有找到“培训名称 / 培训人数”表头，请使用学校下发的《社会培训人次统计表》')


def _col(names, *keys):
    for i, n in enumerate(names):
        if any(k in n for k in keys): return i
    return None


def parse(data: bytes, filename: str) -> dict:
    if not data: raise HTTPException(422, '文件是空的')
    if len(data) > MAX_BYTES: raise HTTPException(413, '文件超过 32 MB，请删掉表里的照片后再上传（照片放在资料文件夹即可）')
    lower = (filename or '').lower()
    if lower.endswith('.xlsx'): rows = _cells_xlsx(data)
    elif lower.endswith('.xls'): rows = _cells_xls(data)
    else: raise HTTPException(422, '只支持 .xlsx 或 .xls 格式的学院统计表')
    head, names = _header(rows)
    title = ' '.join(_text(v) for row in rows[:head] for v, _ in row if _text(v))
    meta = {'title': title}
    for key, label in (('college', r'学院/部门名称[:：]\s*([^\s:：]+?)(?=\s|部门负责人|$)'), ('leader', r'部门负责人[:：]\s*(\S+?)(?=\s|经办人|$)'),
                       ('handler', r'经办人[:：]\s*(\S+?)(?=\s|日期|$)'), ('date', r'日期[:：]\s*(\S+)')):
        match = re.search(label, title)
        meta[key] = match.group(1) if match and not re.search('负责人|经办人|日期', match.group(1)) else ''
    if not meta['college']:
        match = re.search(r'([一-龥]{1,12}?学院)', title.replace(' ', '').split('学院/部门名称')[0])
        meta['college'] = match.group(1) if match else ''
    year = re.search(r'(20\d\d)\s*年', title)
    meta['year'] = int(year.group(1)) if year else None
    c = {'seq': _col(names, '序号'), 'name': _col(names, '培训名称', '培训项目名称'), 'people': _col(names, '人数'),
         'hours': _col(names, '课时'), 'coeff': _col(names, '当量'), 'claimed': _col(names, '人次数'),
         'alloc': _col(names, '分配', '模块'), 'materials': _col(names, '资料'), 'remark': _col(names, '备注')}
    if c['claimed'] is None or c['hours'] is None: raise HTTPException(422, '表头缺少“培训课时”或“人次数”列')
    template = '2025' if c['coeff'] is not None else '2026'
    items, sheet_total, last_name = [], None, ''

    def cell(row, key, formula=False):
        i = c[key]
        if i is None or i >= len(row): return None
        return row[i][1 if formula else 0]

    for offset, row in enumerate(rows[head + 1:], start=head + 2):
        first = _text(row[0][0]) if row else ''
        if '资料要求' in first: break
        if first in ('总结', '合计', '总计') or (not _text(cell(row, 'name')) and not _text(cell(row, 'people')) and _text(cell(row, 'claimed'))):
            value, _ = _number(cell(row, 'claimed'))
            if value: sheet_total = value
            if first in ('总结', '合计', '总计'): break
            continue
        name, people_raw = _text(cell(row, 'name')), _text(cell(row, 'people'))
        if not name and not people_raw: continue
        people, _ = _number(cell(row, 'people'))
        hours, _ = _number(cell(row, 'hours'))
        if hours is None: hours = _formula_value(cell(row, 'hours', True), None, None, None)
        coeff = None
        if template == '2025':
            coeff, _ = _number(cell(row, 'coeff'))
            if coeff is None: coeff = D(1)
        claimed, claimed_raw = _number(cell(row, 'claimed'))
        if claimed is None: claimed = _formula_value(cell(row, 'claimed', True), people, hours, coeff)
        materials = _text(cell(row, 'materials'))
        if 'DISPIMG' in materials.upper(): materials = '（表格里嵌入了图片）'
        items.append({'row_no': offset, 'name': name or (last_name + '（续行）'), 'people_raw': people_raw,
                      'people': people, 'hours': hours, 'coeff': coeff, 'claimed': claimed, 'claimed_raw': claimed_raw,
                      'allocation': _text(cell(row, 'alloc')), 'materials': materials, 'remark': _text(cell(row, 'remark'))})
        last_name = name or last_name
        if len(items) >= MAX_ROWS: raise HTTPException(413, f'单张表最多 {MAX_ROWS} 行')
    if not items: raise HTTPException(422, '表里没有找到培训项目行')
    return {'meta': meta, 'template': template, 'sheet_total': sheet_total, 'items': items}


def check_row(item: dict, template: str) -> dict:
    """按制度复算一行，返回推荐人次和问题清单（大白话）。"""
    p, h, k, claimed = item['people'], item['hours'], item['coeff'], item['claimed']
    issues, mode, basis = [], '', ''

    def add(level, code, text): issues.append({'level': level, 'code': code, 'text': text})

    if p is None or h is None or claimed is None:
        add('严重', 'MISSING', '培训人数、课时或人次数没有填数字，无法核算')
        return {**_public(item), 'person_days': None, 'mode': '', 'recommended': '0', 'basis': '缺少数字，推荐暂按 0', 'issues': issues}
    person_days = p * h / 8
    online_text = any(w in item['people_raw'] for w in ONLINE_WORDS)
    if online_text:
        add('注意' if template == '2025' else '严重', 'ONLINE_VIEWS', f'人数填为线上浏览/观看量（{item["people_raw"]}），请核对实际参训依据及适用年度规则；浏览量不能直接证明有效人次')
    elif p >= 10000:
        add('提示', 'LARGE_COUNT', '填报人数较大，请核对名单、批次与重复人员；不能仅凭人数判断为线上浏览量')
    elif any(w in item['name'].replace('直播销售', '') for w in ONLINE_HINTS):
        add('提示', 'ONLINE', '线上项目：要有实名学习记录和结业考核才能算人次，请核对资料')
    if template == '2025':
        standard = p * h * k / 8
        mode, recommended, basis = '2025 旧口径', standard, '人数×课时×当量÷8（2025 年统计说明第 5 条）'
        if abs(claimed - standard) > D('0.5'):
            if abs(claimed - p * h * k) <= D('0.5') and h * k != 1:
                add('严重', 'MISSING_DIV8', f'填报 {q2(claimed)} = 人数×课时×当量，漏了“÷8”，按公式应为 {q2(standard)}')
            else:
                add('注意', 'FORMULA_MISMATCH', f'填报 {q2(claimed)} 和公式算出的 {q2(standard)} 对不上')
    else:
        if '人天' in item['claimed_raw'] or abs(claimed - person_days) <= D('0.5'):
            mode, recommended, basis = '人天式', person_days, '人天式：人数×课时÷8'
        elif claimed == p:
            mode, recommended, basis = '模块式（1个模块）', p, '模块式：1 个模块，1 名学员计 1 人次'
        elif p and claimed == p * h and h > 1:
            add('注意', 'PEOPLE_X_HOURS', f'填报 {q2(claimed)} = 人数×课时，也可能对应多个模块；请核对备案方式及模块明细。人天式试算 {q2(person_days)}，单模块试算 {q2(p)}')
            mode, recommended = '口径不清', claimed
            basis = '暂列学院填报值，待确认核算方式与模块明细，不据数值关系自动核减'
        elif p and claimed % p == 0 and 1 < claimed / p <= max(h, 1):
            modules = int(claimed / p)
            mode, recommended, basis = f'模块式（{modules}个模块）', claimed, f'模块式：{modules} 个模块 × {q2(p)} 人'
            add('提示', 'MODULES', f'按 {modules} 个模块填报，备案时要写明每个模块和授课老师')
        else:
            add('注意', 'MODE_UNCLEAR', f'填报 {q2(claimed)} 与当前人数的人天式试算 {q2(person_days)}、单模块试算 {q2(p)} 不一致，请核对各单元人数及备案方式')
            mode, recommended = '口径不清', claimed
            basis = '暂列学院填报值，待核算单元与人数依据确认，不自动选用核算方式'
    if recommended > claimed:
        recommended, basis = claimed, basis + '；按规则算出的数比学院填报多，推荐值不超过学院填报'
    if online_text and template != '2025':
        recommended, basis = D(0), '线上浏览量不计人次（推荐口径）'
    allocations = parse_allocations(item['allocation'])
    if allocations:
        total = sum(v for _, v in allocations)
        if total > claimed + D('0.5'):
            add('严重' if template == '2025' else '注意', 'ALLOC_OVER', f'教师填报合计 {q2(total)} 超过项目填报 {q2(claimed)}，请区分原分配数与多师同堂参与人次；教师分别核算不能重复累计学院人次')
        elif total < claimed - D('0.5'):
            add('提示', 'ALLOC_UNDER', f'分给老师的人次合计 {q2(total)}，少于本项目总数 {q2(claimed)}')
    materials = item['materials']
    if materials == '（表格里嵌入了图片）':
        add('提示', 'MATERIAL_IMAGE', '资料情况是图片，请人工核对有没有学员名单或签到表')
    elif not (any(w in materials for w in ROSTER_WORDS) or re.search(r'(?<!\d)3(?!\d)', materials)):
        add('注意', 'NO_ROSTER', '资料栏未列名单或签到，请核验原件；2025旧口径可核对能基本佐证人数的材料，新积分认定须核对实名、出勤与结业依据')
    return {**_public(item), 'person_days': q2(person_days), 'mode': mode, 'recommended': q2(recommended), 'basis': basis,
            'needs_mode_confirmation': template != '2025' and '人天' not in item['claimed_raw'],
            'trial_candidates': {'person_day': q2(person_days), 'single_module': q2(p)} if template != '2025' else {},
            'allocations': [{'name': n, 'value': q2(v)} for n, v in allocations], 'issues': issues}


def _public(item):
    return {k: (q2(v) if isinstance(v, Decimal) else v) for k, v in item.items()}


def check(parsed: dict) -> dict:
    rows = [check_row(item, parsed['template']) for item in parsed['items']]
    claimed = sum((D(r['claimed']) for r in rows if r['claimed'] is not None), D(0))
    recommended = sum((D(r['recommended']) for r in rows), D(0))
    summary = {'rows': len(rows), 'claimed_total': q2(claimed), 'recommended_total': q2(recommended),
               'serious': sum(any(i['level'] == '严重' for i in r['issues']) for r in rows),
               'attention': sum(any(i['level'] == '注意' for i in r['issues']) for r in rows)}
    notes = []
    if parsed['sheet_total'] is not None and abs(parsed['sheet_total'] - claimed) > D('0.5'):
        notes.append(f'表底“总结”写的是 {q2(parsed["sheet_total"])}，逐行加起来是 {q2(claimed)}')
    return {'meta': parsed['meta'], 'template': parsed['template'],
            'template_label': '2025 版（人数×课时×当量÷8）' if parsed['template'] == '2025' else '2026 版（模块式 / 人天式）',
            'sheet_total': None if parsed['sheet_total'] is None else q2(parsed['sheet_total']),
            'summary': summary, 'notes': notes, 'rows': rows}


def totals(rows: list) -> dict:
    """核定后汇总：已核定行用核定值，未核定行仍按推荐值，单独列出未核定行数。"""
    approved = [r for r in rows if r.get('approved') is not None]
    final = sum((D(r['approved']) if r.get('approved') is not None else D(r['recommended']) for r in rows), D(0))
    return {'claimed_total': q2(sum((D(r['claimed']) for r in rows if r['claimed'] is not None), D(0))),
            'recommended_total': q2(sum((D(r['recommended']) for r in rows), D(0))),
            'approved_rows': len(approved), 'pending_rows': len(rows) - len(approved), 'final_total': q2(final),
            'state': '已核定' if len(approved) == len(rows) else ('部分核定' if approved else '待核定')}
