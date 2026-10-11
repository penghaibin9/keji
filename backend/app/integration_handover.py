"""Human-readable integration handover workbook.

The workbook is derived only from the already-sanitized handover payload.
It never reads environment variables, credential files or upstream raw bodies.
"""
from __future__ import annotations
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def _sheet(workbook, title, headers):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    return ws


def _finish(ws):
    # Handover values contain school-provided text. Prevent spreadsheet formula
    # execution when the generated workbook is opened by implementation staff.
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.lstrip(' \t\r\n').startswith(('=','+','-','@')):
                cell.value = "'" + cell.value
    for column in range(1, ws.max_column + 1):
        width = 12
        for cell in ws.iter_cols(min_col=column, max_col=column):
            for item in cell:
                width = max(width, min(50, len(str(item.value or '')) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

def build_handover_workbook(data):
    workspace = data.get('workspace') or {}
    workbook = Workbook()
    workbook.remove(workbook.active)

    status = _sheet(workbook, '状态总览', ['项目', '当前状态'])
    labels = {'materials':'资料齐备','configuration':'配置有效','tests':'测试结果',
              'published':'配置发布','health':'运行健康','acceptance':'学校验收'}
    for key, label in labels.items():
        status.append([label, (workspace.get('statuses') or {}).get(key, '未记录')])
    callbacks = workspace.get('callbacks') or {}
    for key, label in [('login','正式登录回调'),('test_login','实施试登录回调'),('logout_return','退出后返回地址')]:
        if callbacks.get(key):
            status.append([label, callbacks[key]])

    materials = _sheet(workbook, '资料清单', ['资料项目','责任人','文档或证据编号','确认日期'])
    material_data = workspace.get('materials') or {}
    current = {x.get('name'):x for x in material_data.get('items', []) if isinstance(x, dict)}
    for name in material_data.get('required', []):
        item = current.get(name, {})
        materials.append([name, item.get('owner',''), item.get('reference',''), item.get('confirmed_on','')])

    mappings = _sheet(workbook, '字段映射', ['数据类型','源字段','目标字段','转换规则','必填','默认值'])
    for item in (data.get('draft') or data.get('active') or {}).get('mappings', []):
        if not isinstance(item, dict):
            continue
        mappings.append([item.get('data_type',''), item.get('source',''), item.get('target',''),
                         item.get('transform',''), '是' if item.get('required') else '否',
                         '' if item.get('default') is None else str(item.get('default'))])

    credential = _sheet(workbook, '凭据证书元数据', ['项目','登记值'])
    safe = (workspace.get('credential') or {}).get('metadata') or {}
    names = {'credential_type':'凭据类型','managed_by':'安全保管方式','owner':'轮换责任人',
             'certificate_fingerprint_sha256':'证书SHA-256指纹','certificate_subject':'证书主题',
             'certificate_issuer':'证书签发方','valid_from':'证书生效日','valid_until':'证书到期日',
             'last_rotated_on':'上次轮换日期','next_rotation_on':'下次计划轮换日期',
             'credential_reference':'服务器凭据引用'}
    for key, label in names.items():
        credential.append([label, safe.get(key,'')])
    credential.append(['真实密码/令牌/私钥', '不进入本工作簿；按学校安全渠道交接'])

    roles = _sheet(workbook, '统一认证角色确认', ['本地角色层级','学校外部角色/身份标识','本地授权策略','数据范围','学校确认人','确认日期'])
    for item in (workspace.get('identity_policy') or {}).get('record', {}).get('items', []):
        if not isinstance(item, dict):
            continue
        roles.append([item.get('kind',''), item.get('external_roles',''),
                      item.get('authorization_strategy',''), item.get('scope',''),
                      item.get('confirmed_by',''), item.get('confirmed_on','')])

    sources = _sheet(workbook, '权威来源台账',
        ['数据类型','权威来源系统','数据提供方式','同步/调用频率','历史数据范围','责任人','责任边界','确认依据/文档编号','确认日期'])
    for item in (workspace.get('source_registry') or {}).get('items', []):
        if not isinstance(item, dict):
            continue
        sources.append([item.get('data_type',''), item.get('source_system',''), item.get('data_method',''),
                        item.get('frequency',''), item.get('history_range',''), item.get('owner',''),
                        item.get('responsibility',''), item.get('reference',''), item.get('confirmed_on','')])

    acceptance = _sheet(workbook, '学校验收台账', ['验收项目','结果','证据编号','学校联系人','核对日期'])
    acceptance_data = workspace.get('acceptance') or {}
    record = acceptance_data.get('record') or {}
    current = {x.get('code'):x for x in record.get('items', []) if isinstance(x, dict)}
    status_names = {'PASSED':'通过','FAILED':'未通过','PENDING':'待确认'}
    for required in (acceptance_data.get('summary') or {}).get('requirements', []):
        item = current.get(required.get('code'), {})
        acceptance.append([required.get('label',''), status_names.get(item.get('status'),'待确认'),
                           item.get('evidence_ref',''), item.get('school_contact',''), item.get('checked_on','')])
    acceptance.append([])
    acceptance.append(['当前验收状态', (acceptance_data.get('summary') or {}).get('state','待学校签认')])
    acceptance.append(['最终签认人', record.get('signed_by','')])
    acceptance.append(['签认日期', record.get('signed_on','')])
    acceptance.append(['签认依据', record.get('signed_reference','')])

    for ws in workbook.worksheets:
        _finish(ws)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
