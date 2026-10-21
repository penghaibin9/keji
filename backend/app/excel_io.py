"""Production-oriented XLSX import/export and open-format handover helpers.

No import commits during preview.  XLSX parsing rejects formulas, active content,
external links and oversized ZIP payloads.  Every confirm revalidates against the
live authorization/database state and commits the whole batch atomically.
"""
from __future__ import annotations
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from io import BytesIO, StringIO
from urllib.parse import quote
import csv, json, os, re, secrets, stat, zipfile
from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import select, func
from sqlalchemy.inspection import inspect as sa_inspect
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .core import serial, digest, audit, workflow, Actor
from .models import *
from .schemas import RecordIn, QualificationIn, ParticipationIn
from .services import insert_record, qualification_revision
from .passwords import hash_password
from .read_queries import record_query, assessment_query, authorized_org_descendants, RECORD_FILTERS, ASSESSMENT_FILTERS, audit_query
from .org_assessment import OrganizationAssessment, OrganizationAssessmentRule, OrganizationAssessmentMetricEvidence, OrganizationRuleTrial
from .settings_catalog import business_dictionary,import_row_limit
from .integrations import SyncItem,SyncCallLog,ConnectionCheck,DATA_CATALOG
from .delivery import configuration_snapshot,database_dictionary
from .reference_data import TeacherParticipation,TrainingProjectMaster,TrainingCourseMaster,TrainingClassMaster,TeacherHrProfile,TeacherCourseTerm,SchoolCourse,InboundBatch,InboundRow,ensure_reference_from_training_record,record_reference_errors
from .inbound_sync import IntegrationSchedule
from .social_models import SocialProject,SocialUnit,SocialPolicy,SocialClaim,SocialAnnual,SocialResource

MAX_ROWS=5000
EXPORT_FORMAT_VERSION=2
MAX_XLSX_BYTES=10*1024*1024
MAX_UNCOMPRESSED=50*1024*1024
MAX_ZIP_ENTRIES=1000
DANGEROUS_PREFIXES=('=','+','-','@')

TEMPLATES={
 'records':[
  ('来源唯一编号','external_key'),('教师工号','teacher_staff_no'),('业务日期','event_date'),('项目编码','project_code'),('项目名称','project'),
  ('课程编码','course_code'),('课程名称','course'),('班级编码','class_code'),('班级名称','class_name'),('授课学时','hours'),('课程人次','participants'),
  ('归属培训收入','income'),('数据来源','source')],
 'qualifications':[
  ('教师工号','teacher_staff_no'),('资质类别','category'),('证书名称','name'),('证书编号','certificate_no'),('发证机构','issuer'),('取得日期','obtained'),('有效期至','expires')],
 'teachers':[
  ('登录账号','username'),('姓名','name'),('教师工号','staff_no'),('组织编码','org_code'),('职称/岗位','title')],
 'participations':[
  ('来源唯一编号','external_key'),('教师工号','teacher_staff_no'),('参训日期','event_date'),('项目编码','project_code'),('项目名称','project'),
  ('课程编码','course_code'),('课程名称','course'),('参训学时','hours'),('数据来源','source')],
}

TEMPLATE_NAMES={'records':'培训授课数据','participations':'教师参训记录','qualifications':'教师资质草稿','teachers':'教师账号'}
TEMPLATE_COLUMN_HELP={
 '来源唯一编号':'必填，按文本填写。同类记录中必须唯一，保留前导零；重复编号会被拒绝，本模板不覆盖已有记录。',
 '教师工号':'必填，按文本填写学校账号中的完整工号，保留前导零。教师须为当前授权范围内的启用账号，不能用姓名代替工号。',
 '业务日期':'必填，填写实际授课业务发生日期，使用“2026-10-19”格式，不填写日期说明文字。',
 '项目编码':'必填，按文本填写并保留前导零，与项目名称对应。仅限字母、数字、下划线或短横线，最多64个字符；优先核对已有项目编码。',
 '项目名称':'必填，填写实际培训项目名称，与项目编码对应，最多160个字符。导入台账不代表完成项目申报、备案或结项。',
 '课程编码':'必填，按文本填写并保留前导零，与课程名称及项目对应。仅限字母、数字、下划线或短横线，最多64个字符；不能使用已绑定其他项目的课程编码。',
 '课程名称':'必填，填写实际培训课程名称，与课程编码对应，最多160个字符。',
 '班级编码':'必填，按文本填写并保留前导零。仅限字母、数字、下划线或短横线，最多64个字符；同一编码须属于同一项目和课程。',
 '班级名称':'必填，填写实际培训班级名称，与班级编码对应，最多160个字符。',
 '授课学时':'必填，填写教师实际授课学时，使用非负数字，最多两位小数，可填0。单位为学时，不在单元格内附加单位文字。',
 '课程人次':'必填，填写非负整数，单位为人次。按学校确认的课程人次口径核对，不能直接当作项目去重人数，也不要重复统计同一业务。',
 '归属培训收入':'必填，单位为元，最多两位小数。填写归属本条教师授课明细的培训收入，须有核对依据；不能用缺失值猜为0，也不能重复填入项目总收入。',
 '数据来源':'选填，可留空。正式记录统一标记为“表格导入”；如填写，须使用学校当前业务字典允许的来源，原填写值保留在导入预览中。',
 '参训日期':'必填，填写教师作为学员参加培训的实际日期，使用“2026-10-19”格式。不能用本表登记教师授课。',
 '参训学时':'必填，填写教师作为学员参加培训的学时，使用非负数字，最多两位小数，可填0。不计入授课收入或授课人次。',
 '资质类别':'必填，使用学校“资质字典与到期提醒”中当前配置的类别，不得自行编写新的类别名称。',
 '证书名称':'必填，与证书及证明材料上的名称一致，最多160个字符。导入后为草稿，仍需教师本人补充证明材料并提交审核。',
 '证书编号':'必填，按文本填写完整证书编号，保留前导零，最多128个字符。同一教师的相同证书编号不能重复导入。',
 '发证机构':'必填，填写证书上的发证机构完整名称，最多160个字符。',
 '取得日期':'必填，填写证书实际取得日期，使用“2026-10-19”格式，并与证明材料一致。',
 '有效期至':'选填，使用“2026-10-19”格式，不能早于取得日期。仅长期有效的证书可留空；留空时系统按长期有效处理。',
}
TEMPLATE_TEXT_FIELDS={'external_key','teacher_staff_no','project_code','course_code','class_code','certificate_no'}


def _plain(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,Decimal):return str(value)
    return value


def _safe_text(value,*,limit=200):
    text=str(value if value is not None else '').strip()
    if len(text)>limit or any(ord(c)<32 and c not in '\t\n\r' for c in text):raise ValueError('文本格式或长度不正确')
    if text.startswith(DANGEROUS_PREFIXES):raise ValueError('文本不能以公式控制字符开头')
    return text


def _safe_excel_value(value):
    if isinstance(value,Decimal):return value
    value=value.isoformat() if isinstance(value,datetime) else _plain(value)
    if isinstance(value,str) and value.lstrip(' \t\r\n').startswith(DANGEROUS_PREFIXES):return "'"+value
    return value


def _date(value,required=True):
    if value in (None,''):
        if required:raise ValueError('日期必填')
        return None
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    text=_safe_text(value,limit=20)
    if match:=re.fullmatch(r'(\d{4})/(\d{1,2})/(\d{1,2})',text):
        try:return date(*map(int,match.groups())).isoformat()
        except ValueError:raise ValueError('日期不存在，请核对年月日') from None
    try:return date.fromisoformat(text).isoformat()
    except ValueError:raise ValueError('日期格式应为 YYYY-MM-DD 或 YYYY/M/D') from None


def _decimal(value,label,*,minimum=None):
    if value in (None,''):raise ValueError(f'{label}必填')
    try:v=Decimal(str(value))
    except (InvalidOperation,ValueError):raise ValueError(f'{label}必须是数字') from None
    if not v.is_finite() or (minimum is not None and v<Decimal(str(minimum))):raise ValueError(f'{label}数值无效')
    return str(v)


def _integer(value,label,*,minimum=0):
    if isinstance(value,bool):raise ValueError(f'{label}必须是整数')
    try:
        dec=Decimal(str(value));ival=int(dec)
        if dec!=ival:raise ValueError
    except Exception:raise ValueError(f'{label}必须是整数') from None
    if ival<minimum:raise ValueError(f'{label}数值无效')
    return ival


def normalize_row(kind,raw):
    try:
        if kind=='records':
            return {'external_key':_safe_text(raw['来源唯一编号'],limit=128),'teacher_staff_no':_safe_text(raw['教师工号'],limit=64),
              'event_date':_date(raw['业务日期']),'project_code':_safe_text(raw['项目编码'],limit=64),'project':_safe_text(raw['项目名称'],limit=160),
              'course_code':_safe_text(raw['课程编码'],limit=64),'course':_safe_text(raw['课程名称'],limit=160),'class_code':_safe_text(raw['班级编码'],limit=64),
              'class_name':_safe_text(raw['班级名称'],limit=160),'hours':_decimal(raw['授课学时'],'授课学时',minimum=0),'participants':_integer(raw['课程人次'],'课程人次'),
              'income':_decimal(raw['归属培训收入'],'归属培训收入'),'source':_safe_text(raw.get('数据来源') or 'Excel导入',limit=100)}
        if kind=='qualifications':
            return {'teacher_staff_no':_safe_text(raw['教师工号'],limit=64),'category':_safe_text(raw['资质类别'],limit=80),'name':_safe_text(raw['证书名称'],limit=160),
              'certificate_no':_safe_text(raw['证书编号'],limit=128),'issuer':_safe_text(raw['发证机构'],limit=160),'obtained':_date(raw['取得日期']),
              'expires':_date(raw.get('有效期至'),required=False)}
        if kind=='teachers':
            return {'username':_safe_text(raw['登录账号'],limit=64),'name':_safe_text(raw['姓名'],limit=100),'staff_no':_safe_text(raw['教师工号'],limit=64),
              'org_code':_safe_text(raw['组织编码'],limit=64),'title':_safe_text(raw.get('职称/岗位') or '',limit=80)}
        if kind=='participations':
            return {'external_key':_safe_text(raw['来源唯一编号'],limit=128),'teacher_staff_no':_safe_text(raw['教师工号'],limit=64),
              'event_date':_date(raw['参训日期']),'project_code':_safe_text(raw['项目编码'],limit=64),'project':_safe_text(raw['项目名称'],limit=160),
              'course_code':_safe_text(raw['课程编码'],limit=64),'course':_safe_text(raw['课程名称'],limit=160),'hours':_decimal(raw['参训学时'],'参训学时',minimum=0),
              'source':_safe_text(raw.get('数据来源') or 'Excel导入',limit=100)}
    except KeyError as exc:raise ValueError('缺少固定表头字段') from exc
    raise ValueError('不支持的导入类型')


def inspect_xlsx(data,filename,kind,*,max_rows=MAX_ROWS):
    if kind not in TEMPLATES:raise HTTPException(422,'导入类型无效')
    if not filename or Path(filename).suffix.lower()!='.xlsx' or '/' in filename or '\\' in filename or len(filename)>200:
        raise HTTPException(422,'只接收标准 .xlsx 文件')
    if not data or len(data)>MAX_XLSX_BYTES:raise HTTPException(413,'Excel 文件超过10 MiB限制')
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            infos=zf.infolist()
            if len(infos)>MAX_ZIP_ENTRIES or sum(x.file_size for x in infos)>MAX_UNCOMPRESSED:raise ValueError('zip bomb')
            for x in infos:
                n=x.filename.lower()
                if x.file_size>25*1024*1024 or '..' in Path(x.filename).parts or n.endswith('.bin') or 'vbaproject' in n or '/externallinks/' in n or '/embeddings/' in n:
                    raise ValueError('active content')
    except (zipfile.BadZipFile,ValueError):raise HTTPException(422,'Excel 文件损坏、包含活动内容或异常压缩数据') from None
    try:wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except Exception:raise HTTPException(422,'Excel 工作簿无法解析') from None
    try:
        if len(wb.sheetnames)!=1:raise HTTPException(422,'导入文件只能包含一个工作表')
        ws=wb[wb.sheetnames[0]]
        rows=ws.iter_rows()
        try:first=next(rows)
        except StopIteration:raise HTTPException(422,'Excel 文件没有表头') from None
        expected=[x[0] for x in TEMPLATES[kind]]
        headers=[str(c.value).strip() if c.value is not None else '' for c in first]
        # Ignore only truly empty trailing cells, never unknown headers in between.
        while headers and headers[-1]=='':headers.pop()
        if headers!=expected:raise HTTPException(422,'Excel 表头与标准模板不一致')
        parsed=[]
        for excel_row,cells in enumerate(rows,start=2):
            vals=[];present=False;formula=False;errors=[]
            for i,c in enumerate(cells[:len(expected)]):
                if c.data_type=='f':formula=True
                value=c.value
                if value not in (None,''):present=True
                if TEMPLATES[kind][i][1] in {'staff_no','teacher_staff_no'} and c.data_type=='n' and isinstance(value,(int,float)):
                    if isinstance(value,float) and not value.is_integer():errors.append('工号必须是文本')
                    else:
                        value=str(int(value))
                        if re.fullmatch(r'0+',c.number_format):value=value.zfill(len(c.number_format))
                vals.append(_plain(value))
            if not present:continue
            if type(max_rows) is not int or not 100<=max_rows<=10000:raise HTTPException(503,'Excel导入行数策略未正确初始化')
            if len(parsed)>=max_rows:raise HTTPException(422,f'单次导入不能超过{max_rows}行')
            raw=dict(zip(expected,vals+['']*(len(expected)-len(vals))))
            if formula:errors.append('禁止公式单元格')
            try:normalized=normalize_row(kind,raw)
            except ValueError as exc:normalized={};errors.append(str(exc))
            parsed.append({'row_no':excel_row,'raw':raw,'normalized':normalized,'errors':errors})
        if not parsed:raise HTTPException(422,'Excel 没有可导入的数据行')
        return parsed
    finally:wb.close()


def template_bytes(kind):
    if kind not in TEMPLATES:raise HTTPException(404,'模板类型不存在')
    wb=Workbook();ws=wb.active;ws.title='导入数据';headers=[x[0] for x in TEMPLATES[kind]];ws.append(headers);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}1'
    fill=PatternFill('solid',fgColor='D9EAF7')
    for i,cell in enumerate(ws[1],1):
        cell.font=Font(bold=True);cell.fill=fill;cell.alignment=Alignment(horizontal='center')
        help_text='请勿修改表头；禁止公式。日期使用 YYYY-MM-DD。' if kind=='teachers' else '请勿修改表头名称或顺序；禁止公式。\n'+TEMPLATE_COLUMN_HELP[headers[i-1]]
        cell.comment=Comment(help_text,'跃科');cell.comment.width=420;cell.comment.height=140
        column=ws.column_dimensions[get_column_letter(i)];column.width=max(14,min(28,len(headers[i-1])*2+8))
        if kind!='teachers' and TEMPLATES[kind][i-1][1] in TEMPLATE_TEXT_FIELDS:
            column.number_format='@';cell.number_format='@'
    wb.properties.title=f'{TEMPLATE_NAMES[kind]}模板';wb.properties.creator='湖南跃科信息工程有限公司'
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()


def _teachers_by_staff(db,actor,staffs):
    q=select(User).where(User.staff_no.in_(staffs),User.id.in_(actor.teacher_ids))
    return {u.staff_no:u for u in db.scalars(q)}


def validate_rows(db,actor,kind,rows):
    """Return list of row results with live DB authorization and uniqueness checks."""
    out=[]
    normalized=[dict(r['normalized']) for r in rows]
    if kind in {'records','qualifications','participations'}:
        staffs={x.get('teacher_staff_no') for x in normalized if x.get('teacher_staff_no')}
        teachers=_teachers_by_staff(db,actor,staffs)
    else:teachers={}
    existing_ext=set()
    existing_qual=set()
    existing_usernames=set();existing_staff=set();orgs={};teacher_role=None
    if kind=='records':
        vals=[x.get('external_key') for x in normalized if x.get('external_key')]
        if vals:existing_ext=set(db.scalars(select(TrainingRecord.external_key).where(TrainingRecord.external_key.in_(vals))))
        record_sources=set(business_dictionary(db)['record_sources'])
    elif kind=='participations':
        vals=[x.get('external_key') for x in normalized if x.get('external_key')]
        if vals:existing_ext=set(db.scalars(select(TeacherParticipation.external_key).where(TeacherParticipation.external_key.in_(vals))))
        record_sources=set(business_dictionary(db)['record_sources'])
    elif kind=='qualifications':
        certs=[x.get('certificate_no') for x in normalized if x.get('certificate_no')]
        if certs:
            existing_qual={(t,c) for t,c in db.execute(select(User.staff_no,Qualification.certificate_no).join(User,User.id==Qualification.teacher_id).where(Qualification.certificate_no.in_(certs)))}
        setting=db.get(Setting,'qualification');categories=set(setting.value.get('categories',[])) if setting and isinstance(setting.value,dict) else set()
    elif kind=='teachers':
        actor.require_admin()
        vals=[x.get('username') for x in normalized if x.get('username')];staffs=[x.get('staff_no') for x in normalized if x.get('staff_no')]
        if vals:existing_usernames=set(db.scalars(select(User.username).where(User.username.in_(vals))))
        if staffs:existing_staff=set(db.scalars(select(User.staff_no).where(User.staff_no.in_(staffs))))
        orgs={o.code:o for o in db.scalars(select(Org).where(Org.active.is_(True)))}
        teacher_role=db.scalar(select(Role).where(Role.kind=='teacher',Role.active.is_(True)).order_by(Role.code).limit(1))
        if teacher_role is None:raise HTTPException(503,'教师角色尚未初始化')
    seen1=set();seen2=set();pending_references={}
    for source,n in zip(rows,normalized):
        errors=list(source.get('errors') or []);resolved=dict(n)
        try:
            if not n:
                raise ValueError('该行基础格式无效')
            if kind=='records':
                teacher=teachers.get(n['teacher_staff_no'])
                if not teacher or not teacher.active:errors.append('教师工号不存在、停用或超出授权范围')
                else:resolved['teacher_id']=teacher.id
                if n['external_key'] in seen1 or n['external_key'] in existing_ext:errors.append('来源唯一编号重复')
                seen1.add(n['external_key'])
                if n['source'] not in record_sources and n['source']!='Excel导入':errors.append('来源说明不在学校业务字典中，可选值：'+'、'.join(sorted(record_sources|{'Excel导入'})))
                resolved['source']='Excel导入'
                if not errors:
                    RecordIn(**{k:v for k,v in resolved.items() if k!='teacher_staff_no'},reason='Excel预览校验',basis='标准模板预览')
                    errors.extend(record_reference_errors(db,resolved,pending=pending_references))
            elif kind=='participations':
                teacher=teachers.get(n['teacher_staff_no'])
                if not teacher or not teacher.active:errors.append('教师工号不存在、停用或超出授权范围')
                else:resolved['teacher_id']=teacher.id
                if n['external_key'] in seen1 or n['external_key'] in existing_ext:errors.append('来源唯一编号重复')
                seen1.add(n['external_key'])
                if n['source'] not in record_sources and n['source']!='Excel导入':errors.append('来源说明不在学校业务字典中，可选值：'+'、'.join(sorted(record_sources|{'Excel导入'})))
                resolved['source']='Excel导入'
                if not errors:
                    ParticipationIn(**{k:v for k,v in resolved.items() if k!='teacher_staff_no'},reason='Excel预览校验',basis='标准模板预览')
                    errors.extend(record_reference_errors(db,resolved,pending=pending_references))
            elif kind=='qualifications':
                teacher=teachers.get(n['teacher_staff_no'])
                if not teacher or not teacher.active:errors.append('教师工号不存在、停用或超出授权范围')
                else:resolved['teacher_id']=teacher.id
                pair=(n['teacher_staff_no'],n['certificate_no'])
                if pair in seen1 or pair in existing_qual:errors.append('同一教师证书编号重复')
                seen1.add(pair)
                if n['category'] not in categories:errors.append('资质类别不在学校字典中')
                if not errors:
                    QualificationIn(**{k:v for k,v in resolved.items() if k!='teacher_staff_no'},reason='Excel预览校验',basis='标准模板预览')
            else:
                if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',n['username']):errors.append('登录账号格式无效')
                if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',n['staff_no']):errors.append('教师工号格式无效')
                if n['username'] in seen1 or n['username'] in existing_usernames:errors.append('登录账号重复')
                if n['staff_no'] in seen2 or n['staff_no'] in existing_staff:errors.append('教师工号重复')
                seen1.add(n['username']);seen2.add(n['staff_no'])
                org=orgs.get(n['org_code'])
                if not org or org.kind!='dept' or org.id not in actor.org_ids:errors.append('组织编码不存在、不是教研室或超出授权范围')
                elif not errors:resolved['org_id']=org.id;resolved['role_id']=teacher_role.id
        except ValidationError as exc:
            for issue in exc.errors(include_input=False,include_url=False):
                message=issue.get('ctx',{}).get('error') if issue['type']=='value_error' else None
                label=next((title for title,key in TEMPLATES[kind] if key in issue['loc']),'本行数据')
                errors.append(str(message) if message else label+'格式或取值不符合要求')
        except (ValueError,HTTPException) as exc:
            errors.append(exc.detail if isinstance(exc,HTTPException) else str(exc))
        out.append({'row_no':source['row_no'],'raw':source['raw'],'normalized':resolved,'errors':list(dict.fromkeys(errors))})
    return out


def import_preview_digest(rows):
    """Content checksum, not a signature against a database administrator."""
    return digest([{k:r.get(k) for k in ('row_no','raw','normalized','errors')} for r in rows])


def commit_import(db,actor,job,rows,reason,basis):
    expected=(job.receipt or {}).get('preview_rows_sha256')
    if (job.state!='PREVIEW' or not rows or len(rows)!=job.total
        or len({r['row_no'] for r in rows})!=len(rows)
        or not expected or expected!=import_preview_digest(rows)):
        raise HTTPException(409,'预览行数或内容校验失败（旧预览须重新上传），未写入任何正式数据')
    # Keep the all-or-nothing contract even for internal service callers.
    with db.begin_nested():
        return _commit_import(db,actor,job,rows,reason,basis,expected)


def _commit_import(db,actor,job,rows,reason,basis,preview_sha):
    proof={'access_proof':(job.receipt or {}).get('access_proof')} if (job.receipt or {}).get('access_proof') else {}
    checked=validate_rows(db,actor,job.kind,rows)
    if any(x['errors'] for x in checked):
        stored={r.row_no:r for r in db.scalars(select(ImportRow).where(ImportRow.job_id==job.id))}
        for x in checked:
            stored[x['row_no']].payload={'raw':x['raw'],'normalized':x['normalized']};stored[x['row_no']].errors=x['errors']
        job.state='INVALID';job.errors=sum(bool(x['errors']) for x in checked);job.version+=1;job.receipt={**proof,'written':0,'stale':True,'message':'确认时数据已变化，请下载错误明细，修正后重新上传；可放弃本批'}
        audit(db,actor,'imports',job.id,'确认前重新校验未通过',after={'errors':job.errors},reason=reason)
        return job
    written=[]
    if job.kind=='records':
        for x in checked:
            n=x['normalized'];payload={k:v for k,v in n.items() if k!='teacher_staff_no'}
            obj=insert_record(db,actor,RecordIn(**payload,reason=reason,basis=basis),batch_id=job.id,channel='import');ensure_reference_from_training_record(db,obj);written.append({'row_no':x['row_no'],'object_id':obj.id})
    elif job.kind=='participations':
        for x in checked:
            n=x['normalized'];payload={k:v for k,v in n.items() if k!='teacher_staff_no'}
            parsed=ParticipationIn(**payload,reason=reason,basis=basis)
            obj=TeacherParticipation(**{**parsed.model_dump(exclude={'reason','basis'}),'source':'Excel导入'},batch_id=job.id);db.add(obj);db.flush()
            workflow(db,actor,'participation',obj.id,'Excel导入教师参训记录',{},serial(obj),reason,basis);written.append({'row_no':x['row_no'],'object_id':obj.id})
    elif job.kind=='qualifications':
        for x in checked:
            n=x['normalized'];data=QualificationIn(**{k:v for k,v in n.items() if k!='teacher_staff_no'},reason=reason,basis=basis)
            obj=Qualification(**data.model_dump(exclude={'reason','basis'}),state='DRAFT');db.add(obj);db.flush();qualification_revision(db,actor,obj,reason)
            workflow(db,actor,'qualification',obj.id,'Excel导入资质草稿',{},serial(obj),reason,basis);written.append({'row_no':x['row_no'],'object_id':obj.id})
    else:
        from .routes.admin import link_school_rows
        actor.require_admin();role=db.scalar(select(Role).where(Role.kind=='teacher',Role.active.is_(True)).order_by(Role.code).limit(1))
        for x in checked:
            n=x['normalized'];obj=User(username=n['username'],name=n['name'],staff_no=n['staff_no'],title=n['title'],org_id=n['org_id'],role_id=role.id,
              password_hash=hash_password(secrets.token_urlsafe(48)),active=False,is_admin=False,must_change_password=True)
            db.add(obj);db.flush();audit(db,actor,'users',obj.id,'Excel导入教师账号（默认停用）',after=serial(obj),reason=reason);written.append({'row_no':x['row_no'],'object_id':obj.id})
            if obj.staff_no:link_school_rows(db,actor,obj)
    job.state='COMMITTED';job.errors=0;job.version+=1;job.receipt={**proof,'preview_rows_sha256':preview_sha,'written':len(written),'rows':written,'objects_sha256':digest(written),'committed_at':utcnow().isoformat()}
    audit(db,actor,'imports',job.id,'确认Excel整批入库',after={'kind':job.kind,'written':len(written),'sha256':job.sha256},reason=reason)
    return job


def rows_to_xlsx(headers,rows,title='导出数据'):
    return sheets_to_xlsx([(title,headers,rows)])


def sheets_to_xlsx(sheets):
    """sheets: [(title, headers, rows), ...] -> one workbook, one sheet each."""
    wb=Workbook(write_only=False);wb.remove(wb.active)
    for title,headers,rows in sheets:
        ws=wb.create_sheet(title[:31]);ws.append(headers);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}1'
        for cell in ws[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='D9EAF7')
        for row in rows:ws.append([_safe_excel_value(x) for x in row])
        for i,h in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width=max(12,min(36,len(h)*2+6))
            if h.startswith('实施人员使用：'):ws.column_dimensions[get_column_letter(i)].hidden=True
            if title=='教师积分' and h in ('项目编号','项目名称'):
                ws.column_dimensions[get_column_letter(i)].width=36 if h=='项目编号' else 54
        if title=='教师积分':ws.auto_filter.ref=ws.dimensions
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()


def import_error_text(value):
    text=str(value)
    if re.match(r'^\d+ validation errors? for ',text):return '该行数据校验未通过，请核对日期、数值和必填项后重新上传'
    return text


def error_xlsx(job_rows):
    # Raw headers are identical because preview requires the exact template.
    if not job_rows:return rows_to_xlsx(['错误原因'],[],'错误行')
    headers=list((job_rows[0].payload or {}).get('raw',{}).keys())+['错误原因']
    rows=[]
    for r in job_rows:
        if not r.errors:continue
        raw=(r.payload or {}).get('raw',{});rows.append([raw.get(h,'') for h in headers[:-1]]+['；'.join(import_error_text(value) for value in r.errors)])
    return rows_to_xlsx(headers,rows,'错误行')


def require_export_scope(actor,kind):
    if kind=='handover':
        actor.require('data.export');actor.require_admin();return
    required={'records':'records.read','participations':'records.read','qualifications':'qualification.read',
        'points':'points.read','assessments':'assessment.read','organization_assessments':'assessment.read','audit':'audit.read'}
    if kind not in required:raise HTTPException(422,'导出类型无效')
    actor.require('export');actor.require(required[kind])
    if kind=='audit':actor.require_admin()


def export_rows(db,actor,kind,filters,policy=None):
    require_export_scope(actor,kind)
    from .sensitive import MaskPolicy
    pol=policy or MaskPolicy(db,actor);staff=lambda u:pol.value('staff_no',u.staff_no,u.id)
    labels=business_dictionary(db)['status_labels'];label=lambda code:labels.get(code,code)  # 3.4: same display names as the UI
    if kind=='records':
        q=record_query(db,actor,filters)
        items=db.scalars(q.order_by(TrainingRecord.event_date.desc(),TrainingRecord.id)).all();users={u.id:u for u in db.scalars(select(User).where(User.id.in_({x.teacher_id for x in items})))}
        return ['业务日期','来源唯一编号','教师工号','教师姓名','项目编码','项目名称','课程编码','课程名称','班级编码','班级名称','学时','人次','归属培训收入','数据来源'],[
            [x.event_date,x.external_key,staff(users[x.teacher_id]),users[x.teacher_id].name,x.project_code,x.project,x.course_code,x.course,x.class_code,x.class_name,x.hours,x.participants,x.income,x.source] for x in items]
    if kind=='participations':
        allowed={'q','teacher_id','org_id','start','end','project_code','course_code'}
        if set(filters)-allowed:raise HTTPException(422,'参训导出筛选字段无效')
        from .routes.reference import _participation_query
        items=db.scalars(_participation_query(db,actor,filters).order_by(TeacherParticipation.event_date.desc(),TeacherParticipation.id)).all();users={u.id:u for u in db.scalars(select(User).where(User.id.in_({x.teacher_id for x in items})))}
        return ['参训日期','来源唯一编号','教师工号','教师姓名','项目编码','项目名称','课程编码','课程名称','参训学时','数据来源'],[[x.event_date,x.external_key,staff(users[x.teacher_id]),users[x.teacher_id].name,x.project_code,x.project,x.course_code,x.course,x.hours,x.source] for x in items]
    if kind=='qualifications':
        allowed={'q','state','teacher_id','org_id'}
        if set(filters)-allowed:raise HTTPException(422,'资质导出筛选字段无效')
        q=select(Qualification).where(actor.scope(Qualification.teacher_id))
        if filters.get('teacher_id'):q=q.where(Qualification.teacher_id==filters['teacher_id'])
        if filters.get('org_id'):q=q.where(Qualification.teacher_id.in_(select(User.id).where(User.org_id.in_(authorized_org_descendants(db,actor,filters['org_id'])))))
        if filters.get('state'):q=q.where(Qualification.state==filters['state'])
        if filters.get('q'):q=q.where(Qualification.name.contains(filters['q'],autoescape=True)|Qualification.certificate_no.contains(filters['q'],autoescape=True))
        items=db.scalars(q.order_by(Qualification.created_at.desc(),Qualification.id)).all();users={u.id:u for u in db.scalars(select(User).where(User.id.in_({x.teacher_id for x in items})))}
        return ['教师工号','教师姓名','资质类别','证书名称','证书编号','发证机构','取得日期','有效期至','状态','版本'],[[staff(users[x.teacher_id]),users[x.teacher_id].name,x.category,x.name,pol.value('certificate_no',x.certificate_no,x.teacher_id),x.issuer,x.obtained,x.expires or '',label(x.state),x.version] for x in items]
    if kind=='points':
        allowed={'teacher_id','period_id','org_id'}
        if set(filters)-allowed:raise HTTPException(422,'积分导出筛选字段无效')
        q=select(PointEntry).where(actor.scope(PointEntry.teacher_id))
        for k in ('teacher_id','period_id'):
            if filters.get(k):q=q.where(getattr(PointEntry,k)==filters[k])
        if filters.get('org_id'):q=q.where(PointEntry.teacher_id.in_(select(User.id).where(User.org_id.in_(authorized_org_descendants(db,actor,filters['org_id'])))))
        items=db.scalars(q.order_by(PointEntry.created_at,PointEntry.id)).all();users={u.id:u for u in db.scalars(select(User).where(User.id.in_({x.teacher_id for x in items})))}
        periods=dict(db.execute(select(Period.id,Period.name).where(Period.id.in_({x.period_id for x in items}))).all())
        claims={x.source_snapshot.get('social_claim') for x in items if x.source_snapshot.get('social_claim')}
        projects=dict(db.execute(select(SocialClaim.id,SocialProject.data).join(SocialProject,SocialProject.id==SocialClaim.project_id).where(SocialClaim.id.in_(claims))).all())
        def project_labels(x):
            data=projects.get(x.source_snapshot.get('social_claim'),{})
            return [data.get('code',''),data.get('name','')]
        def point_name(x):
            if x.kind=='SOCIAL':
                return '社会培训'+('辅助' if 'auxiliary' in x.item_key else '授课')+'积分'+('更正' if x.source_snapshot.get('correction_id') else '')
            components=((x.source_snapshot.get('rule') or {}).get('config') or {}).get('components') or []
            return next((v['label'] for v in components if v.get('key')==x.item_key and v.get('label')),None) or {'hours':'授课贡献','participants':'课程人次','income':'收入贡献','manual_review':'人工复核调整'}.get(x.item_key,'积分调整项')
        return ['教师工号','教师姓名','认定周期','积分项','类型','分值','原始值','公式','实施人员使用：规则编号','实施人员使用：批次编号','时间（北京时间）','原因','依据','实施人员使用：周期编号','实施人员使用：积分项编号','实施人员使用：原始时间（世界协调时）','项目编号','项目名称'],[[staff(users[x.teacher_id]),users[x.teacher_id].name,periods.get(x.period_id,'未命名周期'),point_name(x),label(x.kind),x.amount,x.raw_value,x.formula,x.rule_id or '',x.batch_id or '',x.created_at+timedelta(hours=8),x.reason,x.basis,x.period_id,x.item_key,x.created_at,*project_labels(x)] for x in items]
    if kind=='assessments':
        if set(filters)-ASSESSMENT_FILTERS:raise HTTPException(422,'考核导出筛选字段无效')
        from .assessment_evidence import assessment_projection
        items=db.scalars(assessment_query(db,actor,filters).order_by(Assessment.score.desc(),Assessment.teacher_id)).all()
        rows=[]
        for obj in items:
            d=assessment_projection(db,obj);x=d['assessment'];u=d['teacher']
            rows.append([pol.value('staff_no',u['staff_no'],obj.teacher_id),u['name'],d['period']['name'],Decimal(str(x['score'])),label(x['state']),'是' if x['returned'] else '否',x['rule_id'],x['version'],obj.archive_sha256 or '',x['period_id']])
        return ['教师工号','教师姓名','考核周期','得分','状态','退回','实施人员使用：规则编号','版本','实施人员使用：归档校验值','实施人员使用：周期编号'],rows
    if kind=='organization_assessments':
        if actor.role.kind=='teacher':raise HTTPException(403,'教师不能导出组织考核')
        allowed={'period_id','target_kind','target_id','state','progress'}
        if set(filters)-allowed:raise HTTPException(422,'组织考核导出筛选字段无效')
        q=select(OrganizationAssessment)
        if actor.role.kind!='school':q=q.where(OrganizationAssessment.target_kind==actor.role.kind,OrganizationAssessment.target_id==actor.org.id)
        for k in ('period_id','target_kind','target_id','state'):
            if filters.get(k):q=q.where(getattr(OrganizationAssessment,k)==filters[k])
        from .routes.calculations import _organization_details,org_returned_ids,org_progress_clause
        if filters.get('progress'):
            generated=db.scalars(q.where(OrganizationAssessment.state=='GENERATED').with_only_columns(OrganizationAssessment.id)).all()
            q=q.where(org_progress_clause(actor,filters['progress'],org_returned_ids(db,generated)))
        items=db.scalars(q.order_by(OrganizationAssessment.created_at.desc(),OrganizationAssessment.id)).all()
        rows=[]
        for obj in items:
            d=serial(_organization_details(db,actor,obj.id));x=d['assessment']
            rows.append([d['target']['name'],{'dept':'教研室','college':'学院','school':'学校'}.get(x['target_kind'],x['target_kind']),d['period']['name'],Decimal(str(x['score'])),label(x['state']),x['rule_id'],x['version'],obj.archive_sha256 or '',x['period_id']])
        return ['考核对象','对象层级','考核周期','得分','状态','实施人员使用：规则编号','版本','实施人员使用：归档校验值','实施人员使用：周期编号'],rows
    if kind=='audit':
        q=audit_query(db,actor,filters)
        items=db.scalars(q.order_by(Audit.created_at.desc(),Audit.id)).all()
        return ['时间（北京时间）','用户','组织ID','模块','业务对象','操作','前值','后值','原因','结果','Request-ID'],[[x.created_at+timedelta(hours=8),x.actor_name,x.org_id or '',x.module,x.object_id,x.action,json.dumps(x.before,ensure_ascii=False),json.dumps(x.after,ensure_ascii=False),x.reason,label(x.result),x.request_id] for x in items]
    raise HTTPException(422,'导出类型无效')


def actor_for_task(db,task):
    user=db.get(User,task.owner_id)
    if not user:raise RuntimeError('导出任务用户不存在')
    role,org=db.get(Role,user.role_id),db.get(Org,user.org_id)
    actor=Actor(db,user,role,org,None)
    if actor.fingerprint()!=task.payload.get('authorization'):raise RuntimeError('授权范围已变化，请重新创建导出')
    return actor


def _ensure_export_root(storage_dir):
    root=Path(storage_dir)/'.exports';root.mkdir(mode=0o700,parents=True,exist_ok=True)
    if root.is_symlink() or not root.is_dir():raise RuntimeError('导出存储目录异常')
    os.chmod(root,0o700);return root


def _atomic_write(path,data):
    from .safe_files import open_nofollow
    tmp=path.with_name('.'+path.name+'.'+secrets.token_hex(6)+'.tmp')
    fd=open_nofollow(tmp,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
    try:
        with os.fdopen(fd,'wb') as f:f.write(data);f.flush();os.fsync(f.fileno())
        os.replace(tmp,path);os.chmod(path,0o600)
    finally:
        try:tmp.unlink()
        except OSError:pass


def process_export_task(db,task,storage_dir):
    kind=task.payload.get('kind');filters=task.payload.get('filters') or {}
    task.attempts+=1;attempt=task.attempts
    try:
        actor=actor_for_task(db,task)
        if kind=='handover':
            actor.require('data.export');actor.require_admin();path,filename=_handover_snapshot(db,task,storage_dir)
        else:
            from .sensitive import MaskPolicy
            policy=MaskPolicy(db,actor);actor.require('export');headers,rows=export_rows(db,actor,kind,filters,policy)
            title={'records':'培训授课数据','participations':'教师参训记录','qualifications':'教师资质','points':'教师积分','assessments':'教师考核','organization_assessments':'组织考核','audit':'审计记录'}[kind]
            data=rows_to_xlsx(headers,rows,title);root=_ensure_export_root(storage_dir);filename=f'{title}_{task.id[:8]}.xlsx';path=root/(task.id+'.xlsx');_atomic_write(path,data)
        data_hash=sha256(path.read_bytes()).hexdigest();task.state='DONE';task.error='';task.result={'filename':filename,'storage_name':path.name,'sha256':data_hash,'size':path.stat().st_size,'mime':'application/zip' if kind=='handover' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','report_version':EXPORT_FORMAT_VERSION,'generated_at':utcnow().isoformat()};task.version+=1
        if kind!='handover':task.result.update(sensitive_masked=policy.masked,rows=len(rows))
        db.add(TaskAttempt(task_id=task.id,attempt=attempt,result='SUCCESS',message='导出生成并完成SHA256校验'));return task
    except Exception as exc:
        task.state='FAILED';task.error=('导出校验未通过，请核对权限与数据完整性' if isinstance(exc,HTTPException) else '导出执行失败：'+type(exc).__name__);task.version+=1;db.add(TaskAttempt(task_id=task.id,attempt=attempt,result='ERROR',message=type(exc).__name__))
        return task


def _csv_bytes(headers,rows):
    out=StringIO(newline='');w=csv.writer(out);w.writerow(headers)
    for row in rows:w.writerow([_safe_excel_value(x) for x in row])
    return ('\ufeff'+out.getvalue()).encode('utf-8')


def _handover_snapshot(db, task, storage_dir):
    # A separate read-only transaction gives every exported table the same view.
    # SQLite requires explicit BEGIN; MySQL uses REPEATABLE READ. Neither is a
    # claim that the live school's MySQL deployment has passed acceptance.
    engine=db.get_bind()
    with engine.connect() as conn:
        if conn.dialect.name == 'mysql':
            conn=conn.execution_options(isolation_level='REPEATABLE READ')
            conn.exec_driver_sql('START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY')
        elif conn.dialect.name == 'sqlite':
            conn.exec_driver_sql('BEGIN')
        else:
            raise RuntimeError('开放移交仅支持已定义的 MySQL/隔离 SQLite 快照协议')
        try:
            with Session(bind=conn,expire_on_commit=False) as snapshot:
                actor=actor_for_task(snapshot,task);actor.require('data.export');actor.require_admin()
                return _handover_zip(snapshot,actor,task,storage_dir)
        finally:
            conn.rollback()


def _handover_zip(db,actor,task,storage_dir):
    root=_ensure_export_root(storage_dir);path=root/(task.id+'.zip');tmp=path.with_name('.'+task.id+'.tmp.zip')
    manifest={'format':'Yueke NonDegree open handover v2','generated_at':utcnow().isoformat()+'Z',
        'files':[],'tables':[],'data_encoding':'UTF-8 JSONL: SQL Numeric/Date/DateTime use exact decimal/ISO strings; null/bool/int/JSON retain type',
        'consistency':'single_read_only_database_snapshot','disaster_restore_proven':False,
        'excluded':['password_hash','sessions','login throttles','command receipts','external secret values','installation provenance','worker runtime leases']}
    def add_bytes(z,name,data):
        z.writestr(name,data);manifest['files'].append({'path':name,'size':len(data),'sha256':sha256(data).hexdigest()})
    tables=[
      ('organizations',Org,['id','code','name','kind','parent_id','active','version','created_at']),('roles',Role,['id','code','name','kind','permissions','active','version','created_at']),
      ('users',User,['id','username','name','staff_no','title','org_id','role_id','active','is_admin','version','created_at']),('periods',Period,['id','name','start_date','end_date','state','version','created_at']),
      ('settings',Setting,['key','value','version']),('training_records',TrainingRecord,[c.key for c in sa_inspect(TrainingRecord).columns]),
      *[(model.__tablename__,model,[c.key for c in sa_inspect(model).columns]) for model in (SocialProject,SocialUnit,SocialPolicy,SocialClaim,SocialAnnual,SocialResource)],
      ('qualifications',Qualification,[c.key for c in sa_inspect(Qualification).columns]),('qualification_revisions',QualificationRevision,[c.key for c in sa_inspect(QualificationRevision).columns]),
      ('rules',Rule,[c.key for c in sa_inspect(Rule).columns]),('point_balances',PointBalance,[c.key for c in sa_inspect(PointBalance).columns]),('point_entries',PointEntry,[c.key for c in sa_inspect(PointEntry).columns]),
      ('assessments',Assessment,[c.key for c in sa_inspect(Assessment).columns]),('workflow_events',WorkflowEvent,[c.key for c in sa_inspect(WorkflowEvent).columns]),
      ('organization_assessment_rules',OrganizationAssessmentRule,[c.key for c in sa_inspect(OrganizationAssessmentRule).columns]),('organization_assessments',OrganizationAssessment,[c.key for c in sa_inspect(OrganizationAssessment).columns]),
      ('organization_assessment_metric_evidence',OrganizationAssessmentMetricEvidence,[c.key for c in sa_inspect(OrganizationAssessmentMetricEvidence).columns]),
      ('teacher_hr_profiles',TeacherHrProfile,[c.key for c in sa_inspect(TeacherHrProfile).columns]),('training_projects',TrainingProjectMaster,[c.key for c in sa_inspect(TrainingProjectMaster).columns]),
      ('training_courses',TrainingCourseMaster,[c.key for c in sa_inspect(TrainingCourseMaster).columns]),('training_classes',TrainingClassMaster,[c.key for c in sa_inspect(TrainingClassMaster).columns]),
      ('teacher_participations',TeacherParticipation,[c.key for c in sa_inspect(TeacherParticipation).columns]),
      ('teacher_course_terms',TeacherCourseTerm,[c.key for c in sa_inspect(TeacherCourseTerm).columns]),
      ('school_courses',SchoolCourse,[c.key for c in sa_inspect(SchoolCourse).columns]),
      ('inbound_batches',InboundBatch,[c.key for c in sa_inspect(InboundBatch).columns]),('inbound_rows',InboundRow,[c.key for c in sa_inspect(InboundRow).columns]),
      ('integration_schedules',IntegrationSchedule,[c.key for c in sa_inspect(IntegrationSchedule).columns]),
      ('sync_items',SyncItem,[c.key for c in sa_inspect(SyncItem).columns]),('sync_call_logs',SyncCallLog,[c.key for c in sa_inspect(SyncCallLog).columns]),
      ('record_revisions',RecordRevision,[c.key for c in sa_inspect(RecordRevision).columns]),
      ('calculation_batches',CalculationBatch,[c.key for c in sa_inspect(CalculationBatch).columns]),
      ('rule_trials',RuleTrial,[c.key for c in sa_inspect(RuleTrial).columns]),
      ('organization_rule_trials',OrganizationRuleTrial,[c.key for c in sa_inspect(OrganizationRuleTrial).columns]),
      ('disputes',Dispute,[c.key for c in sa_inspect(Dispute).columns]),
      ('import_jobs',ImportJob,[c.key for c in sa_inspect(ImportJob).columns]),
      ('import_rows',ImportRow,[c.key for c in sa_inspect(ImportRow).columns]),
      ('tasks',Task,[c.key for c in sa_inspect(Task).columns]),
      ('task_attempts',TaskAttempt,[c.key for c in sa_inspect(TaskAttempt).columns]),
      ('notices',Notice,[c.key for c in sa_inspect(Notice).columns]),
      ('connection_checks',ConnectionCheck,[c.key for c in sa_inspect(ConnectionCheck).columns]),
      ('attachments',Attachment,[c.key for c in sa_inspect(Attachment).columns]),
      ('connections',Connection,[c.key for c in sa_inspect(Connection).columns]),('audit_events',Audit,[c.key for c in sa_inspect(Audit).columns]),
    ]
    try:
        with zipfile.ZipFile(tmp,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
            for name,model,columns in tables:
                safe=[c for c in columns if c not in {'password_hash','auth_version','failures','locked_until','storage_key','token_hash','csrf','lease_token','lease_until'}]
                items=db.scalars(select(model).order_by(*[getattr(model,k) for k in ('created_at','id') if hasattr(model,k)])).all()
                # Validate sealed results before allowing them into portable evidence.
                if model is Assessment:
                    from .assessment_evidence import verified_archive
                    for obj in items:
                        if obj.state=='ARCHIVED':verified_archive(obj)
                if model is OrganizationAssessment:
                    from .org_evidence import verified_org_archive
                    for obj in items:
                        if obj.state=='ARCHIVED':verified_org_archive(obj)
                typed=[{c:serial(getattr(obj,c)) for c in safe} for obj in items]
                rows=[[json.dumps(row[c],ensure_ascii=False,separators=(',',':')) if isinstance(row[c],(dict,list)) else row[c]
                    for c in safe] for row in typed]
                csv_path=f'data/{name}.csv';typed_path=f'typed/{name}.jsonl'
                add_bytes(z,csv_path,_csv_bytes(safe,rows))
                add_bytes(z,typed_path,(''.join(json.dumps(row,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\n' for row in typed)).encode('utf-8'))
                manifest['tables'].append({'name':name,'model_table':model.__tablename__,'columns':safe,'rows':len(typed),
                    'csv_path':csv_path,'typed_path':typed_path,
                    'primary_key':[c.key for c in sa_inspect(model).primary_key],
                    'column_types':{c:str(sa_inspect(model).columns[c].type) for c in safe}})
            structure={t.name:{'columns':[{'name':c.name,'type':str(c.type),'nullable':c.nullable,'primary_key':c.primary_key,'foreign_keys':[str(fk.target_fullname) for fk in c.foreign_keys]} for c in t.columns]} for t in Base.metadata.sorted_tables}
            add_bytes(z,'database_structure.json',json.dumps(structure,ensure_ascii=False,indent=2).encode())
            add_bytes(z,'delivery/initialization_configuration.json',json.dumps(configuration_snapshot(db),ensure_ascii=False,indent=2).encode())
            add_bytes(z,'delivery/database_dictionary.json',json.dumps(database_dictionary(),ensure_ascii=False,indent=2).encode())
            add_bytes(z,'delivery/interface_data_catalog.json',json.dumps(DATA_CATALOG,ensure_ascii=False,indent=2).encode())
            attachments=[]
            for item in db.scalars(select(Attachment).order_by(Attachment.id)):
                from .routes.files import verify_file
                data=verify_file(storage_dir,item)
                ext = Path(item.name).suffix.lower()
                safe_ext = ext if re.fullmatch(r'\.[a-z0-9]{1,8}', ext) else ''
                arc = f'attachments/{item.id}{safe_ext}'
                add_bytes(z,arc,data);attachments.append({'id':item.id,'original_name':item.name,'archive_path':arc,'sha256':item.sha256,'size':item.size,'object_type':item.object_type,'object_id':item.object_id,'teacher_id':item.teacher_id,'uploaded_by':item.uploaded_by,'mime':item.mime,'created_at':item.created_at.isoformat()})
            add_bytes(z,'attachments_manifest.json',json.dumps(attachments,ensure_ascii=False,indent=2).encode())
            add_bytes(z,'README.txt','UTF-8 CSV（电子表格安全显示）+ typed/*.jsonl（精确值与JSON结构）+ 原始附件。\nCSV可能用单引号防公式，数据迁移请读取JSONL；SQL金额/日期/时间为精确字符串，按manifest column_types解释。\n不含密码、会话、命令回执、安装来源凭证或外部接口密钥；本包不是灾备文件，禁止直接恢复成可用安装。\n'.encode('utf-8'))
            # Manifest intentionally does not hash itself.
            z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2).encode())
        os.replace(tmp,path);os.chmod(path,0o600);return path,f'学校业务开放格式移交_{task.id}.zip'
    finally:
        try:tmp.unlink()
        except FileNotFoundError:pass


def read_export_file(storage_dir,task):
    from .safe_files import open_nofollow
    result=task.result or {};name=result.get('storage_name')
    if not name or not re.fullmatch(r'[A-Za-z0-9_.-]{1,120}',name):raise HTTPException(404,'导出文件不存在')
    root=_ensure_export_root(storage_dir);path=root/name
    try:
        fd=open_nofollow(path,os.O_RDONLY)
        with os.fdopen(fd,'rb') as f:meta=os.fstat(f.fileno());data=f.read(meta.st_size+1)
    except OSError:raise HTTPException(404,'导出文件不存在') from None
    if not stat.S_ISREG(meta.st_mode) or len(data)!=result.get('size') or sha256(data).hexdigest()!=result.get('sha256'):
        raise HTTPException(503,'导出文件校验失败，请重新生成')
    return data
