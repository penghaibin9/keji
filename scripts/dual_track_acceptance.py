#!/usr/bin/env python3
"""Verify school dual-track point/assessment samples required by procurement.

The workbook must contain two sheets: 积分双轨 and 考核双轨. This verifier does
not create or infer school rules. It only compares a school's manually supplied
result with the system result after both have been calculated under the signed-off
rule version.
"""
from __future__ import annotations
import argparse,json
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

SHEETS={'积分双轨':'points','考核双轨':'assessment'}
HEADERS=['样本编号','对象标识','周期','系统结果','人工结果','规则版本','学校确认人','备注']

def amount(v):
    try:return Decimal(str(v)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
    except (InvalidOperation,ValueError,TypeError):raise ValueError('结果必须是可比较的数字') from None

def verify(path,signed_off=False):
    path=Path(path);errors=[];detail={}
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        for sheet,key in SHEETS.items():
            if sheet not in wb.sheetnames:
                errors.append('缺少工作表：'+sheet);continue
            ws=wb[sheet];headers=[ws.cell(1,i).value for i in range(1,len(HEADERS)+1)]
            if headers!=HEADERS:
                errors.append(sheet+' 表头不匹配');continue
            rows=[];ids=set()
            for row_no,values in enumerate(ws.iter_rows(min_row=2,max_col=len(HEADERS),values_only=True),2):
                if all(v in (None,'') for v in values):continue
                sample=str(values[0] or '').strip();obj=str(values[1] or '').strip();period=str(values[2] or '').strip();rule=str(values[5] or '').strip();signer=str(values[6] or '').strip()
                if not sample or not obj or not period or not rule:
                    errors.append(f'{sheet} 第{row_no}行缺少样本编号/对象/周期/规则版本');continue
                if sample in ids:errors.append(f'{sheet} 样本编号重复：{sample}');continue
                ids.add(sample)
                try:system,manual=amount(values[3]),amount(values[4])
                except ValueError as exc:errors.append(f'{sheet} 第{row_no}行：{exc}');continue
                rows.append({'sample_id':sample,'object':obj,'period':period,'rule_version':rule,'system':str(system),'manual':str(manual),'match':system==manual,'school_signer':signer})
            if not rows:errors.append(sheet+' 没有可验收样本')
            matched=sum(r['match'] for r in rows);total=len(rows)
            detail[key]={'total':total,'matched':matched,'match_rate':(matched/total if total else 0),'mismatches':[r for r in rows if not r['match']],
                         'all_rows_have_school_signer':bool(rows) and all(r['school_signer'] for r in rows)}
    finally:wb.close()
    points=detail.get('points',{});assessment=detail.get('assessment',{})
    points_rate=points.get('match_rate',0);assessment_rate=assessment.get('match_rate',0)
    signers_ok=points.get('all_rows_have_school_signer') is True and assessment.get('all_rows_have_school_signer') is True
    effective_signed=bool(signed_off and signers_ok and not errors)
    passed=not errors and points_rate==1 and assessment_rate==1 and effective_signed
    return {'status':'PASS' if passed else 'BLOCKED','points_match_rate':points_rate,'assessment_match_rate':assessment_rate,
            'signed_off':effective_signed,'requested_signed_off':bool(signed_off),'school_signers_complete':signers_ok,
            'errors':errors,'detail':detail,'note':'Only real school representative samples and school sign-off may make this evidence PASS.'}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('workbook');ap.add_argument('--signed-off',action='store_true');ap.add_argument('--output');a=ap.parse_args()
    result=verify(a.workbook,a.signed_off);payload=json.dumps(result,ensure_ascii=False,indent=2)
    if a.output:Path(a.output).write_text(payload+'\n',encoding='utf-8')
    print(payload);raise SystemExit(0 if result['status']=='PASS' else 2)
if __name__=='__main__':main()
